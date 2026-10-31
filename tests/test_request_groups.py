import os
import tempfile
import unittest
from io import BytesIO
from unittest import mock

from openpyxl import load_workbook

_database_dir = tempfile.TemporaryDirectory()
os.environ.setdefault("TRAFFICLENS_DB", os.path.join(_database_dir.name, "groups.db"))
os.environ.setdefault("OUTSCRAPER_API_KEY", "unit-test-api-key")
os.environ.setdefault("SECRET_KEY", "unit-test-session-secret-that-is-long-enough")
os.environ.setdefault("TRAFFICLENS_OUTSCRAPER_AUTO_POLL", "0")

import app as traffic


def link(offset=0.0, with_c=True):
    payload = {
        "coords": [[30.138239 + offset, 31.395698], [30.137423 + offset, 31.393946]],
        "path": [[30.138239 + offset, 31.395698], [30.137423 + offset, 31.393946]],
        "road_km": 0.19, "junctions": 0, "no_junctions": True,
    }
    if with_c:
        payload.update(point_c=[30.129997 + offset, 31.383281], point_c_distance_km=1.7)
    return payload


PERIODS = [
    {"label": "AM", "start_time": "07:15:00", "end_time": "08:15:00"},
    {"label": "MD", "start_time": "12:00:00", "end_time": "13:00:00"},
]


class RequestGroupTests(unittest.TestCase):
    def setUp(self):
        with traffic.app.app_context():
            d = traffic.db()
            for name, email, role in (("Group admin", "group-admin@example.test", "admin"),
                                      ("Group client", "group-client@example.test", "client"),
                                      ("Other client", "group-other@example.test", "client")):
                d.execute("INSERT OR IGNORE INTO users(name,email,pw,role) VALUES(?,?,'x',?)", (name, email, role))
            ids = {r["email"]: r["id"] for r in d.execute("SELECT id,email FROM users WHERE email LIKE 'group-%'")}
            d.execute("UPDATE periods SET is_active=1")
            d.commit()
        self.admin_id, self.client_id, self.other_id = (
            ids["group-admin@example.test"], ids["group-client@example.test"], ids["group-other@example.test"])

    def login(self, uid, role, name):
        client = traffic.app.test_client()
        with client.session_transaction() as session:
            session.update(uid=uid, role=role, name=name)
        return client

    def submit(self, client, **overrides):
        body = {"mode": "multi_link", "services": ["travel"], "date": "2026-10-24", "title": "Junction",
                "links": [link(0), link(0.01), link(0.02)], "periods": [PERIODS[0]]}
        body.update(overrides)
        with mock.patch.object(traffic, "apply_network", return_value=None):
            return client.post("/api/request-groups", json=body)

    def test_links_times_periods_become_one_group_of_ordinary_requests(self):
        client = self.login(self.client_id, "client", "Group client")
        response = self.submit(client, mode="multi_period", links=[link()], periods=PERIODS)
        self.assertEqual(response.status_code, 200, response.get_json())
        data = response.get_json()
        self.assertEqual(data["count"], 2)
        with traffic.app.app_context():
            rows = traffic.db().execute(
                "SELECT period_label,start_time,end_time,group_id FROM requests WHERE group_id=? ORDER BY id",
                (data["group_id"],)).fetchall()
            admin_notes = traffic.db().execute(
                "SELECT COUNT(*) FROM notifications WHERE message LIKE ?", (f"%GRP-{data['group_id']}%",)).fetchone()[0]
            admin_count = traffic.db().execute("SELECT COUNT(*) FROM users WHERE role='admin'").fetchone()[0]
        self.assertEqual([(r["period_label"], r["start_time"], r["end_time"]) for r in rows],
                         [("AM", "07:15:00", "08:15:00"), ("MD", "12:00:00", "13:00:00")])
        self.assertEqual(admin_notes, admin_count)   # one notification per admin, not one per request

    def test_single_link_single_period_stays_a_plain_request(self):
        client = self.login(self.client_id, "client", "Group client")
        response = self.submit(client, mode="single", links=[link()], periods=[PERIODS[0]])
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertIsNone(response.get_json()["group_id"])
        with traffic.app.app_context():
            row = traffic.db().execute("SELECT group_id,period_label FROM requests WHERE id=?",
                                       (response.get_json()["id"],)).fetchone()
        self.assertIsNone(row["group_id"]); self.assertEqual(row["period_label"], "AM")

    def test_server_enforces_limits_and_mode_rules(self):
        client = self.login(self.client_id, "client", "Group client")
        self.assertEqual(self.submit(client, links=[link(i / 100) for i in range(11)]).status_code, 400)
        self.assertEqual(self.submit(client, mode="single", links=[link(), link(0.01)]).status_code, 400)
        self.assertEqual(self.submit(client, mode="multi_link", periods=PERIODS).status_code, 400)
        self.assertEqual(self.submit(client, mode="multi_period", links=[link(), link(0.01)], periods=PERIODS).status_code, 400)
        self.assertEqual(self.submit(client, mode="multi_period", links=[link()], periods=[PERIODS[0], PERIODS[0]]).status_code, 400)
        bad_time = [{"label": "AM", "start_time": "09:00:00", "end_time": "08:00:00"}]
        self.assertEqual(self.submit(client, mode="single", links=[link()], periods=bad_time).status_code, 400)
        self.assertEqual(self.submit(client, mode="single", links=[link()], periods=[{"label": "XX", "start_time": "07:00:00", "end_time": "08:00:00"}]).status_code, 400)
        self.assertEqual(self.submit(client, links=[link(with_c=False)]).status_code, 400)   # travel needs point C
        self.assertEqual(self.submit(client, mode="bogus").status_code, 400)

    def test_client_sees_one_report_only_after_admin_sends_the_group(self):
        client = self.login(self.client_id, "client", "Group client")
        admin = self.login(self.admin_id, "admin", "Group admin")
        data = self.submit(client).get_json()
        gid, ids = data["group_id"], data["request_ids"]
        self.assertEqual(len(ids), 3)
        # a child result is not reachable by the client before the group is sent
        self.assertEqual(client.get(f"/results/{ids[0]}").status_code, 302)
        self.assertEqual(client.get(f"/groups/{gid}/export.xlsx").status_code, 404)
        self.assertEqual(client.get(f"/groups/{gid}").status_code, 200)
        admin_page = admin.get(f"/admin/groups/{gid}")
        self.assertEqual(admin_page.status_code, 200)
        self.assertIn(b"Send one report to client", admin_page.data)
        self.assertEqual(admin.get(f"/groups/{gid}").status_code, 302)
        self.assertIn(f"GRP-{gid}".encode(), client.get("/requests").data)
        self.assertNotIn(f"REQ-{ids[0]}".encode(), client.get("/requests").data)
        self.assertEqual(self.login(self.other_id, "client", "Other client").get(f"/groups/{gid}").status_code, 404)
        # the admin cannot send while a request is still open
        self.assertEqual(admin.post(f"/api/groups/{gid}/send").status_code, 409)
        with traffic.app.app_context():
            d = traffic.db()
            for rid in ids:
                d.execute("UPDATE requests SET status='Completed', data=? WHERE id=?",
                          ('{"travel":{"rows":[{"value":42.5}],"note":"ok","source":"admin"}}', rid))
            d.commit()
        self.assertEqual(admin.post(f"/api/groups/{gid}/send").status_code, 200)
        self.assertEqual(admin.post(f"/api/groups/{gid}/send").status_code, 409)
        with traffic.app.app_context():
            notes = traffic.db().execute(
                "SELECT kind FROM notifications WHERE user_id=? AND message LIKE ?",
                (self.client_id, f"%GRP-{gid}%")).fetchall()
        self.assertEqual([n["kind"] for n in notes], ["group_ready"])
        self.assertEqual(client.get(f"/results/{ids[0]}").status_code, 200)
        self.assertIn(b"Download Excel report", client.get(f"/groups/{gid}").data)
        report = client.get(f"/groups/{gid}/export.xlsx")
        self.assertEqual(report.status_code, 200)
        workbook = load_workbook(BytesIO(report.data))
        self.assertEqual(workbook.sheetnames[:3], ["Order Summary", "Links and Periods", "Service Details"])
        values = [[c for c in row] for row in workbook["Service Details"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(len(values), 3)
        self.assertEqual({row[0] for row in values}, {"Link 1", "Link 2", "Link 3"})
        self.assertEqual({row[5] for row in values}, {42.5})

    def test_client_is_not_notified_about_each_child(self):
        client = self.login(self.client_id, "client", "Group client")
        data = self.submit(client).get_json()
        with traffic.app.app_context():
            before = traffic.db().execute("SELECT COUNT(*) FROM notifications WHERE user_id=?", (self.client_id,)).fetchone()[0]
            traffic._create_notification(self.client_id, data["request_ids"][0], "admin_reply", "Reply", "child level")
            traffic.db().commit()
            after = traffic.db().execute("SELECT COUNT(*) FROM notifications WHERE user_id=?", (self.client_id,)).fetchone()[0]
        self.assertEqual(before, after)

    def test_deleting_all_children_removes_the_group(self):
        client = self.login(self.client_id, "client", "Group client")
        data = self.submit(client).get_json()
        with traffic.app.app_context():
            traffic._delete_requests(data["request_ids"])
            left = traffic.db().execute("SELECT COUNT(*) FROM request_groups WHERE id=?", (data["group_id"],)).fetchone()[0]
        self.assertEqual(left, 0)

    def test_admin_manages_periods_and_one_must_stay_active(self):
        admin = self.login(self.admin_id, "admin", "Group admin")
        client = self.login(self.client_id, "client", "Group client")
        self.assertEqual(client.post("/api/periods", json={"label": "EVE"}).status_code, 403)
        self.assertTrue(admin.post("/api/periods", json={"label": " eve "}).get_json()["ok"])
        self.assertEqual(admin.post("/api/periods", json={"label": "EVE"}).status_code, 409)
        self.assertIn(b"EVE", admin.get("/admin").data)
        self.assertIn(b"EVE", client.get("/study").data)
        with traffic.app.app_context():
            ids = [r["id"] for r in traffic.db().execute("SELECT id FROM periods ORDER BY id")]
        for pid in ids[:-1]:
            self.assertEqual(admin.post(f"/api/periods/{pid}/availability", json={"enabled": False}).status_code, 200)
        self.assertEqual(admin.post(f"/api/periods/{ids[-1]}/availability", json={"enabled": False}).status_code, 409)
        # a disabled period cannot be ordered
        self.assertEqual(self.submit(client, mode="single", links=[link()], periods=[PERIODS[0]]).status_code, 400)


if __name__ == "__main__":
    unittest.main()
