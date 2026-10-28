import datetime as dt
import json
import os
import sqlite3
import tempfile
import time
import unittest
from io import BytesIO
from pathlib import Path
from unittest import mock
import urllib.error
from zoneinfo import ZoneInfo
from openpyxl import Workbook, load_workbook
from PIL import Image


_database_dir = tempfile.TemporaryDirectory()
os.environ["TRAFFICLENS_DB"] = os.path.join(_database_dir.name, "test.db")
os.environ["OUTSCRAPER_API_KEY"] = "unit-test-api-key"
os.environ["SECRET_KEY"] = "unit-test-session-secret-that-is-long-enough"
os.environ["TRAFFICLENS_OUTSCRAPER_AUTO_POLL"] = "0"

import app as traffic
import outscraper_integration as integration


class RequestDeletionTests(unittest.TestCase):
    def setUp(self):
        with traffic.app.app_context():
            traffic.db().execute(
                "INSERT OR IGNORE INTO users(name,email,pw,role) VALUES('Deletion admin','deletion-admin@example.test','x','admin')"
            )
            self.admin_id = traffic.db().execute(
                "SELECT id FROM users WHERE email='deletion-admin@example.test'"
            ).fetchone()["id"]
            traffic.db().commit()
        self.client = traffic.app.test_client()
        with self.client.session_transaction() as session:
            session.update(uid=self.admin_id, name="Deletion admin", role="admin")

    def insert_request(self, email):
        with traffic.app.app_context():
            traffic.db().execute(
                "INSERT OR IGNORE INTO users(name,email,pw,role) VALUES('Deletion client',?,'x','client')",
                (email,),
            )
            user_id = traffic.db().execute("SELECT id FROM users WHERE email=?", (email,)).fetchone()["id"]
            request_id = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path) "
                "VALUES(?,'2026-01-01','2026-01-01','2026-01-01','[]','test','[]','[]','Pending','link','[]','[]')",
                (user_id,),
            ).lastrowid
            traffic.db().commit()
            return request_id

    def test_individual_delete_removes_request_notifications_and_workflow(self):
        request_id = self.insert_request("delete-one@example.test")
        dashboard = self.client.get("/admin")
        self.assertEqual(dashboard.status_code, 200)
        self.assertIn(b"Delete all requests", dashboard.data)
        self.assertIn(f'class="btn alt delete-request" data-id="{request_id}"'.encode(), dashboard.data)
        job_id = "delete-linked-workflow"
        with traffic.app.app_context():
            traffic.db().execute(
                "INSERT INTO notifications(user_id,request_id,kind,title,message,created_at) VALUES(?,?,'test','Test','Test','2026-01-01')",
                (self.admin_id, request_id),
            )
            traffic.db().execute(
                "INSERT INTO outscraper_jobs(id,created_at,updated_at,config_json,config_hash,approval_token_hash) "
                "VALUES(?, '2026-01-01','2026-01-01','{}','hash','token')",
                (job_id,),
            )
            traffic.db().execute(
                "INSERT INTO request_directions_workflows(job_id,request_id,status,interval_seconds,created_at,updated_at) "
                "VALUES(?,?,'failed',60,'2026-01-01','2026-01-01')",
                (job_id, request_id),
            )
            traffic.db().execute(
                "INSERT INTO volume_freeflow_attempts(job_id,request_id,submitted_at,status) "
                "VALUES(?,?,'2026-01-01','failed')",
                (job_id, request_id),
            )
            traffic.db().commit()

        response = self.client.post(f"/api/requests/{request_id}/delete")

        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["deleted"], 1)
        with traffic.app.app_context():
            self.assertIsNone(traffic.db().execute("SELECT 1 FROM requests WHERE id=?", (request_id,)).fetchone())
            self.assertIsNone(traffic.db().execute("SELECT 1 FROM notifications WHERE request_id=?", (request_id,)).fetchone())
            self.assertIsNone(traffic.db().execute("SELECT 1 FROM request_directions_workflows WHERE request_id=?", (request_id,)).fetchone())
            self.assertIsNone(traffic.db().execute("SELECT 1 FROM volume_freeflow_attempts WHERE request_id=?", (request_id,)).fetchone())
            self.assertIsNotNone(traffic.db().execute("SELECT 1 FROM outscraper_jobs WHERE id=?", (job_id,)).fetchone())

    def test_delete_all_requires_explicit_confirmation_and_removes_every_request(self):
        first_id = self.insert_request("delete-all-one@example.test")
        second_id = self.insert_request("delete-all-two@example.test")
        denied = self.client.post("/api/requests/delete-all", json={})
        self.assertEqual(denied.status_code, 400)

        with traffic.app.app_context():
            client_id = traffic.db().execute(
                "SELECT id FROM users WHERE email='delete-all-one@example.test'"
            ).fetchone()["id"]
        with self.client.session_transaction() as session:
            session.update(uid=client_id, name="Deletion client", role="client")
        forbidden = self.client.post(
            "/api/requests/delete-all",
            json={"confirmation": "DELETE_ALL_REQUESTS"},
        )
        self.assertEqual(forbidden.status_code, 403)
        with self.client.session_transaction() as session:
            session.update(uid=self.admin_id, name="Deletion admin", role="admin")

        response = self.client.post(
            "/api/requests/delete-all",
            json={"confirmation": "DELETE_ALL_REQUESTS"},
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertGreaterEqual(response.get_json()["deleted"], 2)
        with traffic.app.app_context():
            self.assertEqual(traffic.db().execute("SELECT COUNT(*) FROM requests").fetchone()[0], 0)
            self.assertIsNone(traffic.db().execute(
                "SELECT 1 FROM requests WHERE id IN (?,?)", (first_id, second_id),
            ).fetchone())


class DashboardAndProfileTests(unittest.TestCase):
    def setUp(self):
        with traffic.app.app_context():
            self.admin_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES(?,?,?,'admin')",
                ("Profile admin", f"profile-admin-{time.time_ns()}@example.test", traffic.generate_password_hash("admin123")),
            ).lastrowid
            self.client_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES(?,?,?,'client')",
                ("Dashboard client", f"dashboard-{time.time_ns()}@example.test", traffic.generate_password_hash("client123")),
            ).lastrowid
            traffic.db().commit()

    def logged_client(self, user_id, role, name):
        test_client = traffic.app.test_client()
        with test_client.session_transaction() as session:
            session.update(uid=user_id, name=name, role=role)
        return test_client

    def test_role_specific_navigation_and_routes(self):
        admin_client = self.logged_client(self.admin_id, "admin", "Admin")
        self.assertEqual(admin_client.get("/").headers["Location"], "/admin")
        admin_page = admin_client.get("/admin")
        self.assertEqual(admin_page.status_code, 200)
        navigation = admin_page.data.split(b'<nav class="sidebar-nav">', 1)[1].split(b"</nav>", 1)[0]
        for forbidden in (b"Home", b"Study", b"My requests", b"/study", b"/requests"):
            self.assertNotIn(forbidden, navigation)
        for expected in (b"Dashboard", b"Requests", b"Services", b"Users", b"Profile"):
            self.assertIn(expected, navigation)
        self.assertIn(b"Directions API history", admin_page.data)
        self.assertEqual(admin_client.get("/profile").status_code, 302)
        self.assertEqual(admin_client.get("/admin/profile").status_code, 200)

        user_client = self.logged_client(self.client_id, "client", "Dashboard client")
        self.assertEqual(user_client.get("/").headers["Location"], "/dashboard")
        dashboard = user_client.get("/dashboard")
        self.assertEqual(dashboard.status_code, 200)
        navigation = dashboard.data.split(b'<nav class="sidebar-nav">', 1)[1].split(b"</nav>", 1)[0]
        for expected in (b"Home", b"Services", b"My requests", b"Profile"):
            self.assertIn(expected, navigation)
        self.assertIn(b"/study?service=volumes", dashboard.data)
        selected_service_page = user_client.get("/study?service=volumes")
        self.assertEqual(selected_service_page.status_code, 200)
        self.assertIn(b'value="volumes" checked', selected_service_page.data)
        self.assertEqual(user_client.get("/admin").status_code, 403)
        self.assertEqual(user_client.get("/admin/profile").status_code, 403)

    def test_client_can_submit_request_and_admin_can_review_it(self):
        user_client = self.logged_client(self.client_id, "client", "Dashboard client")
        request_payload = {
            "type": "link",
            "services": ["travel"],
            "date": "2026-10-24",
            "start_time": "08:00:00",
            "end_time": "09:00:00",
            "coords": [[30.138239, 31.395698], [30.137423, 31.393946]],
            "path": [[30.138239, 31.395698], [30.137423, 31.393946]],
            "road_km": 0.19,
            "junctions": 0,
            "no_junctions": True,
            "point_c": [30.129997, 31.383281],
            "point_c_distance_km": 1.7,
        }
        with mock.patch.object(traffic, "apply_network", return_value=None):
            response = user_client.post("/api/requests", json=request_payload)
        self.assertEqual(response.status_code, 200, response.get_json())
        request_id = response.get_json()["id"]
        self.assertIn(f"REQ-{request_id}".encode(), user_client.get("/dashboard").data)
        self.assertIn(f"REQ-{request_id}".encode(), user_client.get("/requests").data)
        with traffic.app.app_context():
            stored = traffic.db().execute(
                "SELECT user_id,status FROM requests WHERE id=?", (request_id,)
            ).fetchone()
        self.assertEqual(stored["user_id"], self.client_id)
        self.assertEqual(stored["status"], "Pending")
        admin_client = self.logged_client(self.admin_id, "admin", "Profile admin")
        dashboard = admin_client.get("/admin#requests")
        self.assertEqual(dashboard.status_code, 200)
        self.assertIn(f"REQ-{request_id}".encode(), dashboard.data)
        self.assertEqual(admin_client.get(f"/admin/requests/{request_id}").status_code, 200)
        self.assertEqual(user_client.get(f"/requests/{request_id}").status_code, 200)
        other_client = self.logged_client(self.admin_id, "admin", "Profile admin")
        self.assertEqual(other_client.get(f"/requests/{request_id}").status_code, 404)
        with traffic.app.app_context():
            traffic._delete_requests([request_id])

    def test_user_can_update_profile_and_change_password(self):
        user_client = self.logged_client(self.client_id, "client", "Dashboard client")
        profile_page = user_client.get("/profile")
        self.assertEqual(profile_page.status_code, 200)
        self.assertIn(b"Change password", profile_page.data)
        update = user_client.post("/profile", data={
            "first_name": "Updated",
            "last_name": "Client",
            "email": f"updated-{time.time_ns()}@example.test",
            "phone_number": "+20 01001234567",
            "job_title": "Researcher",
            "engineering_specialization": "Traffic Engineering",
            "company_name": "Traffic Research Lab",
            "bio": "Researching transport systems.",
        })
        self.assertEqual(update.status_code, 302)
        with traffic.app.app_context():
            user = traffic.db().execute(
                "SELECT name,first_name,last_name,email,phone_number,job_title,"
                "engineering_specialization,company_name,bio,pw FROM users WHERE id=?",
                (self.client_id,),
            ).fetchone()
        self.assertEqual(user["name"], "Updated Client")
        self.assertEqual(user["first_name"], "Updated")
        self.assertEqual(user["last_name"], "Client")
        self.assertTrue(user["email"].startswith("updated-"))
        self.assertEqual(user["phone_number"], "+20 01001234567")
        self.assertEqual(user["job_title"], "Researcher")
        self.assertEqual(user["engineering_specialization"], "")
        self.assertEqual(user["company_name"], "Traffic Research Lab")
        self.assertEqual(user["bio"], "Researching transport systems.")
        self.assertEqual(user_client.get("/dashboard").status_code, 200)
        with user_client.session_transaction() as session:
            self.assertEqual(session["name"], "Updated Client")

        wrong_password = user_client.post("/profile/password", data={
            "current_password": "wrong",
            "new_password": "updated123",
            "confirm_password": "updated123",
        })
        self.assertEqual(wrong_password.status_code, 302)
        with traffic.app.app_context():
            user = traffic.db().execute("SELECT pw FROM users WHERE id=?", (self.client_id,)).fetchone()
            self.assertTrue(traffic.check_password_hash(user["pw"], "client123"))

        changed = user_client.post("/profile/password", data={
            "current_password": "client123",
            "new_password": "updated123",
            "confirm_password": "updated123",
        })
        self.assertEqual(changed.status_code, 302)
        with traffic.app.app_context():
            user = traffic.db().execute("SELECT email,pw FROM users WHERE id=?", (self.client_id,)).fetchone()
            self.assertTrue(traffic.check_password_hash(user["pw"], "updated123"))

        logout = user_client.get("/logout")
        self.assertEqual(logout.status_code, 302)
        self.assertEqual(user_client.get("/dashboard").status_code, 302)
        login = user_client.post("/login", data={
            "email": user["email"],
            "password": "updated123",
        })
        self.assertEqual(login.status_code, 302)
        self.assertEqual(login.headers["Location"], "/dashboard")

    def test_admin_profile_is_self_scoped_and_logout_ends_session(self):
        admin_client = self.logged_client(self.admin_id, "admin", "Admin")
        update = admin_client.post("/profile", data={
            "name": "Updated Administrator",
            "email": f"admin-updated-{time.time_ns()}@example.test",
            "phone_number": "+1 202 555 0134",
            "job_title": "Manager",
        })
        self.assertEqual(update.status_code, 302)
        with traffic.app.app_context():
            user = traffic.db().execute(
                "SELECT name,email FROM users WHERE id=?", (self.admin_id,)
            ).fetchone()
        self.assertEqual(user["name"], "Updated Administrator")
        self.assertTrue(user["email"].startswith("admin-updated-"))
        self.assertEqual(admin_client.get("/admin/profile").status_code, 200)
        self.assertEqual(admin_client.get("/dashboard").status_code, 302)
        logout = admin_client.get("/logout")
        self.assertEqual(logout.status_code, 302)
        self.assertEqual(admin_client.get("/admin").status_code, 302)

    def test_profile_validates_fields_and_preserves_submitted_values(self):
        user_client = self.logged_client(self.client_id, "client", "Dashboard client")
        duplicate_email = f"duplicate-{time.time_ns()}@example.test"
        with traffic.app.app_context():
            traffic.db().execute(
                "UPDATE users SET email=? WHERE id=?",
                (duplicate_email, self.admin_id),
            )
            traffic.db().commit()

        invalid = user_client.post("/profile", data={
            "first_name": "  Given  ",
            "last_name": "Family",
            "email": duplicate_email,
            "phone_number": "+1 (202) 555-0134",
            "job_title": "Engineer",
            "engineering_specialization": "Transportation Engineering",
            "company_name": "Road Lab",
            "bio": "Keep this text after an error.",
        })
        self.assertEqual(invalid.status_code, 400)
        self.assertIn(b"This email address is already in use.", invalid.data)
        self.assertIn(b"Keep this text after an error.", invalid.data)
        with traffic.app.app_context():
            user = traffic.db().execute(
                "SELECT email,phone_number,company_name,bio FROM users WHERE id=?",
                (self.client_id,),
            ).fetchone()
        self.assertFalse(user["email"] == duplicate_email)
        self.assertIsNone(user["phone_number"])
        self.assertIsNone(user["company_name"])
        self.assertIsNone(user["bio"])

        invalid_phone = user_client.post("/profile", data={
            "first_name": "Given",
            "last_name": "Family",
            "email": f"valid-{time.time_ns()}@example.test",
            "phone_number": "555-12",
            "job_title": "Researcher",
        })
        self.assertEqual(invalid_phone.status_code, 400)
        self.assertIn(b"Enter a valid international phone number", invalid_phone.data)
        self.assertIn(b'value="555-12"', invalid_phone.data)

    def test_profile_engineering_specialization_is_only_saved_for_engineers(self):
        user_client = self.logged_client(self.client_id, "client", "Dashboard client")
        saved = user_client.post("/profile", data={
            "first_name": "Engineer",
            "last_name": "User",
            "email": f"engineer-{time.time_ns()}@example.test",
            "phone_number": "00123456789",
            "job_title": "Engineer",
            "engineering_specialization": "Transportation Engineering",
            "company_name": "",
            "bio": "",
        })
        self.assertEqual(saved.status_code, 302)
        with traffic.app.app_context():
            engineer = traffic.db().execute(
                "SELECT phone_number,job_title,engineering_specialization FROM users WHERE id=?",
                (self.client_id,),
            ).fetchone()
        self.assertEqual(engineer["phone_number"], "00123456789")
        self.assertEqual(engineer["engineering_specialization"], "Transportation Engineering")

        changed = user_client.post("/profile", data={
            "first_name": "Engineer",
            "last_name": "User",
            "email": f"engineer-{time.time_ns()}@example.test",
            "phone_number": "00123456789",
            "job_title": "Student",
            "engineering_specialization": "Transportation Engineering",
            "company_name": "",
            "bio": "",
        })
        self.assertEqual(changed.status_code, 302)
        with traffic.app.app_context():
            student = traffic.db().execute(
                "SELECT job_title,engineering_specialization FROM users WHERE id=?",
                (self.client_id,),
            ).fetchone()
        self.assertEqual(student["job_title"], "Student")
        self.assertEqual(student["engineering_specialization"], "")

    def test_profile_database_migration_preserves_and_splits_legacy_names(self):
        connection = sqlite3.connect(":memory:")
        try:
            connection.execute(
                "CREATE TABLE users(id INTEGER PRIMARY KEY,name TEXT,email TEXT,pw TEXT,role TEXT)"
            )
            connection.execute(
                "INSERT INTO users(id,name,email,pw,role) VALUES(1,?,?,?,?)",
                ("Ada Lovelace", "ada@example.test", "hashed-password", "client"),
            )
            traffic._migrate_profile_columns(connection)
            traffic._migrate_profile_columns(connection)
            migrated = connection.execute(
                "SELECT name,first_name,last_name,email,pw,role,phone_number FROM users WHERE id=1"
            ).fetchone()
        finally:
            connection.close()

        self.assertEqual(
            migrated,
            ("Ada Lovelace", "Ada", "Lovelace", "ada@example.test", "hashed-password", "client", None),
        )

    def test_profile_photo_upload_is_normalized_and_only_available_to_owner(self):
        client = self.logged_client(self.client_id, "client", "Dashboard client")
        image_stream = BytesIO()
        Image.new("RGBA", (120, 80), (13, 148, 136, 180)).save(image_stream, format="PNG")
        image_stream.seek(0)

        uploaded = client.post(
            "/profile/photo",
            data={"photo": (image_stream, "portrait.png")},
            content_type="multipart/form-data",
        )
        image_stream.close()
        self.assertEqual(uploaded.status_code, 302)
        uploaded.close()
        with traffic.app.app_context():
            user = traffic.db().execute(
                "SELECT profile_image FROM users WHERE id=?", (self.client_id,)
            ).fetchone()
        self.assertRegex(
            user["profile_image"],
            rf"user-{self.client_id}-[a-f0-9]{{32}}\.jpg",
        )

        photo_response = client.get("/profile/photo")
        self.assertEqual(photo_response.status_code, 200)
        self.assertEqual(photo_response.mimetype, "image/jpeg")
        self.assertEqual(photo_response.headers["X-Content-Type-Options"], "nosniff")
        with Image.open(BytesIO(photo_response.data)) as stored_image:
            self.assertEqual(stored_image.format, "JPEG")
            self.assertEqual(stored_image.size, (120, 80))
        photo_response.close()

        profile_page = client.get("/profile")
        self.assertIn(b"Your profile photo", profile_page.data)
        self.assertIn(b"Remove photo", profile_page.data)
        profile_page.close()
        other_user = self.logged_client(self.admin_id, "admin", "Profile admin")
        other_photo = other_user.get("/profile/photo")
        self.assertEqual(other_photo.status_code, 404)
        other_photo.close()

        removed = client.post("/profile/photo", data={"remove_photo": "1"})
        self.assertEqual(removed.status_code, 302)
        removed.close()
        missing_photo = client.get("/profile/photo")
        self.assertEqual(missing_photo.status_code, 404)
        missing_photo.close()
        with traffic.app.app_context():
            user = traffic.db().execute(
                "SELECT profile_image FROM users WHERE id=?", (self.client_id,)
            ).fetchone()
        self.assertIsNone(user["profile_image"])

    def test_profile_photo_rejects_invalid_images_and_oversized_uploads(self):
        client = self.logged_client(self.client_id, "client", "Dashboard client")
        invalid_stream = BytesIO(b"<svg onload='alert(1)'></svg>")
        invalid = client.post(
            "/profile/photo",
            data={"photo": (invalid_stream, "avatar.svg")},
            content_type="multipart/form-data",
        )
        invalid_stream.close()
        self.assertEqual(invalid.status_code, 302)
        invalid.close()
        page = client.get("/profile")
        self.assertIn(b"selected image could not be opened", page.data)
        page.close()
        missing_photo = client.get("/profile/photo")
        self.assertEqual(missing_photo.status_code, 404)
        missing_photo.close()

        oversized_stream = BytesIO(b"x" * (traffic.PROFILE_PHOTO_MAX_BYTES + 1))
        oversized = client.post(
            "/profile/photo",
            data={
                "photo": (
                    oversized_stream,
                    "large.jpg",
                )
            },
            content_type="multipart/form-data",
        )
        oversized_stream.close()
        self.assertEqual(oversized.status_code, 302)
        oversized.close()
        page = client.get("/profile")
        self.assertIn(
            b"Profile photos must be 5 MB or smaller.",
            page.data,
        )
        page.close()


class OutscraperValidationTests(unittest.TestCase):
    def setUp(self):
        self.config = {
            "routes": [{"origin": "Kuwait City", "destination": "Salmiya"}],
            "departure_time": None,
            "finish_time": None,
            "interval": None,
            "travel_mode": "car",
            "language": "en",
            "region": "KW",
        }

    def test_validates_required_routes_and_documented_enums(self):
        self.assertEqual(integration.validate_configuration(self.config)["routes"], self.config["routes"])
        for invalid in (
            {**self.config, "routes": []},
            {**self.config, "routes": [{"origin": "A", "destination": ""}]},
            {**self.config, "travel_mode": "helicopter"},
            {**self.config, "language": "not-a-language"},
            {**self.config, "region": "XX"},
            {**self.config, "finish_time": 10, "interval": 2},
            {**self.config, "interval": 2},
        ):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                integration.validate_configuration(invalid)

    def test_csv_is_validated_and_limited_to_origin_destination(self):
        self.assertEqual(
            integration.parse_routes_csv("\ufefforigin,destination\nA,B\nC,D\n"),
            [{"origin": "A", "destination": "B"}, {"origin": "C", "destination": "D"}],
        )
        for text in ("", "origin,destination\nA\n", "origin,origin,destination\nA,B,C\n"):
            with self.subTest(text=text), self.assertRaises(ValueError):
                integration.parse_routes_csv(text)

    def test_server_request_uses_key_header_and_async_mode(self):
        response = mock.MagicMock()
        response.__enter__.return_value = response
        response.read.return_value = b'{"id":"req-1","status":"Pending","data":[]}'
        config = integration.validate_configuration(self.config)
        with mock.patch.object(integration.urllib.request, "urlopen", return_value=response) as open_url:
            result = integration.submit_route(config["routes"][0], config, api_key="private-test-key")
        request = open_url.call_args.args[0]
        self.assertEqual(request.get_header("X-api-key"), "private-test-key")
        self.assertIn("async=true", request.full_url)
        self.assertNotIn("private-test-key", request.full_url)
        self.assertEqual(result["id"], "req-1")
        self.assertEqual(open_url.call_args.kwargs["timeout"], integration.REQUEST_TIMEOUT_SECONDS)

    def test_request_details_match_encoded_url_and_never_contain_api_key(self):
        config = integration.validate_configuration({
            "routes": [{"origin": "Kuwait City & Salmiya", "destination": "Airport"}],
            "departure_time": 1792310400, "finish_time": None, "interval": None,
            "travel_mode": "car", "language": "en", "region": "KW",
        })
        details = integration.directions_request_details(config["routes"][0], config)
        self.assertEqual(details["method"], "GET")
        self.assertIn("origin=Kuwait+City+%26+Salmiya", details["url"])
        self.assertIn("travel_mode=car", details["url"])
        self.assertIn("departure_time=1792310400", details["url"])
        self.assertEqual(details["headers"]["X-API-KEY"], "Configured; value hidden")
        self.assertNotIn(os.environ["OUTSCRAPER_API_KEY"], str(details))

    def test_route_departure_override_is_sent_instead_of_common_departure(self):
        configuration = integration.validate_configuration({
            "routes": [{
                "origin": "Kuwait City", "destination": "Salmiya",
                "departure_time": "2026-10-18T07:00",
            }],
            "departure_time": 1792310400, "finish_time": None, "interval": None,
            "travel_mode": "car", "language": "en", "region": "KW",
        })
        expected = int(dt.datetime(2026, 10, 18, 7, tzinfo=dt.timezone(dt.timedelta(hours=3))).timestamp())
        self.assertEqual(configuration["routes"][0]["departure_time"], expected)
        query = integration._query_parameters(configuration["routes"][0], configuration)
        sent_departure = next(value for name, value in query if name == "departure_time")
        self.assertEqual(sent_departure, str(expected))

    def test_result_lookup_uses_documented_results_path(self):
        response = mock.MagicMock()
        response.__enter__.return_value = response
        response.read.return_value = b'{"id":"request-1","status":"Success","data":[]}'
        with mock.patch.object(integration.urllib.request, "urlopen", return_value=response) as open_url:
            result = integration.fetch_request("request-1", api_key="private-test-key")
        self.assertEqual(result["status"], "Success")
        self.assertTrue(open_url.call_args.args[0].full_url.startswith(integration.RESULTS_BASE_URL + "/requests/"))
        with self.assertRaises(ValueError):
            integration.fetch_request("../secret", api_key="private-test-key")

    def test_rate_limit_and_timeout_are_safe_and_not_retried(self):
        with mock.patch.object(integration.urllib.request, "urlopen", side_effect=urllib.error.HTTPError(
            integration.DIRECTIONS_ENDPOINT, 429, "rate limit", {}, None
        )):
            with self.assertRaisesRegex(integration.OutscraperError, "HTTP 429") as caught:
                integration.submit_route(self.config["routes"][0], self.config, api_key="do-not-leak")
        self.assertNotIn("do-not-leak", str(caught.exception))
        with mock.patch.object(integration.urllib.request, "urlopen", side_effect=TimeoutError):
            with self.assertRaises(integration.SubmissionUncertain):
                integration.submit_route(self.config["routes"][0], self.config, api_key="do-not-leak")

    def test_secret_redaction(self):
        cleaned = integration.redact_provider_data(
            {"message": "key=private-test-key", "X-API-KEY": "private-test-key"},
            api_key="private-test-key",
        )
        self.assertNotIn("private-test-key", str(cleaned))
        self.assertEqual(cleaned["X-API-KEY"], "[REDACTED]")


class OutscraperWorkflowTests(unittest.TestCase):
    def setUp(self):
        with traffic.app.app_context():
            traffic.db().execute("DELETE FROM request_directions_workflows")
            traffic.db().execute("DELETE FROM outscraper_job_routes")
            traffic.db().execute("DELETE FROM outscraper_jobs")
            self.client_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Volume client',?,'x','client')",
                (f"volume-client-{time.time_ns()}@example.test",),
            ).lastrowid
            traffic.db().commit()
        self.client = traffic.app.test_client()
        with self.client.session_transaction() as session:
            session.update(uid=1, name="Admin", role="admin")
        self.config = {
            "routes": [{"origin": "Kuwait City", "destination": "Salmiya"}],
            "departure_time": None,
            "finish_time": None,
            "interval": None,
            "travel_mode": "best",
            "language": "en",
            "region": "KW",
        }

    def test_weekend_reference_is_latest_country_weekend_strictly_before_today(self):
        cairo = ZoneInfo("Africa/Cairo")
        friday_morning = dt.datetime(2026, 10, 30, 7, 0, tzinfo=cairo)
        reference_date, weekday = traffic._latest_previous_weekly_weekend(
            "EG", "Africa/Cairo", now=friday_morning
        )
        self.assertEqual(reference_date, dt.date(2026, 10, 24))
        self.assertEqual(weekday, "Saturday")
        self.assertLess(reference_date, friday_morning.date())

        saturday_morning = dt.datetime(2026, 10, 24, 7, 0, tzinfo=cairo)
        reference_date, weekday = traffic._latest_previous_weekly_weekend(
            "EG", "Africa/Cairo", now=saturday_morning
        )
        self.assertEqual(reference_date, dt.date(2026, 10, 23))
        self.assertEqual(weekday, "Friday")
        self.assertLess(reference_date, saturday_morning.date())

    def test_client_can_create_request_with_volumes_and_other_services(self):
        with self.client.session_transaction() as session:
            session.update(uid=self.client_id, name="Volume client", role="client")
        payload = {
            "services": ["volumes", "speed", "travel"],
            "date": "2026-10-01",
            "start_time": "07:00:00",
            "end_time": "09:00:00",
            "coords": [[30.138239, 31.395698], [30.137423, 31.393946]],
            "path": [[30.138239, 31.395698], [30.137423, 31.393946]],
            "road_km": 0.19,
            "junctions": 0,
            "no_junctions": True,
            "point_c": [30.129997, 31.383281],
            "point_c_distance_km": 1.7,
            "area": "Combined service test",
        }
        with mock.patch.object(traffic, "apply_network", return_value=None):
            response = self.client.post("/api/requests", json=payload)
        self.assertEqual(response.status_code, 200, response.get_json())
        with traffic.app.app_context():
            row = traffic.db().execute(
                "SELECT services FROM requests WHERE id=?", (response.get_json()["id"],)
            ).fetchone()
            workflow = traffic.db().execute(
                "SELECT status FROM volumes_workflows WHERE request_id=?",
                (response.get_json()["id"],),
            ).fetchone()
        self.assertEqual(json.loads(row["services"]), ["volumes", "speed", "travel"])
        self.assertEqual(workflow["status"], "waiting_for_inputs")

    def test_volumes_freeflow_form_endpoint_returns_route_and_reference_context(self):
        request_id = self.insert_volume_request()
        with mock.patch.object(traffic, "_route_holiday_location", return_value=("EG", "Egypt", "Africa/Cairo")), \
             mock.patch.object(traffic.outscraper, "api_key_configured", return_value=True):
            response = self.client.get(f"/api/requests/{request_id}/volumes/freeflow-preview")
        self.assertEqual(response.status_code, 200, response.get_json())
        configuration = response.get_json()["configuration"]
        context = configuration["trafficlens_context"]
        reference_date = dt.date.fromisoformat(context["weekend_date"])
        current_date = dt.datetime.now(ZoneInfo("Africa/Cairo")).date()
        self.assertLess(reference_date, current_date)
        self.assertEqual(context["reference_type"], "weekly_weekend")
        self.assertEqual(context["weekend_day"], reference_date.strftime("%A"))
        self.assertEqual(configuration["routes"][0]["departure_time"], context["departure_time"])

    def preview(self, config=None):
        response = self.client.post("/api/outscraper/preview", json=config or self.config)
        self.assertEqual(response.status_code, 201, response.get_json())
        return response.get_json()

    def wait_for(self, job_id, states):
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline:
            with traffic.app.app_context():
                job = traffic._outscraper_job(job_id)
            if job and job["execution_status"] in states:
                return job
            time.sleep(0.01)
        self.fail(f"Job {job_id} did not reach {states}")

    def approve(self, preview, configuration=None):
        return self.client.post(
            f"/api/outscraper/{preview['id']}/approve",
            json={
                "approval_token": preview["approval_token"],
                "configuration": configuration or preview["configuration"],
                "acknowledge_unknown_cost": True,
            },
        )

    def test_request_directions_sends_a_to_c_extracts_a_to_b_and_waits_for_admin_confirmation(self):
        route_timezone = ZoneInfo("Africa/Cairo")
        start = int(dt.datetime(2025, 12, 18, 7, 15, 30, tzinfo=route_timezone).timestamp())
        finish = int(dt.datetime(2025, 12, 18, 8, 45, 30, tzinfo=route_timezone).timestamp())
        record_one = {
            "timestamp": int(dt.datetime(2025, 12, 18, 7, 30, tzinfo=route_timezone).timestamp()),
            "road_distance_timing(meters:seconds)": {
                "0": 0, "141": 8, "232": 13, "444": 24, "524": 28,
                "673": 35, "1056": 54, "1447": 74, "1506": 77,
            },
        }
        with traffic.app.app_context():
            user_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Directions client','directions-client@example.test','x','client')"
            ).lastrowid
            cursor = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path,road_km,start_time,end_time,point_c,point_c_distance_km,data) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (user_id, "2025-12-18", "2025-12-18", "2025-12-18", '["travel","speed"]',
                 "A to C directions test", "[30.1,31.3,30.2,31.4]", "[[30.138239,31.395698],[30.137423,31.393946]]",
                 "Processing", "link", "[7,8]", "[[30.138239,31.395698],[30.137423,31.393946]]",
                 0.19, "07:15:30", "08:45:30", "[30.129997,31.383281]", 1.7,
                 '{"travel":{"rows":[],"note":""}}'),
            )
            request_id = cursor.lastrowid
            traffic.db().commit()
        self.assertEqual(self.client.get(f"/admin/requests/{request_id}").status_code, 200)

        config = {
            "routes": [{"origin": "30.138239, 31.395698", "destination": "30.129997, 31.383281"}],
            "departure_time": start, "finish_time": finish, "interval": 90,
            "travel_mode": "car", "language": "en", "region": "EG",
            "trafficlens_request": {
                "source_request_id": request_id,
                "point_a": "30.138239, 31.395698",
                "point_b": "30.137423, 31.393946",
                "point_c": "30.129997, 31.383281",
                "road_ab_meters": 190,
                "interval_minutes": 90,
                "country_code": "EG",
                "country_name": "Egypt",
                "timezone": "Africa/Cairo",
            },
        }
        with mock.patch.object(traffic, "_route_holiday_location", return_value=("EG", "Egypt", "Africa/Cairo")):
            location = self.client.get(f"/api/requests/{request_id}/directions/location")
            self.assertEqual(location.status_code, 200)
            self.assertEqual(location.get_json()["location"], {
                "country_code": "EG", "country_name": "Egypt", "timezone": "Africa/Cairo",
            })
            preview = self.preview(config)
            duplicate_preview = self.client.post("/api/outscraper/preview", json=config)
            self.assertEqual(duplicate_preview.status_code, 409)
            self.assertIn("already has a Directions workflow", duplicate_preview.get_json()["error"])
        with traffic.app.app_context():
            initial = traffic.db().execute(
                "SELECT status FROM request_directions_workflows WHERE job_id=?", (preview["id"],)
            ).fetchone()
            original_data = traffic.db().execute(
                "SELECT data FROM requests WHERE id=?", (request_id,)
            ).fetchone()["data"]
        self.assertEqual(initial["status"], "awaiting_approval")
        self.assertEqual(original_data, '{"travel":{"rows":[],"note":""}}')

        provider_response = {"id": "request-a-to-c", "status": "Success", "data": [[record_one]]}
        with mock.patch.object(traffic, "_route_holiday_location", return_value=("EG", "Egypt", "Africa/Cairo")):
          with mock.patch.object(traffic.outscraper, "submit_route", return_value=provider_response) as submit:
            approved = self.approve(preview)
            self.assertEqual(approved.status_code, 202, approved.get_json())
            job = self.wait_for(preview["id"], {"completed"})
        sent_route = submit.call_args.args[0]
        self.assertEqual(sent_route["origin"], "30.138239, 31.395698")
        self.assertEqual(sent_route["destination"], "30.129997, 31.383281")
        self.assertEqual(job["config"]["interval"], 90)
        self.assertEqual(job["config"]["trafficlens_request"]["interval_minutes"], 90)
        request_parameters = dict(integration._query_parameters(sent_route, job["config"]))
        self.assertEqual(request_parameters["interval"], "90")
        self.assertEqual(request_parameters["region"], "EG")
        self.assertEqual(request_parameters["departure_time"], str(start))
        self.assertEqual(request_parameters["finish_time"], str(finish))
        with traffic.app.app_context():
            workflow = traffic.db().execute(
                "SELECT status,result_json FROM request_directions_workflows WHERE job_id=?", (preview["id"],)
            ).fetchone()
            source_data = traffic.db().execute(
                "SELECT data FROM requests WHERE id=?", (request_id,)
            ).fetchone()["data"]
        self.assertEqual(workflow["status"], "ready")
        self.assertEqual(source_data, original_data)
        extracted = __import__("json").loads(workflow["result_json"])
        self.assertEqual(extracted["distance_meters"], 190)
        self.assertAlmostEqual(extracted["time_seconds"], 10.6923076923)
        self.assertAlmostEqual(extracted["speed_kmh"], 190 / 10.6923076923 * 3.6)
        self.assertEqual(extracted["reused_hours"], [8])
        self.assertFalse(extracted["hourly"][0]["reused_observation"])
        self.assertTrue(extracted["hourly"][1]["reused_observation"])
        self.assertAlmostEqual(extracted["hourly"][0]["time_seconds"], 10.6923076923)
        self.assertAlmostEqual(extracted["hourly"][1]["time_seconds"], 10.6923076923)
        self.assertAlmostEqual(extracted["hourly"][0]["speed_kmh"], 190 / 10.6923076923 * 3.6)
        review_page = self.client.get(f"/admin/requests/{request_id}")
        self.assertIn(b"Point A", review_page.data)
        self.assertIn(b"Point C", review_page.data)
        self.assertIn(b"Start time", review_page.data)
        self.assertIn(b"End time", review_page.data)
        self.assertIn(b"Default:", review_page.data)
        self.assertIn(b"Time interval (minutes)", review_page.data)
        self.assertIn(b"View Data Scraper", review_page.data)
        self.assertIn(b"Confirm &amp; Send Results", review_page.data)
        self.assertIn(f"/admin/outscraper/{preview['id']}".encode(), review_page.data)
        workflow_status = self.client.get(f"/api/requests/{request_id}/directions").get_json()["workflow"]
        self.assertEqual(workflow_status["status"], "ready")
        self.assertEqual(workflow_status["view_url"], f"/admin/outscraper/{preview['id']}")

        confirmed = self.client.post(f"/api/requests/{request_id}/directions/confirm", json={"job_id": preview["id"]})
        self.assertEqual(confirmed.status_code, 200, confirmed.get_json())
        self.assertEqual(
            self.client.get(f"/api/requests/{request_id}/directions").get_json()["workflow"]["status"],
            "confirmed",
        )
        duplicate = self.client.post(f"/api/requests/{request_id}/directions/confirm", json={"job_id": preview["id"]})
        self.assertEqual(duplicate.status_code, 409)
        with traffic.app.app_context():
            saved = traffic.db().execute(
                "SELECT data,status FROM requests WHERE id=?", (request_id,)
            ).fetchone()
            notices = traffic.db().execute(
                "SELECT kind FROM notifications WHERE user_id=? AND request_id=?",
                (user_id, request_id),
            ).fetchall()
        saved_data = __import__("json").loads(saved["data"])
        self.assertEqual(saved["status"], "Completed")
        self.assertAlmostEqual(saved_data["travel"]["rows"][0]["value"], 10.6923076923 / 60)
        self.assertAlmostEqual(saved_data["travel"]["rows"][1]["value"], 10.6923076923 / 60)
        self.assertAlmostEqual(saved_data["speed"]["rows"][1]["value"], 190 / 10.6923076923 * 3.6)
        self.assertIn("reused for those hours", saved_data["travel"]["note"])
        self.assertEqual(saved_data["travel"]["directions_result"]["route"], "A → B")
        self.assertIn("directions_result", [notice["kind"] for notice in notices])
        client = traffic.app.test_client()
        with client.session_transaction() as session:
            session.update(uid=user_id, name="Directions client", role="client")
        client_result = client.get(f"/results/{request_id}")
        self.assertIn(b"Route A", client_result.data)
        self.assertIn(b"190.0 meters", client_result.data)
        self.assertNotIn(b"request-a-to-c", client_result.data)
        self.assertNotIn(b"30.129997", client_result.data)
        workbook_response = client.get(f"/results/{request_id}/export.xlsx")
        workbook = load_workbook(BytesIO(workbook_response.data), data_only=True)
        summary = list(workbook["Request Summary"].values)
        self.assertIn(("A-to-B distance (m)", 190), summary)


    def insert_volume_request(self, status="Processing", services=None):
        services = services or ["volumes"]
        with traffic.app.app_context():
            cursor = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path,road_km,start_time,end_time,point_c,point_c_distance_km,data) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (self.client_id, "2025-12-18", "2025-12-18", "2025-12-18", json.dumps(services),
                 "Volume workflow test", "[30.1,31.3,30.2,31.4]", "[[30.138239,31.395698],[30.137423,31.393946]]",
                 status, "link", "[7,8]", "[[30.138239,31.395698],[30.137423,31.393946]]",
                 0.19, "07:00:00", "09:00:00", "[30.129997,31.383281]", 1.7, "{}"),
            )
            request_id = cursor.lastrowid
            now = traffic._timestamp_now()
            traffic.db().execute(
                "INSERT INTO volumes_workflows(request_id,status,created_at,updated_at) VALUES(?,'waiting_for_inputs',?,?)",
                (request_id, now, now),
            )
            traffic.db().commit()
            return request_id

    def test_directions_confirmation_reuses_saved_matlab_result_for_admin_review(self):
        request_id = self.insert_volume_request()
        job_id = f"directions-ready-{request_id}"
        direction_result = {
            "route": "A to B",
            "distance_meters": 440.0,
            "time_seconds": 56.16,
            "speed_kmh": 28.20,
            "point_a": "30.1, 31.3",
            "point_b": "30.2, 31.4",
            "point_c": "30.3, 31.5",
            "hourly": [
                {
                    "hour": 7, "speed_kmh": 27.085714285714285,
                    "distance_m": 440.0, "time_seconds": 58.48,
                },
                {
                    "hour": 8, "speed_kmh": 29.0,
                    "distance_m": 440.0, "time_seconds": 53.84,
                },
            ],
            "source_request_id": request_id,
        }
        with traffic.app.app_context():
            traffic.db().execute(
                "UPDATE volumes_workflows SET lanes=1, freeflow_speed=30,freeflow_accepted=1 "
                "WHERE request_id=?",
                (request_id,),
            )
            traffic.db().execute(
                "INSERT INTO request_directions_workflows(job_id,request_id,status,interval_seconds,"
                "result_json,created_at,updated_at) VALUES(?,?,'ready',3600,?,?,?)",
                (
                    job_id,
                    request_id,
                    json.dumps(direction_result),
                    traffic._timestamp_now(),
                    traffic._timestamp_now(),
                ),
            )
            traffic.db().commit()

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            app_root = temp_root / "app"
            app_root.mkdir()
            shared_workbook = temp_root / "Inquiries_12-10-2026.xlsx"
            models_file = temp_root / "Models.mat"
            models_file.write_bytes(b"test-only placeholder")
            source = Workbook()
            source.active.title = "Inqueries"
            source.active.append(["R", "Lanes", "FFS", "Actual Speed"])
            source.save(shared_workbook)
            source.close()

            with mock.patch.object(traffic.volumes_runner, "APP_DIR", app_root), \
                    mock.patch.object(traffic.volumes_runner, "INPUT_TEMPLATE", shared_workbook), \
                    mock.patch.object(traffic.volumes_runner, "MODELS_FILE", models_file):
                prepared = traffic.volumes_runner.prepare_volume_run(request_id, [
                    {
                        "hour": 7,
                        "lanes": 1,
                        "freeflow_speed": 30,
                        "actual_speed": 27.085714285714285,
                    },
                    {
                        "hour": 8,
                        "lanes": 1,
                        "freeflow_speed": 30,
                        "actual_speed": 29.0,
                    },
                ])
                result_book = Workbook()
                result_sheet = result_book.active
                result_sheet.title = "Volumes"
                result_sheet.append([
                    "R", "Lanes", "FFS", "Actual Speed", "Hour",
                    "Estimated Volume", "RMSE Tolerance",
                ])
                result_sheet.append([
                    request_id, 1, 30, 27.08571428571429, 7, 898.61, 118.18,
                ])
                result_sheet.append([
                    request_id, 1, 30, 29.0, 8, 923.03, 128.37,
                ])
                result_book.save(prepared["result_workbook"])
                result_book.close()

                confirmation = self.client.post(
                    f"/api/requests/{request_id}/directions/confirm",
                    json={"job_id": job_id},
                )
                self.assertEqual(confirmation.status_code, 200, confirmation.get_json())
                self.assertEqual(confirmation.get_json()["status"], "confirmed")
                with mock.patch.object(traffic.volumes_runner.subprocess, "run") as matlab:
                    detail = self.client.get(f"/admin/requests/{request_id}")
                matlab.assert_not_called()

        self.assertEqual(detail.status_code, 200)
        self.assertIn(b"ready_for_review", detail.data)
        self.assertIn(b"898.61", detail.data)
        self.assertIn(b"923.03", detail.data)
        self.assertIn(b"Approve &amp; Send Volumes to Client", detail.data)
        with traffic.app.app_context():
            request = traffic.db().execute(
                "SELECT status FROM requests WHERE id=?", (request_id,)
            ).fetchone()
            workflow = traffic.db().execute(
                "SELECT status,actual_result_json,result_json,workbook_path "
                "FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()
            directions = traffic.db().execute(
                "SELECT status,confirmed_by,confirmed_at FROM request_directions_workflows "
                "WHERE request_id=? AND job_id=?",
                (request_id, job_id),
            ).fetchone()
        self.assertEqual(request["status"], "Processing")
        self.assertEqual(workflow["status"], "ready_for_review")
        self.assertIsNotNone(workflow["actual_result_json"])
        saved_result = json.loads(workflow["result_json"])
        self.assertTrue(saved_result["reused_existing_result"])
        self.assertTrue(workflow["workbook_path"].endswith("-results.xlsx"))
        self.assertEqual(directions["status"], "confirmed")
        self.assertIsNotNone(directions["confirmed_by"])
        self.assertIsNotNone(directions["confirmed_at"])

    def test_volume_calculation_waits_for_every_requested_hour(self):
        request_id = self.insert_volume_request()
        with traffic.app.app_context():
            traffic.db().execute(
                "UPDATE volumes_workflows SET lanes=2, freeflow_speed=40, freeflow_accepted=1, actual_result_json=? "
                "WHERE request_id=?",
                ('{"hourly":[{"hour":7,"speed_kmh":25}]}', request_id),
            )
            traffic.db().commit()
            with mock.patch.object(traffic.volumes_runner, "prepare_volume_run") as prepare:
                started = traffic._start_volume_calculation(request_id)
            workflow = traffic.db().execute(
                "SELECT status,error FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()

        self.assertFalse(started)
        self.assertEqual(workflow["status"], "waiting_for_inputs")
        self.assertIn("08:00", workflow["error"])
        prepare.assert_not_called()

    def test_volume_input_is_saved_before_status_changes_to_running(self):
        request_id = self.insert_volume_request()
        prepared_run = {
            "request_id": request_id,
            "hours": [7, 8],
            "input_workbook": "request-input.xlsx",
        }
        with traffic.app.app_context():
            traffic.db().execute(
                "UPDATE volumes_workflows SET lanes=2, freeflow_speed=40, freeflow_accepted=1, actual_result_json=? "
                "WHERE request_id=?",
                ('{"hourly":[{"hour":7,"speed_kmh":25},{"hour":8,"speed_kmh":30}]}', request_id),
            )
            traffic.db().commit()

            def prepare_inputs(saved_request_id, rows):
                self.assertEqual(saved_request_id, request_id)
                self.assertEqual(
                    [row["hour"] for row in rows],
                    [7, 8],
                )
                status = traffic.db().execute(
                    "SELECT status FROM volumes_workflows WHERE request_id=?",
                    (request_id,),
                ).fetchone()["status"]
                self.assertEqual(status, "waiting_for_inputs")
                return prepared_run

            try:
                with mock.patch.object(
                    traffic.volumes_runner,
                    "prepare_volume_run",
                    side_effect=prepare_inputs,
                ), mock.patch.object(traffic.threading, "Thread") as thread:
                    thread.return_value.start.return_value = None
                    started = traffic._start_volume_calculation(request_id)
                    duplicate_start = traffic._start_volume_calculation(request_id)
                workflow = traffic.db().execute(
                    "SELECT status,workbook_path FROM volumes_workflows WHERE request_id=?",
                    (request_id,),
                ).fetchone()
            finally:
                with traffic._outscraper_workers_lock:
                    traffic._outscraper_workers.discard(f"volumes-{request_id}")

        self.assertTrue(started)
        self.assertFalse(duplicate_start)
        self.assertEqual(workflow["status"], "running")
        self.assertEqual(workflow["workbook_path"], prepared_run["input_workbook"])
        thread.assert_called_once()
        self.assertEqual(thread.call_args.kwargs["args"], (request_id, prepared_run))

    def test_startup_backfills_workflows_for_existing_volume_requests(self):
        request_id = self.insert_volume_request()
        with traffic.app.app_context():
            traffic.db().execute("DELETE FROM volumes_workflows WHERE request_id=?", (request_id,))
            traffic.db().commit()
        traffic.init()
        with traffic.app.app_context():
            workflow = traffic.db().execute(
                "SELECT status FROM volumes_workflows WHERE request_id=?", (request_id,)
            ).fetchone()
        self.assertEqual(workflow["status"], "waiting_for_inputs")

    def test_admin_page_exposes_internal_directions_and_volume_review(self):
        request_id = self.insert_volume_request()
        with mock.patch.object(traffic, "_route_holiday_location", return_value=("EG", "Egypt", "Africa/Cairo")):
            response = self.client.get(f"/admin/requests/{request_id}")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Save Actual Speed for MATLAB", response.data)
        self.assertIn(b"Free flow speed for Volumes", response.data)
        self.assertIn(b'id="freeflowSpeedDialog"', response.data)
        self.assertIn(b'id="volumeFreeflowSpeed"', response.data)
        self.assertIn(b"state.freeflow_speed", response.data)
        self.assertIn(b"/api/outscraper/${encodeURIComponent(state.freeflow_job_id)}/refresh", response.data)
        self.assertNotIn(b'/admin/outscraper?source_request_id=', response.data)
        self.assertIn(b"Number of lanes", response.data)
        self.assertNotIn(b"Confirm &amp; Send Results", response.data)

    def test_volume_only_request_can_prepare_a_separately_approved_freeflow_speed_job(self):
        request_id = self.insert_volume_request()
        with traffic.app.app_context():
            source_request = traffic.db().execute(
                "SELECT * FROM requests WHERE id=?", (request_id,)
            ).fetchone()
            with mock.patch.object(traffic, "_route_holiday_location", return_value=("EG", "Egypt", "Africa/Cairo")):
                with mock.patch.object(traffic, "_nearest_previous_weekly_weekend", return_value=(dt.date(2026, 10, 24), "Saturday")):
                    prepared = traffic._outscraper_prefill_from_request(source_request, "freeflow_speed")
                    response = self.client.post("/api/outscraper/preview", json=prepared["prefill_config"])
        self.assertEqual(response.status_code, 201, response.get_json())
        job_id = response.get_json()["id"]
        with traffic.app.app_context():
            workflow = traffic.db().execute(
                "SELECT freeflow_job_id,status FROM volumes_workflows WHERE request_id=?", (request_id,)
            ).fetchone()
        self.assertEqual(workflow["freeflow_job_id"], job_id)
        self.assertEqual(workflow["status"], "waiting_for_inputs")

    def test_completed_weekend_response_saves_freeflow_speed_and_reference_internally(self):
        request_id = self.insert_volume_request()
        job_id = f"volume-freeflow-{request_id}"
        weekend_date, weekend_day = traffic._latest_previous_weekly_weekend(
            "EG", "Africa/Cairo"
        )
        departure = int(dt.datetime.combine(
            weekend_date, dt.time(7), tzinfo=ZoneInfo("Africa/Cairo")
        ).timestamp())
        context = {
            "source_request_id": request_id,
            "service": "freeflow_speed",
            "study_date": "2025-12-18",
            "country_code": "EG",
            "country_name": "Egypt",
            "timezone": "Africa/Cairo",
            "reference_type": "weekly_weekend",
            "weekend_date": weekend_date.isoformat(),
            "weekend_day": weekend_day,
            "departure_time": departure,
        }
        provider_response = {
            "status": "Success",
            "data": [[{
                "road_distance_timing(meters:seconds)": {"0": 0, "100": 10, "200": 30},
            }]],
        }
        with traffic.app.app_context():
            now = traffic._timestamp_now()
            traffic.db().execute(
                "INSERT INTO outscraper_jobs(id,created_at,updated_at,config_json,config_hash,approval_token_hash) "
                "VALUES(?,?,?,?,?,?)",
                (job_id, now, now, __import__("json").dumps({"trafficlens_context": context}), "hash", "token"),
            )
            traffic.db().execute(
                "INSERT INTO outscraper_job_routes(job_id,route_index,origin,destination,raw_response_json) "
                "VALUES(?,1,'A','B',?)",
                (job_id, __import__("json").dumps(provider_response)),
            )
            with mock.patch.object(traffic, "_start_volume_calculation") as start_model:
                traffic._persist_holiday_result_context(job_id)
            start_model.assert_called_once_with(request_id)
            workflow = traffic.db().execute(
                "SELECT freeflow_speed,freeflow_job_id,status FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()
            request_data = traffic.db().execute(
                "SELECT data FROM requests WHERE id=?", (request_id,)
            ).fetchone()["data"]
        self.assertAlmostEqual(workflow["freeflow_speed"], 190 / 28 * 3.6)
        self.assertEqual(workflow["freeflow_job_id"], job_id)
        self.assertEqual(workflow["status"], "waiting_for_inputs")
        saved_context = __import__("json").loads(request_data)["volumes"]["freeflow_context"]
        self.assertEqual(saved_context["weekend_date"], weekend_date.isoformat())
        self.assertEqual(saved_context["weekend_day"], weekend_day)

    def test_volumes_status_recovers_a_completed_linked_freeflow_job(self):
        request_id = self.insert_volume_request()
        job_id = f"completed-freeflow-{request_id}"
        weekend_date, weekend_day = traffic._latest_previous_weekly_weekend("EG", "Africa/Cairo")
        context = {
            "source_request_id": request_id,
            "service": "freeflow_speed",
            "study_date": "2025-12-18",
            "country_code": "EG",
            "country_name": "Egypt",
            "timezone": "Africa/Cairo",
            "reference_type": "weekly_weekend",
            "weekend_date": weekend_date.isoformat(),
            "weekend_day": weekend_day,
        }
        response_data = {
            "status": "Success",
            "data": [[{
                "road_distance_timing(meters:seconds)": {"0": 0, "100": 10, "200": 30},
            }]],
        }
        with traffic.app.app_context():
            now = traffic._timestamp_now()
            traffic.db().execute(
                "INSERT INTO outscraper_jobs(id,created_at,updated_at,config_json,config_hash,approval_token_hash,execution_status) "
                "VALUES(?,?,?,?,?,?,?)",
                (
                    job_id, now, now,
                    __import__("json").dumps({"trafficlens_context": context}),
                    "hash", "token", "completed",
                ),
            )
            traffic.db().execute(
                "INSERT INTO outscraper_job_routes(job_id,route_index,origin,destination,execution_status,raw_response_json) "
                "VALUES(?,1,'A','B','completed',?)",
                (job_id, __import__("json").dumps(response_data)),
            )
            traffic.db().execute(
                "UPDATE volumes_workflows SET freeflow_job_id=? WHERE request_id=?",
                (job_id, request_id),
            )
            traffic.db().commit()
            with mock.patch.object(traffic, "_start_volume_calculation") as start_model:
                response = self.client.get(f"/api/requests/{request_id}/volumes")
            start_model.assert_called_once_with(request_id)
        self.assertEqual(response.status_code, 200, response.get_json())
        workflow = response.get_json()["workflow"]
        self.assertAlmostEqual(workflow["freeflow_speed"], 190 / 28 * 3.6)
        self.assertEqual(workflow["freeflow_job_id"], job_id)
        self.assertEqual(workflow["freeflow_job_status"], "completed")
        self.assertEqual(workflow["freeflow_context"]["weekend_date"], weekend_date.isoformat())

    def test_volume_freeflow_uses_a_to_c_window_and_nearest_eligible_weekend(self):
        request_id = self.insert_volume_request()
        weekend = dt.date(2026, 10, 9)
        with traffic.app.app_context():
            source = traffic.db().execute(
                "SELECT * FROM requests WHERE id=?", (request_id,)
            ).fetchone()
            with mock.patch.object(traffic, "_route_holiday_location", return_value=("EG", "Egypt", "Africa/Cairo")):
                with mock.patch.object(traffic, "_nearest_previous_weekly_weekend", return_value=(weekend, "Friday")):
                    prepared = traffic._outscraper_prefill_from_request(source, "freeflow_speed")
        configuration = prepared["prefill_config"]
        context = configuration["trafficlens_context"]
        self.assertEqual(configuration["routes"][0]["origin"], "30.138239, 31.395698")
        self.assertEqual(configuration["routes"][0]["destination"], "30.129997, 31.383281")
        self.assertTrue(context["volume_input"])
        self.assertEqual(context["weekend_date"], "2026-10-09")
        self.assertEqual(configuration["interval"], 60)
        start_local = dt.datetime.fromtimestamp(configuration["departure_time"], ZoneInfo("Africa/Cairo"))
        finish_local = dt.datetime.fromtimestamp(configuration["finish_time"], ZoneInfo("Africa/Cairo"))
        self.assertEqual((start_local.date(), start_local.hour), (weekend, 7))
        self.assertEqual((finish_local.date(), finish_local.hour), (weekend, 8))

    def test_freeflow_mismatch_requires_explicit_use_anyway_before_matlab(self):
        request_id = self.insert_volume_request(
            services=["volumes", "speed", "freeflow_speed"],
        )
        weekend_date = dt.date(2026, 10, 24)
        departure = int(dt.datetime.combine(
            weekend_date, dt.time(7), tzinfo=ZoneInfo("Africa/Cairo")
        ).timestamp())
        finish = int(dt.datetime.combine(
            weekend_date, dt.time(8), tzinfo=ZoneInfo("Africa/Cairo")
        ).timestamp())
        context = {
            "source_request_id": request_id,
            "service": "freeflow_speed",
            "volume_input": True,
            "study_date": "2025-12-18",
            "country_code": "EG",
            "country_name": "Egypt",
            "timezone": "Africa/Cairo",
            "reference_type": "weekly_weekend",
            "weekend_date": weekend_date.isoformat(),
            "weekend_day": "Saturday",
            "departure_time": departure,
            "finish_time": finish,
            "interval_minutes": 60,
            "road_ab_meters": 190,
        }
        job_id = f"freeflow-mismatch-{request_id}"
        response_data = {
            "status": "Success",
            "data": [[
                {
                    "timestamp": departure + 900,
                    "road_distance_timing(meters:seconds)": {"0": 0, "190": 10},
                },
                {
                    "timestamp": departure + 2700,
                    "road_distance_timing(meters:seconds)": {"0": 0, "190": 7.2},
                },
            ]],
        }
        with traffic.app.app_context():
            now = traffic._timestamp_now()
            traffic.db().execute(
                "UPDATE volumes_workflows SET lanes=2,actual_result_json=? WHERE request_id=?",
                (json.dumps({"hourly": [
                    {"hour": 7, "speed_kmh": 50},
                    {"hour": 8, "speed_kmh": 96},
                ]}), request_id),
            )
            traffic.db().execute(
                "INSERT INTO outscraper_jobs(id,created_at,updated_at,config_json,config_hash,approval_token_hash,execution_status) "
                "VALUES(?,?,?,?,?,?, 'completed')",
                (job_id, now, now, json.dumps({"trafficlens_context": context}), "hash", "token"),
            )
            traffic.db().execute(
                "INSERT INTO outscraper_job_routes(job_id,route_index,origin,destination,execution_status,raw_response_json) "
                "VALUES(?,1,'A','C','completed',?)",
                (job_id, json.dumps(response_data)),
            )
            traffic.db().execute(
                "UPDATE volumes_workflows SET freeflow_job_id=? WHERE request_id=?",
                (job_id, request_id),
            )
            traffic.db().commit()
            with mock.patch.object(traffic, "_start_volume_calculation") as start_model:
                traffic._persist_holiday_result_context(job_id)
                start_model.assert_not_called()
            workflow = traffic.db().execute(
                "SELECT freeflow_speed,freeflow_accepted,freeflow_validation_json FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()
        self.assertAlmostEqual(workflow["freeflow_speed"], 95)
        self.assertEqual(workflow["freeflow_accepted"], 0)
        validation = json.loads(workflow["freeflow_validation_json"])
        self.assertEqual(validation["mismatched_hours"], ["08:00"])

        prepared_run = {"request_id": request_id, "hours": [7, 8], "input_workbook": "in.xlsx"}
        with mock.patch.object(traffic.volumes_runner, "find_existing_volume_result", return_value=None), \
                mock.patch.object(traffic.volumes_runner, "prepare_volume_run", return_value=prepared_run), \
                mock.patch.object(traffic.threading, "Thread") as worker:
            worker.return_value.start.return_value = None
            response = self.client.post(
                f"/api/requests/{request_id}/volumes/freeflow/accept",
                json={"use_anyway": True},
            )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertTrue(response.get_json()["overridden"])
        with traffic.app.app_context():
            workflow = traffic.db().execute(
                "SELECT status,freeflow_accepted,freeflow_override FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()
            request_data = json.loads(traffic.db().execute(
                "SELECT data FROM requests WHERE id=?", (request_id,)
            ).fetchone()["data"])
        self.assertEqual(workflow["status"], "running")
        self.assertEqual(workflow["freeflow_accepted"], 1)
        self.assertEqual(workflow["freeflow_override"], 1)
        self.assertEqual(request_data["freeflow_speed"]["rows"][0]["value"], 95)
        worker.assert_called_once()

    def test_use_anyway_is_blocked_when_actual_speed_hours_are_missing(self):
        request_id = self.insert_volume_request(services=["volumes", "freeflow_speed"])
        with traffic.app.app_context():
            traffic.db().execute(
                "UPDATE volumes_workflows SET freeflow_speed=95,freeflow_result_json=?,"
                "freeflow_validation_json=?,status='waiting_for_inputs' WHERE request_id=?",
                (
                    json.dumps({"selected_speed_kmh": 95}),
                    json.dumps({
                        "is_valid": False,
                        "mismatched_hours": ["08:00"],
                        "missing_hours": ["09:00"],
                    }),
                    request_id,
                ),
            )
            traffic.db().commit()
        with mock.patch.object(traffic, "_start_volume_calculation") as start_model:
            response = self.client.post(
                f"/api/requests/{request_id}/volumes/freeflow/accept",
                json={"use_anyway": True},
            )
        self.assertEqual(response.status_code, 409)
        self.assertIn("missing", response.get_json()["error"])
        start_model.assert_not_called()
        with traffic.app.app_context():
            workflow = traffic.db().execute(
                "SELECT status,freeflow_accepted FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()
        self.assertEqual(workflow["status"], "waiting_for_inputs")
        self.assertEqual(workflow["freeflow_accepted"], 0)

    def test_accept_freeflow_recovers_hourly_speeds_from_existing_directions_review(self):
        request_id = self.insert_volume_request(services=["volumes", "freeflow_speed"])
        job_id = f"saved-directions-{request_id}"
        actual_result = {
            "hourly": [
                {"hour": 7, "speed_kmh": 68.4},
                {"hour": 8, "speed_kmh": 57.0},
            ],
        }
        with traffic.app.app_context():
            now = traffic._timestamp_now()
            traffic.db().execute(
                "UPDATE volumes_workflows SET lanes=2,freeflow_speed=80,"
                "freeflow_job_id='freeflow-existing',freeflow_result_json=?,"
                "freeflow_validation_json=?,actual_result_json=NULL WHERE request_id=?",
                (
                    json.dumps({"selected_speed_kmh": 80}),
                    json.dumps({
                        "is_valid": False,
                        "mismatched_hours": [],
                        "missing_hours": ["07:00", "08:00"],
                    }),
                    request_id,
                ),
            )
            traffic.db().execute(
                "INSERT INTO request_directions_workflows(job_id,request_id,status,interval_seconds,created_at,updated_at,result_json) "
                "VALUES(?,?,'ready',3600,?,?,?)",
                (job_id, request_id, now, now, json.dumps(actual_result)),
            )
            traffic.db().commit()

        with mock.patch.object(traffic, "_start_volume_calculation", return_value=True) as start_model:
            response = self.client.post(
                f"/api/requests/{request_id}/volumes/freeflow/accept",
                json={},
            )

        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertTrue(response.get_json()["accepted"])
        self.assertFalse(response.get_json()["overridden"])
        start_model.assert_called_once_with(request_id)
        with traffic.app.app_context():
            workflow = traffic.db().execute(
                "SELECT actual_result_json,freeflow_validation_json,freeflow_accepted "
                "FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()
        self.assertEqual(
            [item["speed_kmh"] for item in json.loads(workflow["actual_result_json"])["hourly"]],
            [68.4, 57.0],
        )
        self.assertTrue(json.loads(workflow["freeflow_validation_json"])["is_valid"])
        self.assertEqual(workflow["freeflow_accepted"], 1)

    def test_freeflow_comparison_uses_and_persists_client_report_hourly_speeds(self):
        request_id = self.insert_volume_request(services=["volumes", "speed"])
        report_speeds = [
            {"hour": 7, "value": 50.0},
            {"hour": 8, "value": 45.0},
        ]
        stale_internal_result = {
            "hourly": [
                {"hour": 7, "speed_kmh": 12.0},
                {"hour": 8, "speed_kmh": 13.0},
            ],
        }
        with traffic.app.app_context():
            traffic.db().execute(
                "UPDATE requests SET data=? WHERE id=?",
                (json.dumps({"speed": {"rows": report_speeds}}), request_id),
            )
            traffic.db().execute(
                "UPDATE volumes_workflows SET actual_result_json=?,freeflow_result_json=? "
                "WHERE request_id=?",
                (
                    json.dumps(stale_internal_result),
                    json.dumps({"selected_speed_kmh": 60.0, "context": {}}),
                    request_id,
                ),
            )
            traffic.db().commit()

            self.assertTrue(traffic._evaluate_volume_freeflow(request_id))
            workflow = traffic.db().execute(
                "SELECT actual_result_json,freeflow_validation_json FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()

        stored_actual = json.loads(workflow["actual_result_json"])
        validation = json.loads(workflow["freeflow_validation_json"])
        self.assertEqual(
            [item["speed_kmh"] for item in stored_actual["hourly"]],
            [item["value"] for item in report_speeds],
        )
        self.assertEqual(
            [item["actual_speed_kmh"] for item in validation["rows"]],
            [item["value"] for item in report_speeds],
        )
        self.assertTrue(validation["is_valid"])

    def test_actual_speed_confirmation_is_internal_for_volumes_only_request(self):
        request_id = self.insert_volume_request()
        job_id = f"volume-directions-{request_id}"
        result = {
            "hourly": [
                {"hour": 7, "speed_kmh": 50.0},
                {"hour": 8, "speed_kmh": 45.0},
            ],
            "distance_meters": 190,
            "time_seconds": 13.68,
            "speed_kmh": 50.0,
        }
        with traffic.app.app_context():
            now = traffic._timestamp_now()
            traffic.db().execute(
                "INSERT INTO outscraper_jobs(id,created_at,updated_at,config_json,config_hash,approval_token_hash) "
                "VALUES(?,? ,? ,'{}','hash','token')",
                (job_id, now, now),
            )
            traffic.db().execute(
                "INSERT INTO request_directions_workflows(job_id,request_id,status,interval_seconds,created_at,updated_at,result_json) "
                "VALUES(?,?,'ready',3600,?,?,?)",
                (job_id, request_id, now, now, __import__("json").dumps(result)),
            )
            traffic.db().commit()

        with mock.patch.object(traffic, "_start_volume_calculation") as start_model:
            response = self.client.post(
                f"/api/requests/{request_id}/directions/confirm",
                json={"job_id": job_id},
            )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertIn("no unrequested service data was sent to the client", response.get_json()["message"])
        start_model.assert_not_called()
        with traffic.app.app_context():
            request_row = traffic.db().execute(
                "SELECT status,data FROM requests WHERE id=?", (request_id,)
            ).fetchone()
            volume_row = traffic.db().execute(
                "SELECT actual_result_json,status FROM volumes_workflows WHERE request_id=?", (request_id,)
            ).fetchone()
        self.assertEqual(request_row["status"], "Processing")
        self.assertEqual(__import__("json").loads(request_row["data"]), {})
        self.assertEqual(volume_row["status"], "waiting_for_inputs")
        self.assertEqual(__import__("json").loads(volume_row["actual_result_json"]), result)

    def test_directions_confirmation_shares_one_scrape_with_selected_services_and_volumes(self):
        request_id = self.insert_volume_request(services=["volumes", "speed", "travel"])
        job_id = f"volume-and-speed-{request_id}"
        result = {
            "hourly": [
                {"hour": 7, "speed_kmh": 50.0, "travel_minutes": 0.25, "time_seconds": 15.0, "distance_m": 190.0},
                {"hour": 8, "speed_kmh": 45.0, "travel_minutes": 0.3, "time_seconds": 18.0, "distance_m": 190.0},
            ],
            "distance_meters": 190.0,
            "time_seconds": 16.5,
            "speed_kmh": 41.45,
        }
        with traffic.app.app_context():
            now = traffic._timestamp_now()
            traffic.db().execute(
                "INSERT INTO outscraper_jobs(id,created_at,updated_at,config_json,config_hash,approval_token_hash) "
                "VALUES(?,? ,? ,'{}','hash','token')",
                (job_id, now, now),
            )
            traffic.db().execute(
                "INSERT INTO request_directions_workflows(job_id,request_id,status,interval_seconds,created_at,updated_at,result_json) "
                "VALUES(?,?,'ready',3600,?,?,?)",
                (job_id, request_id, now, now, json.dumps(result)),
            )
            traffic.db().commit()
        response = self.client.post(
            f"/api/requests/{request_id}/directions/confirm",
            json={"job_id": job_id},
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertIn("also saved internally", response.get_json()["message"])
        with traffic.app.app_context():
            request_row = traffic.db().execute(
                "SELECT status,data FROM requests WHERE id=?", (request_id,)
            ).fetchone()
            volume_workflow = traffic.db().execute(
                "SELECT actual_result_json,freeflow_accepted FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()
        request_data = json.loads(request_row["data"])
        self.assertEqual(request_row["status"], "Processing")
        self.assertEqual(request_data["speed"]["rows"][1]["value"], 45.0)
        self.assertEqual(request_data["travel"]["rows"][0]["value"], 0.25)
        self.assertNotIn("volumes", request_data)
        self.assertEqual(json.loads(volume_workflow["actual_result_json"]), result)
        self.assertEqual(volume_workflow["freeflow_accepted"], 0)
        with self.client.session_transaction() as session:
            session.update(uid=self.client_id, name="Volume client", role="client")
        result_page = self.client.get(f"/results/{request_id}")
        self.assertIn(b"Speed", result_page.data)
        self.assertIn(b"Travel time", result_page.data)
        self.assertNotIn(b"actual_result_json", result_page.data)

    def test_directions_can_be_confirmed_after_volumes_were_sent(self):
        request_id = self.insert_volume_request(
            services=["lanes", "volumes", "speed", "travel"],
        )
        job_id = f"directions-after-volumes-{request_id}"
        result = {
            "hourly": [
                {
                    "hour": 7, "speed_kmh": 50.0, "travel_minutes": 0.25,
                    "distance_m": 190.0, "time_seconds": 15.0,
                    "reused_observation": False,
                },
                {
                    "hour": 8, "speed_kmh": 45.0, "travel_minutes": 0.3,
                    "distance_m": 190.0, "time_seconds": 18.0,
                    "reused_observation": False,
                },
            ],
            "distance_meters": 190.0,
            "time_seconds": 16.5,
            "speed_kmh": 41.45,
        }
        existing_volumes = {
            "volumes": {
                "rows": [
                    {"hour": 7, "value": 800, "tolerance": 40},
                    {"hour": 8, "value": 750, "tolerance": 35},
                ],
                "source": "matlab",
            },
        }
        with traffic.app.app_context():
            now = traffic._timestamp_now()
            traffic.db().execute(
                "UPDATE requests SET data=?,status='Processing' WHERE id=?",
                (json.dumps(existing_volumes), request_id),
            )
            traffic.db().execute(
                "UPDATE volumes_workflows SET status='sent',lanes=3,result_json=? WHERE request_id=?",
                (
                    json.dumps({
                        "hourly": existing_volumes["volumes"]["rows"],
                        "inputs": {
                            "lanes": 3,
                            "freeflow_speed": 60,
                            "actual_speed_by_hour": [
                                {"hour": 7, "speed_kmh": 50},
                                {"hour": 8, "speed_kmh": 45},
                            ],
                        },
                    }),
                    request_id,
                ),
            )
            traffic.db().execute(
                "INSERT INTO request_directions_workflows(job_id,request_id,status,interval_seconds,created_at,updated_at,result_json) "
                "VALUES(?,?,'ready',3600,?,?,?)",
                (job_id, request_id, now, now, json.dumps(result)),
            )
            traffic.db().commit()

        pending_admin_page = self.client.get(f"/admin/requests/{request_id}")
        self.assertIn(b"Confirm &amp; Send Speed / Travel to Client", pending_admin_page.data)
        self.assertIn(b"Traffic volumes have already been sent.", pending_admin_page.data)
        response = self.client.post(
            f"/api/requests/{request_id}/directions/confirm",
            json={"job_id": job_id},
        )

        self.assertEqual(response.status_code, 200, response.get_json())
        with traffic.app.app_context():
            request_row = traffic.db().execute(
                "SELECT status,data FROM requests WHERE id=?", (request_id,)
            ).fetchone()
            volume_workflow = traffic.db().execute(
                "SELECT status FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()
            directions_workflow = traffic.db().execute(
                "SELECT status FROM request_directions_workflows WHERE job_id=?",
                (job_id,),
            ).fetchone()
        data = json.loads(request_row["data"])
        self.assertEqual(request_row["status"], "Completed")
        self.assertEqual(data["lanes"]["rows"], [{"value": 3.0}])
        self.assertEqual(data["speed"]["rows"][1]["value"], 45.0)
        self.assertEqual(data["travel"]["rows"][0]["value"], 0.25)
        self.assertEqual(data["volumes"], existing_volumes["volumes"])
        self.assertEqual(volume_workflow["status"], "sent")
        self.assertEqual(directions_workflow["status"], "confirmed")
        admin_page = self.client.get(f"/admin/requests/{request_id}")
        self.assertEqual(admin_page.status_code, 200)
        self.assertIn(b"const REQ=", admin_page.data)
        self.assertIn(b"const rows=hourly?REQ.hours.map", admin_page.data)

    def test_saving_volume_lanes_populates_requested_lanes_service(self):
        request_id = self.insert_volume_request(services=["lanes", "volumes"])
        with mock.patch.object(traffic, "_start_volume_calculation", return_value=False):
            response = self.client.post(
                f"/api/requests/{request_id}/volumes/lanes",
                json={"lanes": "3"},
            )
        self.assertEqual(response.status_code, 200, response.get_json())
        with traffic.app.app_context():
            row = traffic.db().execute(
                "SELECT data FROM requests WHERE id=?", (request_id,)
            ).fetchone()
        self.assertEqual(
            json.loads(row["data"])["lanes"]["rows"],
            [{"value": 3.0}],
        )

    def test_admin_approval_sends_volume_and_rmse_tolerance_to_client(self):
        request_id = self.insert_volume_request(services=["lanes", "volumes"])
        result = {
            "hourly": [
                {"hour": 7, "value": 812.5, "tolerance": 43.2},
                {"hour": 8, "value": 790.0, "tolerance": 38.6},
            ],
            "inputs": {
                "lanes": 2,
                "freeflow_speed": 60.0,
                "freeflow_context": {
                    "study_date": "2026-10-01",
                    "country_code": "EG",
                    "country_name": "Egypt",
                    "timezone": "Africa/Cairo",
                    "reference_type": "weekly_weekend",
                    "weekend_date": "2026-09-26",
                    "weekend_day": "Saturday",
                    "departure_time": 1790402400,
                },
                "actual_speed_by_hour": [
                    {"hour": 7, "speed_kmh": 50.0},
                    {"hour": 8, "speed_kmh": 45.0},
                ],
            },
        }
        with traffic.app.app_context():
            traffic.db().execute(
                "UPDATE volumes_workflows SET status='ready_for_review', lanes=2, freeflow_speed=60, result_json=? WHERE request_id=?",
                (__import__("json").dumps(result), request_id),
            )
            traffic.db().commit()

        manual_save = self.client.post(
            f"/api/requests/{request_id}/save-all",
            json={"services": {"volumes": {"rows": [{"hour": 7, "value": 1}], "note": ""}}},
        )
        self.assertEqual(manual_save.status_code, 409)
        manual_single_save = self.client.post(
            f"/api/requests/{request_id}/data/volumes",
            json={"rows": [{"hour": 7, "value": 1}], "note": ""},
        )
        self.assertEqual(manual_single_save.status_code, 409)
        response = self.client.post(f"/api/requests/{request_id}/volumes/confirm", json={})
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["status"], "sent")
        self.assertEqual(response.get_json()["request_status"], "Completed")
        with traffic.app.app_context():
            request_data = json.loads(traffic.db().execute(
                "SELECT data FROM requests WHERE id=?", (request_id,)
            ).fetchone()["data"])
        self.assertEqual(request_data["lanes"]["rows"], [{"value": 2.0}])
        with self.client.session_transaction() as session:
            session.update(uid=self.client_id, name="Volume client", role="client")
        client_result = self.client.get(f"/results/{request_id}")
        self.assertEqual(client_result.status_code, 200)
        self.assertIn(b"RMSE tolerance", client_result.data)
        self.assertIn(b"43.2", client_result.data)
        self.assertIn(b"813", client_result.data)
        self.assertNotIn(b"812.5", client_result.data)
        self.assertEqual(
            [row["value"] for row in request_data["volumes"]["rows"]], [813, 790]
        )
        self.assertIn(b"2026-09-26", client_result.data)
        self.assertIn(b"Saturday", client_result.data)
        self.assertIn(b"Number of lanes", client_result.data)
        export_response = self.client.get(f"/results/{request_id}/export.xlsx")
        workbook = load_workbook(BytesIO(export_response.data), data_only=True)
        details = list(workbook["Service Details"].values)
        self.assertIn("Tolerance", details[0])
        self.assertTrue(any(row[2] == 813 and row[3] == 43.2 and row[5] == "MATLAB" for row in details[1:]))
        summary = list(workbook["Request Summary"].values)
        self.assertTrue(any("weekly weekend used" in str(row[0]) and "Saturday · 2026-09-26" in str(row[1]) for row in summary))

    def test_request_cannot_be_marked_completed_before_every_selected_service_is_ready(self):
        request_id = self.insert_volume_request(services=["volumes", "speed"])
        direct = self.client.post(
            f"/api/requests/{request_id}/status",
            json={"status": "Completed"},
        )
        self.assertEqual(direct.status_code, 409)
        self.assertIn("every selected service", direct.get_json()["error"])
        bulk = self.client.post(
            f"/api/requests/{request_id}/save-all",
            json={"status": "Completed", "services": {}},
        )
        self.assertEqual(bulk.status_code, 409)

    def test_request_directions_interpolation_requires_profile_to_reach_point_b(self):
        self.assertEqual(traffic._interpolate_route_seconds({"0": 0, "100": 10, "200": 30}, 150), 20)
        with self.assertRaisesRegex(ValueError, "does not reach Point B"):
            traffic._interpolate_route_seconds({"0": 0, "100": 10}, 150)

    def test_completed_directions_with_no_data_still_fails_extraction(self):
        job_id = "completed-empty-directions"
        config = {
            "departure_time": 1766034000,
            "finish_time": 1766041200,
            "trafficlens_request": {"timezone": "Africa/Cairo", "road_ab_meters": 190},
        }
        with traffic.app.app_context():
            request_id = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path,road_km,point_c) "
                "VALUES(1,'2025-12-18','2025-12-18','2025-12-18','[\"travel\"]','test','[]',"
                "'[[30.138239,31.395698],[30.137423,31.393946]]','Processing','link','[7,8]','[]',0.19,"
                "'[30.129997,31.383281]')"
            ).lastrowid
            traffic.db().execute(
                "INSERT INTO outscraper_jobs(id,created_at,updated_at,config_json,config_hash,approval_token_hash) "
                "VALUES(?,?,?,?,?,?)",
                (job_id, "2025-12-18", "2025-12-18", __import__("json").dumps(config), "hash", "token"),
            )
            traffic.db().execute(
                "INSERT INTO outscraper_job_routes(job_id,route_index,origin,destination,execution_status,raw_response_json) "
                "VALUES(?,1,'A','C','completed','{\"status\":\"Success\",\"data\":[]}')",
                (job_id,),
            )
            traffic.db().execute(
                "INSERT INTO request_directions_workflows(job_id,request_id,status,interval_seconds,created_at,updated_at) "
                "VALUES(?,?,'running',7200,'2025-12-18','2025-12-18')",
                (job_id, request_id),
            )
            traffic.db().commit()
            with self.assertRaisesRegex(ValueError, "no usable timing data"):
                traffic._extract_request_directions_result(job_id)

    def test_ready_directions_hourly_speeds_are_saved_for_volumes_before_client_confirmation(self):
        request_id = self.insert_volume_request(services=["volumes", "speed"])
        job_id = f"ready-volume-directions-{request_id}"
        timezone = ZoneInfo("Africa/Cairo")
        start = int(dt.datetime(2025, 12, 18, 16, 0, tzinfo=timezone).timestamp())
        finish = int(dt.datetime(2025, 12, 18, 18, 0, tzinfo=timezone).timestamp())
        config = {
            "departure_time": start,
            "finish_time": finish,
            "trafficlens_request": {
                "point_a": "30.138239, 31.395698",
                "point_b": "30.137423, 31.393946",
                "point_c": "30.129997, 31.383281",
                "timezone": "Africa/Cairo",
                "road_ab_meters": 190,
            },
        }
        response = {
            "status": "Success",
            "data": [[
                {
                    "timestamp": start + 300,
                    "road_distance_timing(meters:seconds)": {"0": 0, "190": 10},
                },
                {
                    "timestamp": start + 3600 + 300,
                    "road_distance_timing(meters:seconds)": {"0": 0, "190": 12},
                },
            ]],
        }
        with traffic.app.app_context():
            now = traffic._timestamp_now()
            traffic.db().execute(
                "UPDATE requests SET hours='[16,17]' WHERE id=?",
                (request_id,),
            )
            traffic.db().execute(
                "UPDATE volumes_workflows SET lanes=2,freeflow_result_json=? "
                "WHERE request_id=?",
                (
                    json.dumps({"selected_speed_kmh": 80, "context": {}}),
                    request_id,
                ),
            )
            traffic.db().execute(
                "INSERT INTO outscraper_jobs(id,created_at,updated_at,config_json,config_hash,approval_token_hash) "
                "VALUES(?,?,?,?,?,?)",
                (job_id, now, now, json.dumps(config), "hash", "token"),
            )
            traffic.db().execute(
                "INSERT INTO outscraper_job_routes(job_id,route_index,origin,destination,execution_status,raw_response_json) "
                "VALUES(?,1,'A','C','completed',?)",
                (job_id, json.dumps(response)),
            )
            traffic.db().execute(
                "INSERT INTO request_directions_workflows(job_id,request_id,status,interval_seconds,created_at,updated_at) "
                "VALUES(?,?,'running',3600,?,?)",
                (job_id, request_id, now, now),
            )
            traffic.db().commit()

            with mock.patch.object(traffic, "_start_volume_calculation") as start_calculation:
                traffic._extract_request_directions_result(job_id)
                start_calculation.assert_called_once_with(request_id)

            saved = traffic.db().execute(
                "SELECT actual_result_json,freeflow_validation_json FROM volumes_workflows WHERE request_id=?",
                (request_id,),
            ).fetchone()
            directions = traffic.db().execute(
                "SELECT status FROM request_directions_workflows WHERE job_id=?",
                (job_id,),
            ).fetchone()
        actual = json.loads(saved["actual_result_json"])
        validation = json.loads(saved["freeflow_validation_json"])
        self.assertEqual(
            [item["speed_kmh"] for item in actual["hourly"]],
            [68.4, 57.0],
        )
        self.assertEqual(
            [item["actual_speed_kmh"] for item in validation["rows"]],
            [68.4, 57.0],
        )
        self.assertTrue(validation["is_valid"])
        self.assertEqual(directions["status"], "ready")

    def test_preview_cost_unknown_and_exact_configuration_is_persisted(self):
        payload = {**self.config, "routes": self.config["routes"] * 2}
        preview = self.preview(payload)
        with traffic.app.app_context():
            job = traffic._outscraper_job(preview["id"])
        self.assertEqual(job["config"], preview["configuration"])
        self.assertEqual(preview["planned_http_submissions"], 2)
        self.assertIsNone(preview["cost_estimate"])
        self.assertEqual(preview["estimate_confidence"], "unknown")
        self.assertIsNone(job["actual_cost"])

    def test_rejection_never_calls_provider(self):
        preview = self.preview()
        with mock.patch.object(traffic.outscraper, "submit_route") as submit:
            response = self.client.post(
                f"/api/outscraper/{preview['id']}/reject",
                json={"approval_token": preview["approval_token"]},
            )
            self.assertEqual(response.status_code, 200)
            submit.assert_not_called()
        with traffic.app.app_context():
            job = traffic._outscraper_job(preview["id"])
        self.assertEqual(job["approval_status"], "rejected")
        self.assertIsNotNone(job["rejected_at"])

    def test_execution_requires_unknown_cost_acknowledgement_and_matching_configuration(self):
        preview = self.preview()
        with mock.patch.object(traffic.outscraper, "submit_route") as submit:
            denied = self.client.post(f"/api/outscraper/{preview['id']}/approve", json={
                "approval_token": preview["approval_token"],
                "configuration": preview["configuration"],
            })
            self.assertEqual(denied.status_code, 409)
            modified = {**preview["configuration"], "routes": [{"origin": "Changed", "destination": "Salmiya"}]}
            changed = self.approve(preview, modified)
            self.assertEqual(changed.status_code, 409)
            submit.assert_not_called()

    def _holiday_config(self, request_id):
        holiday, holiday_name = traffic._nearest_public_holiday("KW", dt.date(2026, 10, 18))
        departure = int(dt.datetime.combine(
            holiday, dt.time(7), tzinfo=ZoneInfo("Asia/Kuwait")
        ).timestamp())
        origin, destination = "29.300000, 47.900000", "29.400000, 48.000000"
        return {
            "routes": [{"origin": origin, "destination": destination, "departure_time": departure}],
            "departure_time": None, "finish_time": None, "interval": None,
            "travel_mode": "car", "language": "en", "region": "KW",
            "trafficlens_context": {
                "source_request_id": request_id, "service": "freeflow_time",
                "study_date": "2026-10-18", "country_code": "KW", "country_name": "Kuwait",
                "timezone": "Asia/Kuwait", "holiday_date": holiday.isoformat(),
                "holiday_name": holiday_name, "departure_time": departure,
                "origin": origin, "destination": destination,
            },
        }

    def _insert_holiday_request(self, service="freeflow_time"):
        with traffic.app.app_context():
            cursor = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path,road_km,junctions,start_time,end_time) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (1, "2026-10-18", "2026-10-18", "2026-10-18", __import__("json").dumps([service]),
                 "Free flow route", "[29.3,47.9,29.4,48.0]", "[[29.3,47.9],[29.4,48.0]]",
                 "Pending", "link", "[7]", "[[29.3,47.9],[29.35,47.95],[29.4,48.0]]",
                 1.2, 0, "07:00:00", "08:00:00"),
            )
            traffic.db().commit()
            return traffic.db().execute("SELECT * FROM requests WHERE id=?", (cursor.lastrowid,)).fetchone()

    def test_freeflow_services_are_available_and_excluded_from_static_network_values(self):
        with mock.patch.object(traffic, "get_setting", return_value={}):
            catalog = traffic.service_catalog()
        self.assertTrue(catalog["freeflow_time"]["enabled"])
        self.assertTrue(catalog["freeflow_speed"]["enabled"])
        row = self._insert_holiday_request("freeflow_speed")
        with mock.patch.object(traffic.network, "HAVE_GEO", True), \
             mock.patch.object(traffic, "get_setting", return_value={"file": "network.gpkg"}), \
             mock.patch.object(traffic.network, "lookup", return_value={
                 "values": {"freeflow_speed": {"value": 60, "note": "static"}}, "coverage": 1,
             }):
            with traffic.app.app_context():
                outcome = traffic.apply_network(row["id"])
        self.assertEqual(outcome["filled"], [])
        with traffic.app.app_context():
            saved = traffic.db().execute("SELECT data FROM requests WHERE id=?", (row["id"],)).fetchone()
        self.assertEqual(__import__("json").loads(saved["data"] or "{}"), {})

    def test_freeflow_prefill_uses_nearest_country_holiday_at_7am_and_shows_date(self):
        row = self._insert_holiday_request("freeflow_time")
        with mock.patch.object(traffic, "_route_holiday_location", return_value=("KW", "Kuwait", "Asia/Kuwait")):
            with traffic.app.app_context():
                prepared = traffic._outscraper_prefill_from_request(row, "freeflow_time")
            page = self.client.get(f"/admin/outscraper?source_request_id={row['id']}&service=freeflow_time")
        context = prepared["prefill_config"]["trafficlens_context"]
        expected_date, expected_name = traffic._nearest_public_holiday("KW", dt.date.fromisoformat(row["d1"]))
        departure = dt.datetime.fromtimestamp(context["departure_time"], ZoneInfo(context["timezone"]))
        self.assertEqual(context["holiday_date"], expected_date.isoformat())
        self.assertEqual(context["holiday_name"], expected_name)
        self.assertEqual((departure.hour, departure.minute), (7, 0))
        self.assertEqual(prepared["prefill_config"]["routes"][0]["departure_time"], context["departure_time"])
        self.assertEqual(page.status_code, 200)
        self.assertIn(expected_date.isoformat().encode(), page.data)
        self.assertIn(b"07:00", page.data)

    def test_holiday_configuration_is_bound_to_route_date_and_time(self):
        row = self._insert_holiday_request()
        config = self._holiday_config(row["id"])
        clean = integration.validate_configuration(config)
        self.assertEqual(clean["trafficlens_context"]["holiday_date"], config["trafficlens_context"]["holiday_date"])
        self.assertEqual(clean["routes"][0]["departure_time"], clean["trafficlens_context"]["departure_time"])
        changed_route = {**config, "routes": [{**config["routes"][0], "origin": "Changed"}]}
        with self.assertRaisesRegex(ValueError, "Reference route and departure time"):
            integration.validate_configuration(changed_route)
        changed_time = {**config, "routes": [{**config["routes"][0], "departure_time": config["routes"][0]["departure_time"] + 3600}]}
        with self.assertRaisesRegex(ValueError, "Reference route and departure time"):
            integration.validate_configuration(changed_time)

    def test_completed_holiday_api_result_persists_date_in_request_and_exports(self):
        row = self._insert_holiday_request()
        config = self._holiday_config(row["id"])
        used_date = config["trafficlens_context"]["holiday_date"]
        preview = self.preview(config)
        with traffic.app.app_context():
            traffic.db().execute(
                "UPDATE outscraper_job_routes SET execution_status='completed', provider_request_id=?, provider_status=?, raw_response_json=?, result_count=1 WHERE job_id=?",
                ("holiday-provider", "Success", '{"id":"holiday-provider","status":"Success","data":[[{"duration":"10 mins"}]]}', preview["id"]),
            )
            traffic._refresh_job_status(preview["id"])
            saved = traffic.db().execute("SELECT data FROM requests WHERE id=?", (row["id"],)).fetchone()
        service_data = __import__("json").loads(saved["data"])["freeflow_time"]
        self.assertEqual(service_data["source"], "outscraper")
        self.assertEqual(service_data["holiday_context"]["holiday_date"], used_date)
        self.assertEqual(service_data["outscraper_job_id"], preview["id"])
        result_page = self.client.get(f"/results/{row['id']}")
        self.assertIn(b"Nearest public holiday used", result_page.data)
        self.assertIn(used_date.encode(), result_page.data)
        export = self.client.get(f"/results/{row['id']}/export.xlsx")
        workbook = load_workbook(BytesIO(export.data), data_only=True)
        summary_values = [value for record in workbook["Request Summary"].values for value in record]
        self.assertTrue(any(used_date in str(value) for value in summary_values))
        csv_export = self.client.get(f"/admin/outscraper/{preview['id']}/results.csv")
        self.assertIn(b"holiday_date_used", csv_export.data)
        self.assertIn(used_date.encode(), csv_export.data)

    def test_configuration_change_invalidates_approval(self):
        preview = self.preview()
        invalidated = self.client.post(
            f"/api/outscraper/{preview['id']}/invalidate",
            json={"approval_token": preview["approval_token"]},
        )
        self.assertEqual(invalidated.status_code, 200)
        with mock.patch.object(traffic.outscraper, "submit_route") as submit:
            response = self.approve(preview)
            self.assertEqual(response.status_code, 409)
            submit.assert_not_called()

    def test_duplicate_approval_submits_once(self):
        preview = self.preview()
        with mock.patch.object(traffic.outscraper, "submit_route", return_value={
            "id": "provider-1", "status": "Pending", "data": [],
        }) as submit:
            first = self.approve(preview)
            self.assertEqual(first.status_code, 202, first.get_json())
            job = self.wait_for(preview["id"], {"running"})
            second = self.approve(preview)
            self.assertEqual(second.status_code, 409)
            submit.assert_called_once()
        self.assertEqual(job["routes"][0]["provider_request_id"], "provider-1")
        self.assertEqual(job["routes"][0]["request_details"]["method"], "GET")
        self.assertIn("travel_mode=best", job["routes"][0]["request_details"]["url"])
        self.assertNotIn("unit-test-api-key", str(job["routes"][0]["request_details"]))
        detail = self.client.get(f"/admin/outscraper/{preview['id']}")
        self.assertIn(b"API request sent for each route", detail.data)
        self.assertIn(b"travel_mode=best", detail.data)
        self.assertNotIn(b"unit-test-api-key", detail.data)

    def test_timeout_marks_submission_unknown_and_does_not_resubmit(self):
        preview = self.preview()
        with mock.patch.object(traffic.outscraper, "submit_route", side_effect=integration.SubmissionUncertain(
            "Connection to Outscraper failed; submission status is uncertain"
        )) as submit:
            self.assertEqual(self.approve(preview).status_code, 202)
            job = self.wait_for(preview["id"], {"unknown"})
            self.assertEqual(self.approve(preview).status_code, 409)
            submit.assert_called_once()
        self.assertIn("uncertain", job["error"])
        with mock.patch.object(traffic.outscraper, "fetch_request") as poll:
            self.client.post(f"/api/outscraper/{preview['id']}/refresh")
            poll.assert_not_called()

    def test_provider_failure_is_saved_without_retry(self):
        preview = self.preview()
        with mock.patch.object(traffic.outscraper, "submit_route", side_effect=integration.OutscraperError(
            "Outscraper returned HTTP 400"
        )) as submit:
            self.assertEqual(self.approve(preview).status_code, 202)
            job = self.wait_for(preview["id"], {"failed"})
            submit.assert_called_once()
        self.assertEqual(job["routes"][0]["error"], "Outscraper returned HTTP 400")

    def test_async_results_are_polled_persisted_and_exported(self):
        preview = self.preview()
        pending = {"id": "provider-async-1", "status": "Pending", "data": []}
        completed = {"id": "provider-async-1", "status": "Success", "data": [[
            {"road": "Route 1", "distance(meters)": 1200},
        ]]}
        with mock.patch.object(traffic.outscraper, "submit_route", return_value=pending):
            self.assertEqual(self.approve(preview).status_code, 202)
            self.wait_for(preview["id"], {"running"})
        with mock.patch.object(traffic.outscraper, "fetch_request", return_value=completed) as poll:
            response = self.client.post(f"/api/outscraper/{preview['id']}/refresh")
            self.assertEqual(response.status_code, 200, response.get_json())
            poll.assert_called_once_with("provider-async-1")
        with traffic.app.app_context():
            job = traffic._outscraper_job(preview["id"])
        self.assertEqual(job["execution_status"], "completed")
        self.assertEqual(job["result_count"], 1)
        self.assertEqual(job["routes"][0]["result"]["data"][0][0]["road"], "Route 1")
        self.assertEqual(self.client.get(f"/admin/outscraper/{preview['id']}").status_code, 200)
        exported = self.client.get(f"/admin/outscraper/{preview['id']}/results.csv")
        self.assertEqual(exported.status_code, 200)
        self.assertIn(b"Route 1", exported.data)
        self.assertIn(b"distance(meters)", exported.data)

    def test_empty_completed_results_are_retained_as_successful_empty_result(self):
        preview = self.preview()
        with mock.patch.object(traffic.outscraper, "submit_route", return_value={
            "id": "provider-empty", "status": "Success", "data": [],
        }):
            self.assertEqual(self.approve(preview).status_code, 202)
            job = self.wait_for(preview["id"], {"completed"})
        self.assertEqual(job["result_count"], 0)
        self.assertEqual(job["routes"][0]["result"]["data"], [])

    def test_background_monitor_persists_results_without_an_open_page(self):
        preview = self.preview()
        with traffic.app.app_context():
            traffic.db().execute("UPDATE outscraper_jobs SET approval_status='approved',execution_status='running' WHERE id=?", (preview["id"],))
            traffic.db().execute(
                "UPDATE outscraper_job_routes SET execution_status='running',provider_request_id='provider-background' WHERE job_id=?",
                (preview["id"],))
            traffic.db().commit()
        completed = {"id": "provider-background", "status": "Success", "data": [[{"road": "Background route"}]]}
        with mock.patch.object(traffic.time, "sleep") as wait, mock.patch.object(
            traffic.outscraper, "fetch_request", return_value=completed
        ) as poll:
            traffic._monitor_outscraper_job(preview["id"])
        wait.assert_called_once_with(15)
        poll.assert_called_once_with("provider-background")
        with traffic.app.app_context():
            job = traffic._outscraper_job(preview["id"])
        self.assertEqual(job["execution_status"], "completed")
        self.assertEqual(job["routes"][0]["result"]["data"][0][0]["road"], "Background route")

    def test_startup_migrates_legacy_route_columns_before_status_refresh(self):
        directory = tempfile.TemporaryDirectory()
        original_db = traffic.DB
        legacy_db = os.path.join(directory.name, "legacy.db")
        try:
            traffic.DB = legacy_db
            legacy_schema = traffic.SCHEMA.replace(
                "last_polled_at TEXT,\n  request_details_json TEXT, ",
                "",
            ).replace(",\n  start_time TEXT, end_time TEXT, point_c TEXT", "")
            legacy_schema = legacy_schema.replace("interval_minutes INTEGER,\n", "")
            self.assertNotEqual(legacy_schema, traffic.SCHEMA)
            connection = sqlite3.connect(legacy_db)
            connection.executescript(legacy_schema)
            connection.execute(
                "INSERT INTO users(id,name,email,pw,role) VALUES(1,'Admin','migration-admin@example.test','x','admin')"
            )
            connection.execute("INSERT INTO settings(key,value) VALUES('bundled_network_setup_v2','true')")
            connection.execute(
                "INSERT INTO outscraper_jobs(id,created_at,updated_at,config_json,config_hash,approval_token_hash,approval_status,execution_status,provider_ids) "
                "VALUES('legacy-job','2026-10-18T00:00:00+00:00','2026-10-18T00:00:00+00:00','{}','hash','token','approved','running','[\"legacy-provider\"]')"
            )
            connection.execute(
                "INSERT INTO outscraper_job_routes(job_id,route_index,origin,destination,execution_status,provider_request_id) "
                "VALUES('legacy-job',1,'Origin','Destination','running','legacy-provider')"
            )
            connection.commit()
            connection.close()

            traffic.init()
            migrated = sqlite3.connect(legacy_db)
            columns = {row[1] for row in migrated.execute("PRAGMA table_info(outscraper_job_routes)")}
            self.assertIn("last_polled_at", columns)
            self.assertIn("request_details_json", columns)
            workflow_columns = {row[1] for row in migrated.execute("PRAGMA table_info(request_directions_workflows)")}
            self.assertIn("interval_minutes", workflow_columns)
            request_columns = {row[1] for row in migrated.execute("PRAGMA table_info(requests)")}
            self.assertIn("start_time", request_columns)
            self.assertIn("end_time", request_columns)
            self.assertIn("point_c", request_columns)
            migrated.close()

            with mock.patch.object(traffic.outscraper, "fetch_request", return_value={
                "id": "legacy-provider", "status": "Success", "data": [],
            }) as poll:
                response = self.client.post("/api/outscraper/legacy-job/refresh")
            self.assertEqual(response.status_code, 200, response.get_json())
            poll.assert_called_once_with("legacy-provider")
            self.assertEqual(response.get_json()["status"], "completed")
        finally:
            traffic.DB = original_db
            directory.cleanup()

    def test_csv_export_neutralizes_spreadsheet_formulas(self):
        preview = self.preview()
        response_data = {"id": "provider-1", "status": "Success", "data": [[
            {"road": '=HYPERLINK("bad")'},
        ]]}
        with traffic.app.app_context():
            traffic.db().execute(
                "UPDATE outscraper_job_routes SET execution_status='completed', provider_request_id=?, provider_status=?, raw_response_json=?, result_count=1 WHERE job_id=?",
                ("provider-1", "Success", __import__("json").dumps(response_data), preview["id"]))
            traffic.db().execute("UPDATE outscraper_jobs SET execution_status='completed', provider_ids='[\"provider-1\"]', result_count=1 WHERE id=?", (preview["id"],))
            traffic.db().commit()
        exported = self.client.get(f"/admin/outscraper/{preview['id']}/results.csv")
        self.assertIn(b"'=HYPERLINK", exported.data)

    def test_admin_page_renders_and_client_is_forbidden(self):
        preview = self.preview()
        self.assertEqual(self.client.get("/admin/outscraper").status_code, 200)
        self.assertEqual(self.client.get("/admin").status_code, 200)
        self.assertEqual(self.client.get("/study").status_code, 200)
        detail = self.client.get(f"/admin/outscraper/{preview['id']}")
        self.assertEqual(detail.status_code, 200)
        self.assertIn(b"Continue approval", detail.data)
        self.assertIn(b"Approve despite unknown cost", detail.data)
        self.assertNotIn(b"unit-test-api-key", detail.data)
        client = traffic.app.test_client()
        with client.session_transaction() as session:
            session.update(uid=999, name="Client", role="client")
        with traffic.app.app_context():
            traffic.db().execute("INSERT INTO users(id,name,email,pw,role) VALUES(999,'Client','outscraper-client@example.test','x','client')")
            traffic.db().commit()
        self.assertEqual(client.get("/admin/outscraper").status_code, 403)
        self.assertTrue(preview["id"])

    def test_my_requests_can_repeat_own_request_with_editable_prefill(self):
        with traffic.app.app_context():
            user_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Repeat client','repeat-client@example.test','x','client')"
            ).lastrowid
            cursor = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path,road_km,junctions,start_time,end_time,point_c) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (user_id, "2026-10-21", "2026-10-21", "2026-10-21", '["travel","speed"]',
                 "Old road study", "[29.3,47.9,29.4,48.0]", "[[29.31,47.91],[29.32,47.92]]",
                 "Completed", "link", "[7,8]", "[[29.31,47.91],[29.32,47.92]]",
                 1.2, 0, "07:15:30", "08:45:10", "[29.33,47.93]"),
            )
            request_id = cursor.lastrowid
            traffic.db().commit()
        with self.client.session_transaction() as session:
            session.update(uid=user_id, name="Repeat client", role="client")

        requests_page = self.client.get("/requests")
        self.assertEqual(requests_page.status_code, 200)
        self.assertIn(b"Repeat / edit", requests_page.data)
        repeat_page = self.client.get(f"/study?repeat_request_id={request_id}")
        self.assertEqual(repeat_page.status_code, 200)
        self.assertIn(b"Repeating request REQ-", repeat_page.data)
        self.assertIn(b'"date": "2026-10-21"', repeat_page.data)
        self.assertIn(b'"start_time": "07:15:30"', repeat_page.data)
        self.assertIn(b'"end_time": "08:45:10"', repeat_page.data)
        self.assertIn(b'"services": ["travel", "speed"]', repeat_page.data)
        self.assertIn(b'"coords": [[29.31, 47.91], [29.32, 47.92]]', repeat_page.data)
        self.assertEqual(self.client.get(f"/study?repeat_request_id={request_id + 9999}").status_code, 404)

    def test_legacy_arabic_notifications_are_converted_to_english(self):
        with traffic.app.app_context():
            db = traffic.db()
            user_id = db.execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Legacy client','legacy-client@example.test','x','client')"
            ).lastrowid
            legacy = (
                ("new_request", "\u0637\u0644\u0628 \u062c\u062f\u064a\u062f",
                 "\u0642\u062f\u0645 Sara \u0637\u0644\u0628\u064b\u0627 \u062c\u062f\u064a\u062f\u064b\u0627: REQ-7."),
                ("volume_result", "\u0646\u062a\u0627\u0626\u062c \u0623\u062d\u062c\u0627\u0645 \u0627\u0644\u0645\u0631\u0648\u0631 \u062c\u0627\u0647\u0632\u0629",
                 "\u0627\u0639\u062a\u0645\u062f \u0627\u0644\u0645\u0633\u0624\u0648\u0644 \u0646\u062a\u0627\u0626\u062c \u0623\u062d\u062c\u0627\u0645 \u0627\u0644\u0645\u0631\u0648\u0631 \u0644\u0644\u0637\u0644\u0628 REQ-1."),
            )
            for kind, title, message in legacy:
                db.execute(
                    "INSERT INTO notifications(user_id,request_id,kind,title,message,created_at) VALUES(?,?,?,?,?,'2026-01-01')",
                    (user_id, 1, kind, title, message),
                )
            db.commit()
            raw = sqlite3.connect(traffic.DB)
            try:
                traffic._translate_legacy_notifications(raw)
                raw.commit()
                rows = raw.execute(
                    "SELECT title,message FROM notifications WHERE user_id=? ORDER BY id", (user_id,)
                ).fetchall()
            finally:
                raw.close()
        self.assertEqual(rows[0], ("New request", "Sara submitted a new request: REQ-7."))
        self.assertEqual(rows[1], (
            "Traffic volume results are ready",
            "The administrator approved the traffic volume results for request REQ-1.",
        ))

    def test_new_request_notifies_admin_and_admin_reply_notifies_only_request_owner(self):
        with traffic.app.app_context():
            client_user_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Notification client','notification-client@example.test','x','client')"
            ).lastrowid
            other_user_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Other client','notification-other@example.test','x','client')"
            ).lastrowid
            traffic.db().commit()
        with self.client.session_transaction() as session:
            session.update(uid=client_user_id, name="Notification client", role="client")

        admin_client = traffic.app.test_client()
        with admin_client.session_transaction() as session:
            session.update(uid=1, name="Admin", role="admin")
        unread_before = admin_client.get("/api/notifications/unread-count").get_json()["count"]

        request_payload = {
            "services": ["travel"], "date": "2026-10-21",
            "start_time": "07:00:00", "end_time": "08:00:00",
            "coords": [[29.3, 47.9], [29.4, 48.0]], "path": [[29.3, 47.9], [29.4, 48.0]],
            "road_km": 1.2, "junctions": 0, "no_junctions": True, "point_c": [29.41, 48.01],
            "point_c_distance_km": 1.6,
        }
        with mock.patch.object(traffic, "apply_network", return_value=None):
            created = self.client.post("/api/requests", json=request_payload)
        self.assertEqual(created.status_code, 200, created.get_json())
        request_id = created.get_json()["id"]

        admin_notifications = admin_client.get("/notifications")
        self.assertEqual(admin_notifications.status_code, 200)
        self.assertIn(b"New request", admin_notifications.data)
        self.assertIn(f"REQ-{request_id}".encode(), admin_notifications.data)
        admin_count = admin_client.get("/api/notifications/unread-count").get_json()
        self.assertEqual(admin_count["count"], unread_before + 1)
        with traffic.app.app_context():
            notification_id = traffic.db().execute(
                "SELECT id FROM notifications WHERE user_id=1 AND request_id=? AND kind='new_request'",
                (request_id,),
            ).fetchone()["id"]
        opened = admin_client.get(f"/notifications/{notification_id}/open")
        self.assertEqual(opened.status_code, 302)
        self.assertTrue(opened.headers["Location"].endswith(f"/admin/requests/{request_id}"))
        self.assertEqual(admin_client.get("/api/notifications/unread-count").get_json()["count"], unread_before)

        with admin_client.session_transaction() as session:
            session.update(uid=1, name="Admin", role="admin")
        reply = admin_client.post(f"/api/requests/{request_id}/reply", json={"reply": "Please review the attached travel data."})
        self.assertEqual(reply.status_code, 200)
        duplicate = admin_client.post(f"/api/requests/{request_id}/reply", json={"reply": "Please review the attached travel data."})
        self.assertEqual(duplicate.status_code, 200)
        saved_reply = admin_client.post(
            f"/api/requests/{request_id}/save-all",
            json={"services": {}, "reply": "Updated reply from Save all."},
        )
        self.assertEqual(saved_reply.status_code, 200)

        with self.client.session_transaction() as session:
            session.update(uid=client_user_id, name="Notification client", role="client")
        client_count = self.client.get("/api/notifications/unread-count").get_json()
        self.assertEqual(client_count["count"], 2)
        client_notifications = self.client.get("/notifications")
        self.assertIn(b"Updated reply from Save all.", client_notifications.data)
        self.assertIn(b"Enable device notifications", client_notifications.data)
        device_feed = self.client.get("/api/notifications/device-feed?after=0")
        self.assertEqual(device_feed.status_code, 200)
        feed_items = device_feed.get_json()["notifications"]
        self.assertEqual(len(feed_items), 2)
        self.assertTrue(all(item["url"] == f"/results/{request_id}" for item in feed_items))
        self.assertIn("Please review the attached travel data.", feed_items[0]["message"])
        worker_asset = self.client.get("/static/notification-sw.js")
        self.assertEqual(worker_asset.status_code, 200)
        worker_asset.close()
        self.assertEqual(self.client.get("/api/notifications/device-feed?after=-1").status_code, 400)
        with traffic.app.app_context():
            client_notification_id = traffic.db().execute(
                "SELECT id FROM notifications WHERE user_id=? AND request_id=? AND kind='admin_reply' ORDER BY id DESC LIMIT 1",
                (client_user_id, request_id),
            ).fetchone()["id"]
        with admin_client.session_transaction() as session:
            session.update(uid=other_user_id, name="Other client", role="client")
        self.assertEqual(admin_client.get(f"/notifications/{client_notification_id}/open").status_code, 404)

    def test_result_excel_export_contains_summary_details_and_safe_text(self):
        with traffic.app.app_context():
            user_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Excel client','excel-client@example.test','x','client')"
            ).lastrowid
            cursor = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path,road_km,junctions,start_time,end_time,point_c,data) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (user_id, "2026-10-21", "2026-10-21", "2026-10-21", '["travel","speed"]',
                 "Excel test link", "[29.3,47.9,29.4,48.0]", "[[29.31,47.91],[29.32,47.92]]",
                 "Completed", "link", "[7,8]", "[[29.31,47.91],[29.32,47.92]]",
                 1.2, 0, "07:15:30", "08:45:10", "[29.33,47.93]",
                 __import__("json").dumps({
                     "travel": {
                         "rows": [{"hour": 7, "value": 25.5}, {"hour": 8, "value": '=HYPERLINK("bad")'}],
                         "note": "Travel time reviewed", "source": "admin", "updated": "2026-10-21 12:00",
                     },
                     "speed": {"rows": [{"hour": 7, "value": 40}], "source": "network"},
                 })),
            )
            request_id = cursor.lastrowid
            traffic.db().commit()
        with self.client.session_transaction() as session:
            session.update(uid=user_id, name="Excel client", role="client")

        result_page = self.client.get(f"/results/{request_id}")
        self.assertEqual(result_page.status_code, 200)
        result_html = result_page.data
        self.assertLess(result_html.index(b'id="rmap"'), result_html.index(b"Technical summary"))
        self.assertLess(result_html.index(b"Technical summary"), result_html.index(b"Service details"))
        self.assertLess(result_html.index(b"Service details"), result_html.index(b"Download Excel report"))

        response = self.client.get(f"/results/{request_id}/export.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.data), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Request Summary", "Service Details"])
        summary_rows = list(workbook["Request Summary"].values)
        self.assertIn(("Time range", "07:15:30–08:45:10"), summary_rows)
        self.assertIn(("Point C", None), summary_rows)
        details = workbook["Service Details"]
        self.assertEqual(details["A2"].value, "Travel time")
        self.assertEqual(details["B2"].value, "07:00-08:00")
        self.assertEqual(details["C2"].value, 25.5)
        self.assertEqual(details["C3"].value, "'=HYPERLINK(\"bad\")")
        self.assertEqual(details["C3"].data_type, "s")

        other_client = traffic.app.test_client()
        with other_client.session_transaction() as session:
            session.update(uid=1, name="Admin", role="admin")
        self.assertEqual(other_client.get(f"/results/{request_id}/export.xlsx").status_code, 200)
        with traffic.app.app_context():
            other_user_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Other client','other-excel@example.test','x','client')"
            ).lastrowid
            traffic.db().commit()
        with other_client.session_transaction() as session:
            session.update(uid=other_user_id, name="Other client", role="client")
        self.assertEqual(other_client.get(f"/results/{request_id}/export.xlsx").status_code, 404)

    def test_admin_controls_point_c_visibility_and_directions_buttons_are_available_from_open_result(self):
        with traffic.app.app_context():
            user_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Point C client','point-c-client@example.test','x','client')"
            ).lastrowid
            cursor = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path,point_c,point_c_distance_km) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (user_id, "2026-10-21", "2026-10-21", "2026-10-21", '["travel"]',
                 "Point C visibility route", "[29.3,47.9,29.4,48.0]", "[[29.31,47.91],[29.32,47.92]]",
                 "Completed", "link", "[7]", "[[29.31,47.91],[29.32,47.92]]", "[29.33,47.93]", 1.6),
            )
            request_id = cursor.lastrowid
            traffic.db().commit()

        client = traffic.app.test_client()
        with client.session_transaction() as session:
            session.update(uid=user_id, name="Point C client", role="client")
        hidden_result = client.get(f"/results/{request_id}")
        self.assertEqual(hidden_result.status_code, 200)
        self.assertNotIn(b"29.33,47.93", hidden_result.data)
        hidden_export = client.get(f"/results/{request_id}/export.xlsx")
        hidden_book = load_workbook(BytesIO(hidden_export.data), data_only=True)
        point_c_row = next(row for row in hidden_book["Request Summary"].values if row[0] == "Point C")
        self.assertIsNone(point_c_row[1])

        response = self.client.post(
            f"/api/requests/{request_id}/save-all",
            json={"services": {}, "point_c_visible": True},
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        shown_result = client.get(f"/results/{request_id}")
        self.assertIn(b"29.33", shown_result.data)
        self.assertIn(b"47.93", shown_result.data)
        shown_export = client.get(f"/results/{request_id}/export.xlsx")
        shown_book = load_workbook(BytesIO(shown_export.data), data_only=True)
        point_c_row = next(row for row in shown_book["Request Summary"].values if row[0] == "Point C")
        self.assertEqual(point_c_row[1], "29.33, 47.93")

        admin_result = self.client.get(f"/results/{request_id}")
        self.assertIn(b"API Directions", admin_result.data)
        self.assertIn(f"/admin/requests/{request_id}".encode(), admin_result.data)

    def test_travel_request_prefills_one_a_to_b_route_with_selected_hour_range(self):
        with traffic.app.app_context():
            user_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Travel client','travel-prefill@example.test','x','client')"
            ).lastrowid
            traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path,start_time,end_time) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (user_id, "2025-10-12", "2025-10-12", "2025-10-12", '["travel","speed"]',
                 "Link A to B", "[29.3,47.9,29.4,48.0]", "[[29.375912,47.977401],[29.380000,47.982100]]",
                 "Pending", "link", "[7,8]", "[[29.375912,47.977401],[29.380000,47.982100]]",
                 "07:15:30", "08:45:10"),
            )
            traffic.db().commit()
            source_request_id = traffic.db().execute("SELECT last_insert_rowid()").fetchone()[0]
        response = self.client.get(f"/admin/outscraper?source_request_id={source_request_id}")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Prepared from TrafficLens request", response.data)
        self.assertIn(b"Exact time range: 07:15:30", response.data)
        self.assertIn(b"Hourly periods: 07:00-09:00", response.data)
        self.assertIn(b"Departure and finish use the exact selected times", response.data)
        self.assertIn(b"29.375912, 47.977401", response.data)
        with traffic.app.test_request_context():
            prefill = traffic._outscraper_prefill_from_request(
                traffic.db().execute("SELECT * FROM requests WHERE id=?", (source_request_id,)).fetchone()
            )
        config = integration.validate_configuration(prefill["prefill_config"])
        kuwait = dt.timezone(dt.timedelta(hours=3))
        self.assertEqual(config["routes"], [{
            "origin": "29.375912, 47.977401",
            "destination": "29.380000, 47.982100",
        }])
        self.assertEqual(config["departure_time"],
                         int(dt.datetime(2025, 10, 12, 7, 15, 30, tzinfo=kuwait).timestamp()))
        self.assertEqual(config["finish_time"],
                         int(dt.datetime(2025, 10, 12, 8, 45, 10, tzinfo=kuwait).timestamp()))
        self.assertEqual(config["interval"], 1)
        self.assertEqual(config["travel_mode"], "car")
        api_parameters = dict(integration._query_parameters(config["routes"][0], config))
        self.assertEqual(api_parameters["departure_time"], str(config["departure_time"]))
        self.assertEqual(api_parameters["finish_time"], str(config["finish_time"]))
        self.assertEqual(api_parameters["interval"], "1")

    def test_admin_dashboard_offers_directions_for_travel_and_freeflow_services(self):
        with traffic.app.app_context():
            user_id = traffic.db().execute(
                "INSERT INTO users(name,email,pw,role) VALUES('Button client','directions-button@example.test','x','client')"
            ).lastrowid
            cursor = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path,road_km,point_c,point_c_distance_km) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (user_id, "2026-10-18", "2026-10-18", "2026-10-18", '["travel"]',
                 "Test route", "[29,47,30,48]", "[[29.1,47.1],[29.2,47.2]]",
                 "Pending", "link", "[8]", "[[29.1,47.1],[29.2,47.2]]", 1.2, "[29.3,47.3]", 1.6),
            )
            travel_id = cursor.lastrowid
            cursor = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                (user_id, "2026-10-18", "2026-10-18", "2026-10-18", '["speed"]',
                 "Speed only", "[29,47,30,48]", "[[29.1,47.1],[29.2,47.2]]",
                 "Pending", "link", "[8]", "[[29.1,47.1],[29.2,47.2]]"),
            )
            speed_id = cursor.lastrowid
            cursor = traffic.db().execute(
                "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,status,type,hours,path) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                (user_id, "2026-10-18", "2026-10-18", "2026-10-18", '["freeflow_time","freeflow_speed"]',
                 "Free flow route", "[29,47,30,48]", "[[29.1,47.1],[29.2,47.2]]",
                 "Pending", "link", "[8]", "[[29.1,47.1],[29.2,47.2]]"),
            )
            freeflow_id = cursor.lastrowid
            traffic.db().commit()
        dashboard = self.client.get("/admin")
        self.assertIn(f"/admin/requests/{travel_id}".encode(), dashboard.data)
        self.assertNotIn(f"source_request_id={speed_id}".encode(), dashboard.data)
        self.assertIn(b"API Directions", self.client.get(f"/admin/requests/{travel_id}").data)
        self.assertNotIn(b"API Directions", self.client.get(f"/admin/requests/{speed_id}").data)
        freeflow_page = self.client.get(f"/admin/requests/{freeflow_id}")
        self.assertIn(b"service=freeflow_time", freeflow_page.data)
        self.assertIn(b"service=freeflow_speed", freeflow_page.data)

    def test_study_request_is_rejected_when_route_has_junctions_or_cannot_be_verified(self):
        request_data = {
            "services": ["travel"], "date": "2026-10-21", "start_time": "07:00:00",
            "end_time": "08:00:00", "hours": [7],
            "coords": [[29.3, 47.9], [29.4, 48.0]],
            "path": [[29.3, 47.9], [29.4, 48.0]], "road_km": 1.2,
            "no_junctions": True, "point_c": [29.41, 48.01], "point_c_distance_km": 1.6,
        }
        with self.subTest(reason="junction"):
            response = self.client.post("/api/requests", json={**request_data, "junctions": 1})
            self.assertEqual(response.status_code, 400)
            self.assertIn("intersections", response.get_json()["error"])
        with self.subTest(reason="unverified route"):
            response = self.client.post("/api/requests", json={**request_data, "junctions": None})
            self.assertEqual(response.status_code, 400)
            self.assertIn("could not be verified", response.get_json()["error"])

    def test_traffic_volumes_can_be_requested_with_other_services(self):
        with self.client.session_transaction() as session:
            session.update(uid=self.client_id, name="Volume client", role="client")
        response = self.client.post("/api/requests", json={
            "services": ["volumes", "speed"],
            "date": "2025-12-18",
            "start_time": "07:00:00",
            "end_time": "08:00:00",
            "coords": [[30.138239, 31.395698], [30.137423, 31.393946]],
            "path": [[30.138239, 31.395698], [30.137423, 31.393946]],
            "road_km": 0.19,
            "junctions": 0,
            "no_junctions": True,
            "point_c": [30.129997, 31.383281],
            "point_c_distance_km": 1.7,
        })
        self.assertEqual(response.status_code, 200, response.get_json())
        with traffic.app.app_context():
            saved = traffic.db().execute(
                "SELECT services FROM requests WHERE id=?", (response.get_json()["id"],)
            ).fetchone()
            workflow = traffic.db().execute(
                "SELECT status FROM volumes_workflows WHERE request_id=?",
                (response.get_json()["id"],),
            ).fetchone()
        self.assertEqual(json.loads(saved["services"]), ["volumes", "speed"])
        self.assertEqual(workflow["status"], "waiting_for_inputs")

    def test_study_request_stores_exact_times_and_all_overlapping_hours(self):
        study_page = self.client.get("/study")
        self.assertEqual(study_page.status_code, 200)
        self.assertIn(b'id="startTime" step="1"', study_page.data)
        self.assertIn(b'id="endTime" step="1"', study_page.data)
        request_data = {
            "services": ["travel"], "date": "2026-10-21",
            "start_time": "07:15:30", "end_time": "09:00:00",
            "coords": [[29.3, 47.9], [29.4, 48.0]],
            "path": [[29.3, 47.9], [29.4, 48.0]], "road_km": 1.2,
            "junctions": 0, "no_junctions": True, "point_c": [29.41, 48.01], "point_c_distance_km": 1.6,
        }
        invalid_time = self.client.post("/api/requests", json={
            **request_data, "start_time": "09:00:00", "end_time": "07:15:30",
        })
        self.assertEqual(invalid_time.status_code, 400)
        response = self.client.post("/api/requests", json=request_data)
        self.assertEqual(response.status_code, 200, response.get_json())
        with traffic.app.app_context():
            saved = traffic.db().execute(
                "SELECT hours,start_time,end_time,point_c,point_c_distance_km,show_point_c FROM requests WHERE id=?",
                (response.get_json()["id"],),
            ).fetchone()
        self.assertEqual(__import__("json").loads(saved["hours"]), [7, 8])
        self.assertEqual(saved["start_time"], "07:15:30")
        self.assertEqual(saved["end_time"], "09:00:00")
        self.assertEqual(__import__("json").loads(saved["point_c"]), [29.41, 48.01])
        self.assertEqual(saved["point_c_distance_km"], 1.6)
        self.assertEqual(saved["show_point_c"], 0)
        detail = self.client.get(f"/admin/requests/{response.get_json()['id']}")
        self.assertEqual(detail.status_code, 200)
        self.assertIn(b"Point C", detail.data)
        self.assertIn(b"29.410000, 48.010000", detail.data)

    def test_time_or_speed_request_requires_valid_point_c(self):
        payload = {
            "services": ["speed"], "date": "2026-10-21",
            "start_time": "07:00:00", "end_time": "08:00:00",
            "coords": [[29.3, 47.9], [29.4, 48.0]],
            "path": [[29.3, 47.9], [29.4, 48.0]], "road_km": 1.2,
            "junctions": 0, "no_junctions": True,
        }
        response = self.client.post("/api/requests", json=payload)
        self.assertEqual(response.status_code, 400)
        self.assertIn("point C", response.get_json()["error"])
        too_short = self.client.post("/api/requests", json={
            **payload, "point_c": [29.41, 48.01], "point_c_distance_km": 1.5,
        })
        self.assertEqual(too_short.status_code, 400)
        valid = self.client.post("/api/requests", json={
            **payload, "point_c": [29.41, 48.01], "point_c_distance_km": 1.5001,
        })
        self.assertEqual(valid.status_code, 200, valid.get_json())

    def test_existing_directions_job_can_be_copied_for_edit_and_resubmission(self):
        preview = self.preview()
        response = self.client.get(f"/admin/outscraper?duplicate_job_id={preview['id']}")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Copied configuration from Directions API request", response.data)
        self.assertIn(b"Routes and options below are editable", response.data)


if __name__ == "__main__":
    unittest.main()
    