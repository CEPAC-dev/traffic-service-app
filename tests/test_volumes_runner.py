import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

import volumes_runner as runner


class VolumeRunnerTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.workbook_path = self.root / "Inquiries_12-10-2026.xlsx"
        self.models_path = self.root / "Models.mat"
        self.models_path.write_bytes(b"test model placeholder")
        self.run_root = self.root / "app"
        self.run_root.mkdir()
        self.create_workbook()
        self.patches = [
            mock.patch.object(runner, "APP_DIR", self.run_root),
            mock.patch.object(runner, "INPUT_TEMPLATE", self.workbook_path),
            mock.patch.object(runner, "MODELS_FILE", self.models_path),
        ]
        for patcher in self.patches:
            patcher.start()

    def tearDown(self):
        for patcher in reversed(self.patches):
            patcher.stop()
        self.temp_dir.cleanup()

    def create_workbook(self, rows=None, sheet_name="Inqueries", headers=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(headers or ["R", "Lanes", "FFS", "Actual Speed"])
        for row in rows or []:
            sheet.append(row)
        workbook.save(self.workbook_path)
        workbook.close()

    def read_rows(self, path, sheet_name="Inqueries"):
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(workbook[sheet_name].iter_rows(values_only=True))
        finally:
            workbook.close()

    def write_result_workbook(self, path):
        result_book = Workbook()
        sheet = result_book.active
        sheet.title = "Volumes"
        sheet.append([
            "R", "Lanes", "FFS", "Actual Speed", "Hour",
            "Estimated Volume", "RMSE Tolerance",
        ])
        sheet.append([7, 2, 40, 25.5, 8, 120.5, 6.2])
        sheet.append([7, 2, 40, 28, 9, 140.75, 7.1])
        result_book.save(path)
        result_book.close()

    def request_rows(self):
        return [
            {"hour": 8, "lanes": 2, "freeflow_speed": 40, "actual_speed": 25.5},
            {"hour": 9, "lanes": 2, "freeflow_speed": 40, "actual_speed": 28.0},
        ]

    def test_prepare_stages_rows_in_the_clients_own_request_folder(self):
        self.create_workbook(rows=[
            [7, 1, 35, 22],
            [42, 3, 55, 40],
        ])
        before = self.workbook_path.read_bytes()

        prepared = runner.prepare_volume_run(7, self.request_rows(), 5)

        self.assertEqual(self.workbook_path.read_bytes(), before)
        run_dir = self.run_root / "instance" / "volume_runs" / "client-5" / "REQ-7"
        self.assertEqual(Path(prepared["input_workbook"]).parent, run_dir)
        self.assertEqual(Path(prepared["result_workbook"]).parent, run_dir)
        self.assertEqual(prepared["hours"], [8, 9])
        self.assertEqual(
            self.read_rows(prepared["input_workbook"]),
            [
                ("R", "Lanes", "FFS", "Actual Speed", "Hour"),
                (7, 2, 40, 25.5, 8),
                (7, 2, 40, 28, 9),
            ],
        )

    def test_each_clients_request_gets_a_separate_workbook(self):
        first = runner.prepare_volume_run(7, self.request_rows(), 5)
        other_rows = [{"hour": 10, "lanes": 3, "freeflow_speed": 60, "actual_speed": 44.0}]
        second = runner.prepare_volume_run(8, other_rows, 6)

        self.assertNotEqual(Path(first["input_workbook"]).parent, Path(second["input_workbook"]).parent)
        self.assertEqual([row[0] for row in self.read_rows(first["input_workbook"])[1:]], [7, 7])
        self.assertEqual(self.read_rows(second["input_workbook"])[1:], [(8, 3, 60, 44, 10)])
        self.assertEqual(self.read_rows(self.workbook_path), [("R", "Lanes", "FFS", "Actual Speed")])

    def test_saved_result_is_never_reused_across_clients(self):
        prepared = runner.prepare_volume_run(7, self.request_rows(), 5)
        self.write_result_workbook(prepared["result_workbook"])
        self.assertIsNotNone(runner.find_existing_volume_result(7, self.request_rows(), 5))
        self.assertIsNone(runner.find_existing_volume_result(7, self.request_rows(), 6))

    def test_client_id_is_required(self):
        for bad in (None, 0, -1, True, "5"):
            with self.assertRaisesRegex(ValueError, "client ID"):
                runner.prepare_volume_run(7, self.request_rows(), bad)
        self.assertFalse((self.run_root / "instance" / "volume_runs").exists())

    def test_template_open_in_excel_does_not_block_a_run(self):
        lock_path = self.workbook_path.with_name("~$" + self.workbook_path.name)
        lock_path.touch()
        try:
            prepared = runner.prepare_volume_run(7, self.request_rows(), 5)
        finally:
            lock_path.unlink()
        self.assertTrue(Path(prepared["input_workbook"]).is_file())

    def test_missing_sheet_and_wrong_headers_are_rejected_without_mutation(self):
        self.create_workbook(sheet_name="Other")
        before = self.workbook_path.read_bytes()
        with self.assertRaisesRegex(RuntimeError, "Inqueries sheet"):
            runner.prepare_volume_run(7, self.request_rows(), 5)
        self.assertEqual(self.workbook_path.read_bytes(), before)

        self.create_workbook(headers=["Lanes", "R", "FFS", "Actual Speed"])
        before = self.workbook_path.read_bytes()
        with self.assertRaisesRegex(RuntimeError, "must begin with columns"):
            runner.prepare_volume_run(7, self.request_rows(), 5)
        self.assertEqual(self.workbook_path.read_bytes(), before)

    def test_missing_workbook_and_invalid_values_are_rejected(self):
        self.workbook_path.unlink()
        with self.assertRaisesRegex(RuntimeError, "workbook is missing"):
            runner.prepare_volume_run(7, self.request_rows(), 5)
        self.assertFalse((self.run_root / "instance" / "volume_runs").exists())

        self.create_workbook(rows=[[42, 3, 55, 40]])
        before = self.workbook_path.read_bytes()
        invalid_rows = self.request_rows()
        invalid_rows[0]["lanes"] = 0
        with self.assertRaisesRegex(ValueError, "greater than zero"):
            runner.prepare_volume_run(7, invalid_rows, 5)
        self.assertEqual(self.workbook_path.read_bytes(), before)

    def test_unsorted_hour_inputs_are_staged_in_hour_order(self):
        rows = list(reversed(self.request_rows()))
        prepared = runner.prepare_volume_run(7, rows, 5)
        staged = self.read_rows(prepared["input_workbook"])
        self.assertEqual(prepared["hours"], [8, 9])
        self.assertEqual(staged[0][4], "Hour")
        self.assertEqual([row[3] for row in staged[1:]], [25.5, 28])
        self.assertEqual([row[4] for row in staged[1:]], [8, 9])

    def test_excel_float_round_trip_precision_does_not_fail_write_verification(self):
        precise_speed = 29.420689655172414
        prepared = runner.prepare_volume_run(7, [
            {
                "hour": 8,
                "lanes": 1.0,
                "freeflow_speed": precise_speed,
                "actual_speed": precise_speed,
            },
        ], 5)
        self.assertEqual(prepared["hours"], [8])
        self.assertAlmostEqual(
            self.read_rows(prepared["input_workbook"])[1][2],
            precise_speed,
            places=12,
        )

    def test_output_without_hour_values_is_mapped_to_current_request_hours(self):
        prepared = runner.prepare_volume_run(7, self.request_rows(), 5)

        def fake_matlab(*args, **kwargs):
            result_path = Path(prepared["result_workbook"])
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Volumes"
            sheet.append([
                "R", "Lanes", "FFS", "Actual Speed", "Hour",
                "Estimated Volume", "RMSE Tolerance",
            ])
            sheet.append([7, 2, 40, 25.5, None, 100, 5])
            sheet.append([7, 2, 40, 28, None, 120, 6])
            workbook.save(result_path)
            workbook.close()
            return mock.Mock(returncode=0, stdout="", stderr="")

        with mock.patch.object(runner, "_matlab_executable", return_value="matlab"), \
                mock.patch.object(runner.subprocess, "run", side_effect=fake_matlab) as matlab:
            result = runner.run_prepared_volume_estimation(prepared)

        self.assertEqual([item["hour"] for item in result["hourly"]], [8, 9])
        self.assertEqual([item["value"] for item in result["hourly"]], [100, 120])
        self.assertEqual(matlab.call_count, 1)

    def test_matlab_failure_is_reported(self):
        prepared = runner.prepare_volume_run(7, self.request_rows(), 5)
        failed_run = mock.Mock(returncode=1, stdout="", stderr="model error")
        with mock.patch.object(runner, "_matlab_executable", return_value="matlab"), \
                mock.patch.object(runner.subprocess, "run", return_value=failed_run):
            with self.assertRaisesRegex(RuntimeError, "MATLAB volume calculation failed.*model error"):
                runner.run_prepared_volume_estimation(prepared)

    def test_matching_saved_matlab_result_is_reused_without_running_matlab(self):
        prepared = runner.prepare_volume_run(7, self.request_rows(), 5)
        result_book = Workbook()
        sheet = result_book.active
        sheet.title = "Volumes"
        sheet.append([
            "R", "Lanes", "FFS", "Actual Speed", "Hour",
            "Estimated Volume", "RMSE Tolerance",
        ])
        sheet.append([7, 2, 40, 25.5, 8, 120.5, 6.2])
        sheet.append([7, 2, 40, 28, 9, 140.75, 7.1])
        result_book.save(prepared["result_workbook"])
        result_book.close()

        with mock.patch.object(runner.subprocess, "run") as matlab:
            result = runner.find_existing_volume_result(7, self.request_rows(), 5)

        self.assertIsNotNone(result)
        self.assertTrue(result["reused_existing_result"])
        self.assertEqual(result["hourly"], [
            {"hour": 8, "value": 120.5, "tolerance": 6.2},
            {"hour": 9, "value": 140.75, "tolerance": 7.1},
        ])
        matlab.assert_not_called()

    def test_saved_matlab_result_is_not_reused_for_different_inputs(self):
        prepared = runner.prepare_volume_run(7, self.request_rows(), 5)
        result_book = Workbook()
        sheet = result_book.active
        sheet.title = "Volumes"
        sheet.append([
            "R", "Lanes", "FFS", "Actual Speed", "Hour",
            "Estimated Volume", "RMSE Tolerance",
        ])
        sheet.append([7, 2, 40, 25.5, 8, 120.5, 6.2])
        sheet.append([7, 2, 40, 28, 9, 140.75, 7.1])
        result_book.save(prepared["result_workbook"])
        result_book.close()

        changed_rows = self.request_rows()
        changed_rows[0]["actual_speed"] = 26.0
        self.assertIsNone(runner.find_existing_volume_result(7, changed_rows, 5))


if __name__ == "__main__":
    unittest.main()
