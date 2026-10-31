"""Prepare and run the MATLAB traffic-volume model."""

import math
import os
import shutil
import subprocess
import time
import uuid
from pathlib import Path

from openpyxl import load_workbook


APP_DIR = Path(__file__).resolve().parent
INPUT_TEMPLATE = APP_DIR / "Inquiries_12-10-2026.xlsx"
MODELS_FILE = APP_DIR / "Models.mat"
MATLAB_FUNCTION = "Link_Volumes_12_10_2026"
RUN_TIMEOUT_SECONDS = 300
REQUIRED_COLUMNS = ("R", "Lanes", "FFS", "Actual Speed")


def _matlab_executable():
    configured = os.environ.get("TRAFFICLENS_MATLAB_EXECUTABLE", "").strip()
    if configured:
        executable = Path(configured)
        if executable.is_file():
            return str(executable)
        raise RuntimeError("TRAFFICLENS_MATLAB_EXECUTABLE does not point to a MATLAB executable.")

    found = shutil.which("matlab")
    if found:
        return found

    program_files = Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
    installations = program_files / "MATLAB"
    candidates = sorted(
        installations.glob("R*/bin/matlab.exe"),
        key=lambda item: item.parent.parent.name,
        reverse=True,
    )
    if candidates:
        return str(candidates[0])
    raise RuntimeError(
        "MATLAB was not found. Install MATLAB on the app server or configure "
        "TRAFFICLENS_MATLAB_EXECUTABLE."
    )


def _matlab_string(value):
    return str(value).replace("'", "''")


def _finite_number(value, field):
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field} must be numeric.") from exc
    if not math.isfinite(number):
        raise ValueError(f"{field} must be finite.")
    return number


def volume_runs_root():
    return APP_DIR / "instance" / "volume_runs"


def _positive_int(value, field):
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        raise ValueError(f"A valid {field} is required for the MATLAB input.")
    return value


def request_run_dir(client_id, request_id):
    """Each client's request gets its own folder, so no workbook is shared between clients."""
    client_id = _positive_int(client_id, "client ID")
    request_id = _positive_int(request_id, "request ID")
    return volume_runs_root() / f"client-{client_id}" / f"REQ-{request_id}"


def _paired_input_name(path):
    """The input workbook that belongs to a ...-results.xlsx file."""
    path = Path(path)
    if path.name.endswith("-results.xlsx"):
        return path.with_name(path.name[: -len("-results.xlsx")] + ".xlsx")
    return path


def cleanup_request_runs(client_id, request_id, keep_paths):
    """Delete the superseded run workbooks of one request, keeping only the given results."""
    run_dir = request_run_dir(client_id, request_id)
    if not run_dir.is_dir():
        return 0
    keep = set()
    for path in keep_paths:
        if path:
            keep.add(Path(path).resolve())
            keep.add(_paired_input_name(path).resolve())
    removed = 0
    for path in run_dir.iterdir():
        if path.is_file() and path.suffix == ".xlsx" and path.resolve() not in keep:
            try:
                path.unlink()
                removed += 1
            except OSError:
                pass
    return removed


def purge_stale_runs(referenced_paths, max_age_days, now=None):
    """Delete run workbooks older than max_age_days that no workflow still points to."""
    if not max_age_days or max_age_days <= 0:
        return 0
    root = volume_runs_root()
    if not root.is_dir():
        return 0
    keep = set()
    for path in referenced_paths:
        if path:
            keep.add(Path(path).resolve())
            keep.add(_paired_input_name(path).resolve())
    cutoff = (now if now is not None else time.time()) - max_age_days * 86400
    removed = 0
    for path in root.rglob("*.xlsx"):
        try:
            if path.is_file() and path.resolve() not in keep and path.stat().st_mtime < cutoff:
                path.unlink()
                removed += 1
        except OSError:
            continue
    return removed


def _normalize_rows(request_id, rows):
    if isinstance(request_id, bool) or not isinstance(request_id, int) or request_id <= 0:
        raise ValueError("A valid request ID is required for the MATLAB input.")
    if not isinstance(rows, list) or not rows:
        raise ValueError("At least one hourly input row is required.")

    normalized = []
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError("Each volume input row must be an object.")
        hour = row.get("hour")
        if isinstance(hour, bool) or not isinstance(hour, int) or not 0 <= hour <= 23:
            raise ValueError("Each volume input must have a valid study hour.")
        lanes = _finite_number(row.get("lanes"), "Lanes")
        freeflow_speed = _finite_number(row.get("freeflow_speed"), "Free Flow Speed")
        actual_speed = _finite_number(row.get("actual_speed"), "Actual Speed")
        if lanes <= 0 or freeflow_speed <= 0 or actual_speed <= 0:
            raise ValueError("Lanes and speeds must be greater than zero.")
        normalized.append((request_id, lanes, freeflow_speed, actual_speed))

    hours = [row["hour"] for row in rows]
    if len(hours) != len(set(hours)):
        raise ValueError("Each selected study hour must have exactly one actual speed.")
    ordered_inputs = sorted(zip(hours, normalized), key=lambda item: item[0])
    return [row for _, row in ordered_inputs], [hour for hour, _ in ordered_inputs]


def _verify_headers(sheet):
    headers = tuple(sheet.cell(row=1, column=index).value for index in range(1, 5))
    if headers != REQUIRED_COLUMNS:
        raise RuntimeError(
            "The Inqueries sheet must begin with columns R, Lanes, FFS, Actual Speed "
            "in that order."
        )


def _read_request_rows(workbook_path, request_id):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Inqueries" not in workbook.sheetnames:
            raise RuntimeError("The input workbook does not contain the Inqueries sheet.")
        sheet = workbook["Inqueries"]
        _verify_headers(sheet)
        return [
            tuple(values[:4])
            for values in sheet.iter_rows(min_row=2, max_col=4, values_only=True)
            if values[0] == request_id
        ]
    finally:
        workbook.close()


def _rows_match_after_excel_round_trip(expected_rows, actual_rows):
    if len(expected_rows) != len(actual_rows):
        return False
    for expected, actual in zip(expected_rows, actual_rows):
        if len(expected) != len(actual):
            return False
        for expected_value, actual_value in zip(expected, actual):
            if isinstance(expected_value, (int, float)) and not isinstance(expected_value, bool):
                try:
                    if not math.isclose(
                        float(expected_value),
                        float(actual_value),
                        rel_tol=1e-12,
                        abs_tol=1e-12,
                    ):
                        return False
                except (TypeError, ValueError):
                    return False
            elif expected_value != actual_value:
                return False
    return True


def _read_result_workbook(output_path, request_id, expected_rows, expected_hours):
    result_book = load_workbook(output_path, read_only=True, data_only=True)
    try:
        if "Volumes" not in result_book.sheetnames:
            raise RuntimeError("MATLAB output workbook is missing the Volumes sheet.")
        result_sheet = result_book["Volumes"]
        header_values = next(
            result_sheet.iter_rows(min_row=1, max_row=1, values_only=True),
            (),
        )
        headers = [str(value or "").strip() for value in header_values]
        header_indexes = {name.casefold(): index for index, name in enumerate(headers)}
        required = {
            "r", "lanes", "ffs", "actual speed", "hour",
            "estimated volume", "rmse tolerance",
        }
        if not required.issubset(header_indexes):
            raise RuntimeError("MATLAB output is missing required input or result columns.")

        outputs = []
        for values in result_sheet.iter_rows(min_row=2, values_only=True):
            if not values or all(value is None for value in values):
                continue
            output_hour = values[header_indexes["hour"]]
            if output_hour is None or (
                isinstance(output_hour, float) and math.isnan(output_hour)
            ):
                output_index = len(outputs)
                if output_index >= len(expected_hours):
                    raise RuntimeError("MATLAB returned more rows than the current request contains.")
                hour = expected_hours[output_index]
            else:
                hour = int(_finite_number(output_hour, "Output hour"))
            row_inputs = (
                _finite_number(values[header_indexes["r"]], "Output route ID"),
                _finite_number(values[header_indexes["lanes"]], "Output lanes"),
                _finite_number(values[header_indexes["ffs"]], "Output free-flow speed"),
                _finite_number(values[header_indexes["actual speed"]], "Output actual speed"),
            )
            expected_index = len(outputs)
            if expected_index >= len(expected_rows):
                raise RuntimeError("MATLAB returned more rows than the current request contains.")
            if hour != expected_hours[expected_index] or not _rows_match_after_excel_round_trip(
                [expected_rows[expected_index]], [row_inputs]
            ):
                raise RuntimeError("MATLAB output does not match the current request inputs.")
            if row_inputs[0] != request_id:
                raise RuntimeError("MATLAB output contains a different request ID.")
            volume = _finite_number(
                values[header_indexes["estimated volume"]], "Estimated volume"
            )
            tolerance = _finite_number(
                values[header_indexes["rmse tolerance"]], "RMSE tolerance"
            )
            outputs.append({"hour": hour, "value": volume, "tolerance": tolerance})

        if [row["hour"] for row in outputs] != expected_hours:
            raise RuntimeError("MATLAB output does not match the requested study hours.")
        return outputs
    finally:
        result_book.close()


def find_existing_volume_result(request_id, rows, client_id):
    """Return a saved result only when its input and output match this request exactly."""
    normalized, hours = _normalize_rows(request_id, rows)
    run_dir = request_run_dir(client_id, request_id)
    if not run_dir.is_dir():
        return None

    expected_rows = [
        (*row, hour) for row, hour in zip(normalized, hours)
    ]
    for input_path in sorted(run_dir.glob(f"REQ-{request_id}-*.xlsx")):
        if input_path.stem.endswith("-results"):
            continue
        result_path = input_path.with_name(f"{input_path.stem}-results.xlsx")
        if not result_path.is_file():
            continue
        try:
            input_book = load_workbook(input_path, read_only=True, data_only=True)
            try:
                if "Inqueries" not in input_book.sheetnames:
                    continue
                input_sheet = input_book["Inqueries"]
                _verify_headers(input_sheet)
                header = tuple(
                    input_sheet.cell(row=1, column=index).value
                    for index in range(1, 6)
                )
                if header != (*REQUIRED_COLUMNS, "Hour"):
                    continue
                actual_rows = [
                    tuple(values[:5])
                    for values in input_sheet.iter_rows(
                        min_row=2, max_col=5, values_only=True
                    )
                    if any(value is not None for value in values)
                ]
            finally:
                input_book.close()
            if not _rows_match_after_excel_round_trip(expected_rows, actual_rows):
                continue
            outputs = _read_result_workbook(
                result_path, request_id, normalized, hours
            )
            return {
                "request_id": request_id,
                "hourly": outputs,
                "input_workbook": str(input_path),
                "result_workbook": str(result_path),
                "reused_existing_result": True,
            }
        except (OSError, ValueError, RuntimeError):
            continue
    return None


def prepare_volume_run(request_id, rows, client_id):
    """Stage a request-specific MATLAB input workbook inside the owning client's folder.

    The template workbook is only read. Nothing from any request is written back to it,
    so rows from different clients never end up in the same file.
    """
    if not INPUT_TEMPLATE.is_file():
        raise RuntimeError(f"The MATLAB input workbook is missing: {INPUT_TEMPLATE.name}")
    if not MODELS_FILE.is_file():
        raise RuntimeError(f"The MATLAB model data is missing: {MODELS_FILE.name}")
    normalized, hours = _normalize_rows(request_id, rows)
    run_dir = request_run_dir(client_id, request_id)

    run_id = uuid.uuid4().hex
    workbook_path = run_dir / f"REQ-{int(request_id)}-{run_id}.xlsx"
    output_path = run_dir / f"REQ-{int(request_id)}-{run_id}-results.xlsx"
    temporary_run = run_dir / f".REQ-{int(request_id)}-{run_id}.tmp.xlsx"

    run_dir.mkdir(parents=True, exist_ok=True)
    try:
        template = load_workbook(INPUT_TEMPLATE)
        try:
            if "Inqueries" not in template.sheetnames:
                raise RuntimeError("The input workbook does not contain the Inqueries sheet.")
            _verify_headers(template["Inqueries"])
            run_sheet = template["Inqueries"]
            if run_sheet.max_row > 1:
                run_sheet.delete_rows(2, run_sheet.max_row - 1)
            run_sheet.cell(row=1, column=5, value="Hour")
            for row, hour in zip(normalized, hours):
                run_sheet.append((*row, hour))
            template.save(temporary_run)
        finally:
            template.close()
        os.replace(temporary_run, workbook_path)
        saved_rows = _read_request_rows(workbook_path, request_id)
        if not _rows_match_after_excel_round_trip(normalized, saved_rows):
            raise RuntimeError(
                "The MATLAB input rows were written but could not be verified in the request workbook."
            )
    finally:
        try:
            temporary_run.unlink()
        except FileNotFoundError:
            pass

    return {
        "request_id": request_id,
        "client_id": client_id,
        "hours": hours,
        "normalized_rows": normalized,
        "input_workbook": str(workbook_path),
        "result_workbook": str(output_path),
    }


def run_prepared_volume_estimation(prepared_run):
    """Run MATLAB using a previously saved and verified request-specific workbook."""
    workbook_path = Path(prepared_run["input_workbook"])
    output_path = Path(prepared_run["result_workbook"])
    if not workbook_path.is_file():
        raise RuntimeError("The saved MATLAB input workbook is no longer available.")
    if not MODELS_FILE.is_file():
        raise RuntimeError(f"The MATLAB model data is missing: {MODELS_FILE.name}")

    call = (
        f"{MATLAB_FUNCTION}('{_matlab_string(workbook_path)}',"
        f"'{_matlab_string(output_path)}','{_matlab_string(MODELS_FILE)}')"
    )
    completed = subprocess.run(
        [_matlab_executable(), "-batch", call],
        cwd=APP_DIR,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=RUN_TIMEOUT_SECONDS,
        check=False,
    )
    if completed.returncode:
        detail = (completed.stderr or completed.stdout or "").strip()
        raise RuntimeError(
            "MATLAB volume calculation failed."
            + (f" {detail[-1200:]}" if detail else "")
        )
    if not output_path.is_file():
        raise RuntimeError("MATLAB completed without creating the volume results workbook.")

    expected_hours = prepared_run["hours"]
    outputs = _read_result_workbook(
        output_path,
        prepared_run["request_id"],
        prepared_run["normalized_rows"],
        expected_hours,
    )
    return {
        "request_id": prepared_run["request_id"],
        "hourly": outputs,
        "input_workbook": str(workbook_path),
        "result_workbook": str(output_path),
    }


def run_volume_estimation(request_id, rows, client_id):
    """Persist the current request and run the model for its selected hours."""
    prepared_run = prepare_volume_run(request_id, rows, client_id)
    return run_prepared_volume_estimation(prepared_run)
