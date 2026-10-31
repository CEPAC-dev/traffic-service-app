import csv, hashlib, io, json, math, os, re, secrets, sqlite3, threading, time, warnings, datetime as dt
from functools import lru_cache, wraps
from pathlib import Path
from flask import Flask, g, render_template, request, redirect, url_for, session, jsonify, abort, flash
from flask import send_file
from werkzeug.security import generate_password_hash, check_password_hash
from dotenv import load_dotenv
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
import urllib.error
from urllib.parse import urlencode
from urllib.request import Request as UrlRequest, urlopen
import volumes_runner
try:
    import holidays
except ModuleNotFoundError:
    holidays = None
import network
import outscraper_integration as outscraper
try:
    from timezonefinder import TimezoneFinder
except ModuleNotFoundError:
    TimezoneFinder = None
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageOps, UnidentifiedImageError

load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))
app = Flask(__name__)
configured_secret = os.environ.get("SECRET_KEY", "").strip()
if configured_secret == "generate-a-random-secret-for-local-use":
    configured_secret = ""
if configured_secret and len(configured_secret) < 32:
    raise RuntimeError("SECRET_KEY must contain at least 32 characters; generate a random key locally.")
app.secret_key = configured_secret or secrets.token_hex(32)
app.config["EPHEMERAL_SECRET_KEY"] = not bool(configured_secret)
if not configured_secret:
    app.logger.warning("SECRET_KEY is not configured; using an ephemeral key for this local run.")
app.config["MAX_CONTENT_LENGTH"] = 300 * 1024 * 1024   # network files can be large
app.config["OUTSCRAPER_AUTO_POLL"] = os.environ.get("TRAFFICLENS_OUTSCRAPER_AUTO_POLL", "1") != "0"
DB = os.environ.get("TRAFFICLENS_DB") or os.path.join(os.path.dirname(os.path.abspath(__file__)), "mobility.db")
SERVICES = {"lanes": "Number of lanes", "speed": "Speed", "travel": "Travel time", "length": "Length",
            "freeflow_time": "Free flow time", "freeflow_speed": "Free flow speed", "capacity": "Capacity",
            "volumes": "Traffic volumes"}
PROFILE_JOB_TITLES = (
    "Engineer", "Data Analyst", "Researcher", "Student", "Manager", "Other",
)
PROFILE_PHOTO_MAX_BYTES = 5 * 1024 * 1024
PROFILE_PHOTO_SIZE = (512, 512)
ENGINEERING_SPECIALIZATIONS = (
    "Civil Engineering", "Transportation Engineering", "Traffic Engineering",
    "Highway Engineering", "Structural Engineering", "Other",
)
HIDDEN_SERVICES_BY_DEFAULT = {"capacity"}
STATUSES = ["Pending", "Processing", "Completed", "Rejected"]
APP_NAME = "TrafficLens"   # change the app name here (shown in the header, page title and landing page)

# Default explanation of each service. The admin can edit these texts from Admin > Services.
SERVICE_DEFAULTS = {
 "lanes": {"icon": "\U0001F6E3\uFE0F", "desc": "The number of traffic lanes on the link."},
 "speed": {"icon": "\u26A1", "desc": "The average speed of vehicles on the link during each hour you selected."},
 "travel": {"icon": "\u23F1\uFE0F", "desc": "How long it takes to drive the whole link during each hour you selected."},
 "length": {"icon": "\U0001F4CF", "desc": "The length of the link, from its start (A) to its end (B)."},
 "freeflow_time": {"icon": "\U0001F552", "desc": "Travel time at 7:00 AM on the nearest public holiday in the route's country."},
 "freeflow_speed": {"icon": "\U0001F680", "desc": "Travel speed at 7:00 AM on the latest weekly weekend date before today in the route's country."},
 "capacity": {"icon": "\U0001F699", "desc": "The maximum number of vehicles that can pass through the link in one hour."},
 "volumes": {"icon": "\U0001F4CA", "desc": "Estimated hourly traffic volume calculated from lanes, actual speed, and weekly-weekend free-flow speed using the traffic model."},
}

# One form per service. kind="static": one value for the whole link. kind="hourly": one value per selected hour.
def _f(kind, label, unit, hint): return {"kind": kind, "label": label, "unit": unit, "hint": hint}
FORMS = {
 "lanes": _f("static", "Number of lanes", "lanes", "Enter the number of lanes of the link."),
 "length": _f("static", "Length", "km", "Enter the length of the link in kilometres."),
 "freeflow_speed": _f("static", "Free flow speed", "km/h", "Use Directions API on the latest weekly weekend date before today at 7:00 AM in the route's country."),
 "freeflow_time": _f("static", "Free flow time", "min", "Use Directions API for the nearest public holiday at 7:00 AM in the route's country."),
 "capacity": _f("static", "Capacity", "veh/h", "Enter the maximum number of vehicles per hour the link can carry."),
 "volumes": _f("hourly", "Estimated traffic volume", "veh/h", "Calculated from the traffic model after administrator review."),
 "speed": _f("hourly", "Average speed", "km/h", "Enter the average speed for each hour the client selected."),
 "travel": _f("hourly", "Travel time", "min", "Enter the travel time of the whole link for each hour the client selected."),
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, name TEXT, email TEXT UNIQUE, pw TEXT, role TEXT DEFAULT 'client');
CREATE TABLE IF NOT EXISTS requests(id INTEGER PRIMARY KEY, user_id INTEGER, created TEXT, d1 TEXT, d2 TEXT,
  services TEXT, area TEXT, bbox TEXT, coords TEXT, status TEXT DEFAULT 'Pending', results TEXT,
  type TEXT DEFAULT 'area', reply TEXT, reply_at TEXT, data TEXT, hours TEXT, path TEXT, road_km REAL, junctions INTEGER,
  start_time TEXT, end_time TEXT, point_c TEXT, point_c_distance_km REAL,
  show_point_c INTEGER NOT NULL DEFAULT 0, title TEXT,
  junction_quality_ack INTEGER NOT NULL DEFAULT 0, junction_quality_ack_at TEXT);
CREATE TABLE IF NOT EXISTS service_info(key TEXT PRIMARY KEY, description TEXT);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS notifications(
  id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id),
  request_id INTEGER NOT NULL REFERENCES requests(id), kind TEXT NOT NULL,
  title TEXT NOT NULL, message TEXT NOT NULL, created_at TEXT NOT NULL,
  read_at TEXT
);
CREATE INDEX IF NOT EXISTS idx_notifications_user_unread ON notifications(user_id, read_at, id);
CREATE TABLE IF NOT EXISTS outscraper_jobs(
  id TEXT PRIMARY KEY, created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
  config_json TEXT NOT NULL, config_hash TEXT NOT NULL, approval_token_hash TEXT NOT NULL,
  approval_status TEXT NOT NULL DEFAULT 'pending', execution_status TEXT NOT NULL DEFAULT 'awaiting_approval',
  approved_at TEXT, rejected_at TEXT, started_at TEXT, completed_at TEXT,
  estimated_cost REAL, estimated_currency TEXT, estimate_confidence TEXT NOT NULL DEFAULT 'unknown',
  actual_cost REAL, actual_currency TEXT, provider_ids TEXT NOT NULL DEFAULT '[]',
  error TEXT, result_count INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS outscraper_job_routes(
  id INTEGER PRIMARY KEY, job_id TEXT NOT NULL REFERENCES outscraper_jobs(id),
  route_index INTEGER NOT NULL, origin TEXT NOT NULL, destination TEXT NOT NULL,
  execution_status TEXT NOT NULL DEFAULT 'awaiting_approval', provider_request_id TEXT,
  provider_status TEXT, submitted_at TEXT, completed_at TEXT, last_polled_at TEXT,
  request_details_json TEXT, raw_response_json TEXT, result_count INTEGER NOT NULL DEFAULT 0, error TEXT,
  UNIQUE(job_id, route_index)
);
CREATE INDEX IF NOT EXISTS idx_outscraper_jobs_created ON outscraper_jobs(created_at);
CREATE INDEX IF NOT EXISTS idx_outscraper_jobs_status ON outscraper_jobs(approval_status, execution_status);
CREATE INDEX IF NOT EXISTS idx_outscraper_routes_status ON outscraper_job_routes(execution_status);
CREATE TABLE IF NOT EXISTS request_directions_workflows(
  job_id TEXT PRIMARY KEY REFERENCES outscraper_jobs(id),
  request_id INTEGER NOT NULL REFERENCES requests(id),
  status TEXT NOT NULL DEFAULT 'awaiting_response',
  interval_seconds INTEGER NOT NULL,
  interval_minutes INTEGER,
  result_json TEXT,
  error TEXT,
  confirmed_by INTEGER REFERENCES users(id),
  confirmed_at TEXT,
  created_at TEXT NOT NULL,
  updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_request_directions_request ON request_directions_workflows(request_id, created_at);
CREATE TABLE IF NOT EXISTS volumes_workflows(
  request_id INTEGER PRIMARY KEY REFERENCES requests(id),
  status TEXT NOT NULL DEFAULT 'waiting_for_inputs',
  lanes REAL,
  actual_result_json TEXT,
  freeflow_speed REAL,
  freeflow_accepted INTEGER NOT NULL DEFAULT 0,
  freeflow_result_json TEXT,
  freeflow_validation_json TEXT,
  freeflow_override INTEGER NOT NULL DEFAULT 0,
  freeflow_accepted_at TEXT,
  freeflow_job_id TEXT,
  result_json TEXT,
  workbook_path TEXT,
  error TEXT,
  created_at TEXT NOT NULL,
  updated_at TEXT NOT NULL,
  reviewed_by INTEGER REFERENCES users(id),
  reviewed_at TEXT,
  sent_at TEXT
);
CREATE TABLE IF NOT EXISTS volume_freeflow_attempts(
  job_id TEXT PRIMARY KEY REFERENCES outscraper_jobs(id),
  request_id INTEGER NOT NULL REFERENCES requests(id),
  submitted_at TEXT NOT NULL,
  status TEXT NOT NULL,
  result_json TEXT,
  validation_json TEXT,
  error TEXT,
  accepted INTEGER NOT NULL DEFAULT 0,
  override INTEGER NOT NULL DEFAULT 0,
  accepted_at TEXT
);
CREATE INDEX IF NOT EXISTS idx_volume_freeflow_attempts_request
  ON volume_freeflow_attempts(request_id,submitted_at);
CREATE TABLE IF NOT EXISTS request_groups(
  id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id), title TEXT,
  created TEXT NOT NULL, study_date TEXT NOT NULL, services TEXT NOT NULL,
  mode TEXT NOT NULL, sent_at TEXT
);
CREATE TABLE IF NOT EXISTS periods(
  id INTEGER PRIMARY KEY, label TEXT NOT NULL UNIQUE,
  is_active INTEGER NOT NULL DEFAULT 1, sort_order INTEGER NOT NULL DEFAULT 0
);
"""

MAX_GROUP_LINKS = 10
GROUP_MODES = ("single", "multi_period", "multi_link")

def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB); g.db.row_factory = sqlite3.Row
    return g.db

def _create_notification(user_id, request_id, kind, title, message):
    hidden_child = kind != "group_ready" and db().execute(
        "SELECT 1 FROM requests r JOIN request_groups g ON g.id=r.group_id "
        "JOIN users u ON u.id=? WHERE r.id=? AND g.sent_at IS NULL AND u.role!='admin'",
        (user_id, request_id),
    ).fetchone()
    if hidden_child:
        return   # a client hears about a grouped order once, when the admin sends the whole report
    db().execute(
        "INSERT INTO notifications(user_id,request_id,kind,title,message,created_at) VALUES(?,?,?,?,?,?)",
        (user_id, request_id, kind, title, message, dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    )

def _notify_admins_new_request(request_id, client_name):
    admins = db().execute("SELECT id FROM users WHERE role='admin'").fetchall()
    for admin_user in admins:
        _create_notification(
            admin_user["id"], request_id, "new_request", "New request",
            f"{client_name} submitted a new request: REQ-{request_id}.",
        )

def _notify_request_owner_reply(request_id, user_id, old_reply, new_reply):
    if new_reply and new_reply != (old_reply or ""):
        _create_notification(
            user_id, request_id, "admin_reply", "New reply to your request",
            f"The administrator replied to request REQ-{request_id}: {new_reply}",
        )

def _service_ready(service, entry):
    entry = entry or {}
    if entry.get("rows"):
        return True
    if service == "volumes":
        return False
    # A Directions API placeholder note ("awaiting administrator review") is not a result.
    if entry.get("source") == "outscraper":
        return False
    return bool(entry.get("note"))

def _all_requested_services_ready(requested_services, data):
    return bool(requested_services) and all(
        _service_ready(service, data.get(service)) for service in requested_services
    )

def _save_volume_lanes_service(request_data, requested_services, lanes):
    if "lanes" in requested_services and lanes is not None:
        request_data["lanes"] = {
            "rows": [{"value": lanes}],
            "note": "Number of lanes used for the Traffic volumes calculation.",
            "source": "admin",
            "updated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        }

@app.teardown_appcontext
def close(_):
    d = g.pop("db", None)
    if d: d.close()

def _translate_legacy_notifications(c):
    """Rewrite notifications stored in Arabic by earlier versions into English. Safe to run on every start."""
    titles = {
        "طلب جديد": "New request",
        "رد جديد على طلبك": "New reply to your request",
        "نتائج أحجام المرور جاهزة": "Traffic volume results are ready",
        "نتائج الطلب جاهزة": "Request results are ready",
    }
    for arabic, english in titles.items():
        c.execute("UPDATE notifications SET title=? WHERE title=?", (english, arabic))
    patterns = (
        (r"^قدم (.+) طلبًا جديدًا: (REQ-\d+)\.$", r"\1 submitted a new request: \2."),
        (r"^أضاف المسؤول ردًا على الطلب (REQ-\d+): (.*)$", r"The administrator replied to request \1: \2"),
        (r"^اعتمد المسؤول نتائج أحجام المرور للطلب (REQ-\d+)\.$", r"The administrator approved the traffic volume results for request \1."),
        (r"^اعتمد المسؤول نتائج الاتجاهات للطلب (REQ-\d+)\.$", r"The administrator approved the directions results for request \1."),
    )
    rows = c.execute(
        "SELECT id,message FROM notifications WHERE message LIKE 'قدم%' OR message LIKE 'أضاف%' OR message LIKE 'اعتمد%'"
    ).fetchall()
    for notification_id, message in rows:
        for pattern, replacement in patterns:
            translated, count = re.subn(pattern, replacement, message or "", flags=re.S)
            if count:
                c.execute("UPDATE notifications SET message=? WHERE id=?", (translated, notification_id))
                break


def _migrate_profile_columns(c):
    user_cols = {column[1] for column in c.execute("PRAGMA table_info(users)")}
    for column in (
        "first_name", "last_name", "phone_number", "job_title",
        "engineering_specialization", "company_name", "bio", "profile_image",
    ):
        if column not in user_cols:
            c.execute(f"ALTER TABLE users ADD COLUMN {column} TEXT")
    for user_id, name, first_name, last_name in c.execute(
        "SELECT id,name,first_name,last_name FROM users"
    ).fetchall():
        if first_name and last_name:
            continue
        name_parts = (name or "").strip().split(None, 1)
        migrated_first = name_parts[0] if name_parts else ""
        migrated_last = name_parts[1] if len(name_parts) > 1 else ""
        c.execute(
            "UPDATE users SET first_name=COALESCE(NULLIF(first_name,''),?), "
            "last_name=COALESCE(NULLIF(last_name,''),?) WHERE id=?",
            (migrated_first, migrated_last, user_id),
        )


def init():
    c = sqlite3.connect(DB); c.executescript(SCHEMA)
    _migrate_profile_columns(c)
    workflow_cols = {column[1] for column in c.execute("PRAGMA table_info(request_directions_workflows)")}
    if "interval_minutes" not in workflow_cols:
        c.execute("ALTER TABLE request_directions_workflows ADD COLUMN interval_minutes INTEGER")
        c.execute("UPDATE request_directions_workflows SET interval_minutes=MAX(1, CAST((interval_seconds + 59) / 60 AS INTEGER))")
    route_cols = [column[1] for column in c.execute("PRAGMA table_info(outscraper_job_routes)")]
    for column in ("last_polled_at", "request_details_json"):
        if column not in route_cols:
            c.execute(f"ALTER TABLE outscraper_job_routes ADD COLUMN {column} TEXT")
    volume_cols = {column[1] for column in c.execute("PRAGMA table_info(volumes_workflows)")}
    for column, definition in (
        ("freeflow_accepted", "INTEGER NOT NULL DEFAULT 0"),
        ("freeflow_result_json", "TEXT"),
        ("freeflow_validation_json", "TEXT"),
        ("freeflow_override", "INTEGER NOT NULL DEFAULT 0"),
        ("freeflow_accepted_at", "TEXT"),
    ):
        if column not in volume_cols:
            c.execute(f"ALTER TABLE volumes_workflows ADD COLUMN {column} {definition}")
    c.execute(
        "UPDATE volumes_workflows SET freeflow_accepted=1 "
        "WHERE freeflow_speed IS NOT NULL AND freeflow_accepted=0 AND freeflow_result_json IS NULL"
    )
    for job_id, created_at, execution_status, config_json in c.execute(
        "SELECT id,created_at,execution_status,config_json FROM outscraper_jobs"
    ).fetchall():
        try:
            context = json.loads(config_json or "{}").get("trafficlens_context") or {}
            if not context.get("volume_input"):
                continue
            if not c.execute(
                "SELECT 1 FROM requests WHERE id=?", (context["source_request_id"],)
            ).fetchone():
                continue
            c.execute(
                "INSERT OR IGNORE INTO volume_freeflow_attempts"
                "(job_id,request_id,submitted_at,status) VALUES(?,?,?,?)",
                (job_id, context["source_request_id"], created_at, execution_status),
            )
        except (KeyError, TypeError, ValueError, json.JSONDecodeError):
            app.logger.warning("Could not backfill free-flow attempt history for job %s.", job_id)
    if app.config["EPHEMERAL_SECRET_KEY"]:
        c.execute(
            "UPDATE outscraper_jobs SET approval_status='invalidated', execution_status='invalidated', updated_at=? "
            "WHERE approval_status='pending' AND execution_status='awaiting_approval'",
            (_timestamp_now(),))
        c.execute("UPDATE outscraper_job_routes SET execution_status='invalidated' WHERE job_id IN "
                  "(SELECT id FROM outscraper_jobs WHERE approval_status='invalidated' AND execution_status='invalidated')")
    c.execute("UPDATE outscraper_jobs SET execution_status='unknown', error='The app restarted while submission was in progress; check provider status before taking further action.' WHERE execution_status='submitting'")
    c.execute("UPDATE outscraper_job_routes SET execution_status='unknown', error='The app restarted while submission was in progress; do not resubmit automatically.' WHERE execution_status='submitting'")
    c.execute(
        "UPDATE outscraper_job_routes SET execution_status='unknown', error='The app restarted before submission status could be recorded; do not resubmit automatically.' "
        "WHERE execution_status='queued' AND job_id IN (SELECT id FROM outscraper_jobs WHERE execution_status='unknown' AND error LIKE 'The app restarted while submission is in progress%')"
    )
    cols = [x[1] for x in c.execute("PRAGMA table_info(requests)")]
    for col, ddl in (("type", "TEXT DEFAULT 'area'"), ("reply", "TEXT"), ("reply_at", "TEXT"), ("data", "TEXT"), ("hours", "TEXT"), ("path", "TEXT"), ("road_km", "REAL"), ("junctions", "INTEGER"), ("start_time", "TEXT"), ("end_time", "TEXT"), ("point_c", "TEXT"), ("point_c_distance_km", "REAL"), ("show_point_c", "INTEGER NOT NULL DEFAULT 0"), ("title", "TEXT"), ("junction_quality_ack", "INTEGER NOT NULL DEFAULT 0"), ("junction_quality_ack_at", "TEXT"), ("group_id", "INTEGER"), ("period_label", "TEXT")):
        if col not in cols: c.execute(f"ALTER TABLE requests ADD COLUMN {col} {ddl}")
    c.execute("CREATE INDEX IF NOT EXISTS idx_requests_group ON requests(group_id)")
    if not c.execute("SELECT 1 FROM periods").fetchone():
        c.executemany("INSERT INTO periods(label,sort_order) VALUES(?,?)", [("AM", 1), ("MD", 2), ("PM", 3)])
    workflow_now = _timestamp_now()
    for request_id, services_json in c.execute("SELECT id,services FROM requests"):
        if "volumes" in json.loads(services_json or "[]"):
            c.execute(
                "INSERT OR IGNORE INTO volumes_workflows(request_id,status,created_at,updated_at) "
                "VALUES(?,'waiting_for_inputs',?,?)",
                (request_id, workflow_now, workflow_now),
            )
    _translate_legacy_notifications(c)
    if not c.execute("SELECT 1 FROM users").fetchone():
        admin_password = os.environ.get("TRAFFICLENS_ADMIN_PASSWORD", "").strip()
        generated_password = not admin_password
        if generated_password:
            admin_password = secrets.token_urlsafe(12)
        c.execute("INSERT INTO users(name,email,pw,role) VALUES(?,?,?,?)",
                  ("Admin", "admin@demo.com", generate_password_hash(admin_password), "admin"))
        if generated_password:
            print(f"[TrafficLens] First-run admin created: admin@demo.com / {admin_password}  "
                  "(shown once; set TRAFFICLENS_ADMIN_PASSWORD to choose your own)", flush=True)
    bundled_network = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Kwait Links.zip")
    setup_done = c.execute("SELECT 1 FROM settings WHERE key='bundled_network_setup_v2'").fetchone()
    network_row = c.execute("SELECT value FROM settings WHERE key='network'").fetchone()
    cfg = json.loads(network_row[0]) if network_row else {}
    configured_file = cfg.get("file") or ""
    is_bundled = bool(configured_file and os.path.isfile(configured_file) and os.path.samefile(configured_file, bundled_network))
    if not setup_done and os.path.isfile(bundled_network) and network.HAVE_GEO and (not configured_file or is_bundled):
        metadata = network.inspect(bundled_network)
        guesses = network.guess(metadata["columns"])
        numeric = network.numeric_columns(bundled_network)
        cfg.update({"file": bundled_network, "filename": "Kwait Links.zip",
                    "uploaded": cfg.get("uploaded") or dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
                    "tol": cfg.get("tol", 25), "min_cov": cfg.get("min_cov", 80),
                    "speed_unit": cfg.get("speed_unit", "km/h"), **metadata})
        for key in ("lanes", "speed", "capacity"):
            if cfg.get(key) not in numeric: cfg[key] = guesses.get(key, "")
        if cfg.get("road_name") not in metadata["columns"]:
            cfg["road_name"] = guesses.get("road_name", "")
        c.execute("INSERT INTO settings(key,value) VALUES('network',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                  (json.dumps(cfg),))
        c.execute("INSERT INTO settings(key,value) VALUES('bundled_network_setup_v2','true')")
    elif not setup_done and configured_file and not is_bundled:
        c.execute("INSERT INTO settings(key,value) VALUES('bundled_network_setup_v2','true')")
    c.commit(); c.close()

_outscraper_workers_lock = threading.Lock()
_outscraper_workers = set()

def _timestamp_now():
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")

def _token_hash(token):
    return hashlib.sha256(token.encode("utf-8")).hexdigest()

def _forget_outscraper_token(job_id):
    tokens = dict(session.get("outscraper_tokens", {}))
    tokens.pop(job_id, None)
    session["outscraper_tokens"] = tokens

def _outscraper_job(job_id):
    row = db().execute("SELECT * FROM outscraper_jobs WHERE id=?", (job_id,)).fetchone()
    if not row: return None
    job = dict(row)
    job["config"] = json.loads(job.pop("config_json"))
    job["routes"] = [dict(route) for route in db().execute(
        "SELECT * FROM outscraper_job_routes WHERE job_id=? ORDER BY route_index", (job_id,)).fetchall()]
    for route in job["routes"]:
        route["result"] = json.loads(route["raw_response_json"]) if route["raw_response_json"] else None
        route["request_details"] = json.loads(route["request_details_json"]) if route.get("request_details_json") else None
    return job

def _refresh_job_status(job_id):
    routes = db().execute(
        "SELECT execution_status, result_count, error FROM outscraper_job_routes WHERE job_id=?",
        (job_id,)).fetchall()
    states = [route["execution_status"] for route in routes]
    if states and all(state == "completed" for state in states):
        status = "completed"
    elif any(state == "running" for state in states):
        status = "running"
    elif any(state == "unknown" for state in states):
        status = "unknown"
    elif any(state == "failed" for state in states):
        status = "failed"
    elif any(state == "abandoned" for state in states):
        status = "abandoned"
    elif states and all(state in ("rejected", "invalidated") for state in states):
        status = states[0]
    else:
        status = "submitting"
    errors = [route["error"] for route in routes if route["error"]]
    ids = [row["provider_request_id"] for row in db().execute(
        "SELECT provider_request_id FROM outscraper_job_routes WHERE job_id=? AND provider_request_id IS NOT NULL",
        (job_id,)).fetchall()]
    count = sum(route["result_count"] for route in routes)
    completed = _timestamp_now() if status in ("completed", "failed", "unknown") else None
    db().execute(
        "UPDATE outscraper_jobs SET execution_status=?, provider_ids=?, result_count=?, error=?, completed_at=COALESCE(?,completed_at), updated_at=? WHERE id=?",
        (status, json.dumps(ids), count, "; ".join(errors) or None, completed, _timestamp_now(), job_id))
    db().commit()
    db().execute(
        "UPDATE volume_freeflow_attempts SET status=?,error=? WHERE job_id=?",
        (status, "; ".join(errors) or None, job_id),
    )
    db().commit()
    if status == "completed":
        _persist_holiday_result_context(job_id)
        workflow = db().execute(
            "SELECT 1 FROM request_directions_workflows WHERE job_id=?", (job_id,)
        ).fetchone()
        if workflow:
            try:
                _extract_request_directions_result(job_id)
            except (KeyError, TypeError, ValueError, OverflowError, sqlite3.Error) as exc:
                _update_request_directions_workflow(job_id, "failed", str(exc)[:1000])
    elif status in ("failed", "unknown"):
        _update_request_directions_workflow(job_id, status, "Scraper execution ended with status: " + status)
    return status

def _sync_volume_actual_speeds(request_id):
    workflow = db().execute(
        "SELECT actual_result_json FROM volumes_workflows WHERE request_id=?",
        (request_id,),
    ).fetchone()
    request_row = db().execute(
        "SELECT services,hours,data FROM requests WHERE id=?", (request_id,)
    ).fetchone()
    if not workflow or not request_row:
        return False

    actual_result = json.loads(workflow["actual_result_json"] or "{}")
    actual_hourly = {
        int(item["hour"]): item
        for item in actual_result.get("hourly", [])
        if isinstance(item, dict) and str(item.get("hour", "")).isdigit()
    }
    directions_row = db().execute(
        "SELECT result_json FROM request_directions_workflows "
        "WHERE request_id=? AND status IN ('ready','confirmed') AND result_json IS NOT NULL "
        "ORDER BY updated_at DESC,created_at DESC LIMIT 1",
        (request_id,),
    ).fetchone()
    directions_result = json.loads(directions_row["result_json"] or "{}") if directions_row else {}
    changed = False
    for item in directions_result.get("hourly", []):
        if not isinstance(item, dict) or not str(item.get("hour", "")).isdigit():
            continue
        hour = int(item["hour"])
        target = actual_hourly.setdefault(hour, {"hour": hour})
        for key in ("speed_kmh", "distance_m", "time_seconds", "travel_minutes", "reused_observation"):
            if item.get(key) is not None and target.get(key) != item[key]:
                target[key] = item[key]
                changed = True

    services = set(json.loads(request_row["services"] or "[]"))
    request_data = json.loads(request_row["data"] or "{}")
    if "speed" in services:
        for item in (request_data.get("speed") or {}).get("rows", []):
            if not isinstance(item, dict) or item.get("hour") is None or item.get("value") is None:
                continue
            hour = int(item["hour"])
            target = actual_hourly.setdefault(hour, {"hour": hour})
            if target.get("speed_kmh") != item["value"]:
                target["speed_kmh"] = item["value"]
                changed = True

    requested_hours = sorted(set(json.loads(request_row["hours"] or "[]")))
    actual_result["hourly"] = [
        actual_hourly[hour] for hour in sorted(actual_hourly)
    ]
    complete = bool(requested_hours) and all(
        hour in actual_hourly and actual_hourly[hour].get("speed_kmh") is not None
        for hour in requested_hours
    )
    if (changed or not workflow["actual_result_json"]) and actual_result["hourly"]:
        db().execute(
            "UPDATE volumes_workflows SET actual_result_json=?,updated_at=? WHERE request_id=?",
            (json.dumps(actual_result, ensure_ascii=False), _timestamp_now(), request_id),
        )
    return complete


def _evaluate_volume_freeflow(request_id):
    _sync_volume_actual_speeds(request_id)
    workflow = db().execute(
        "SELECT * FROM volumes_workflows WHERE request_id=?", (request_id,)
    ).fetchone()
    request_row = db().execute(
        "SELECT services,hours,data,user_id FROM requests WHERE id=?", (request_id,)
    ).fetchone()
    if not workflow or not request_row or not workflow["freeflow_result_json"]:
        return False
    source_result = json.loads(workflow["freeflow_result_json"])
    candidate_speed = source_result.get("selected_speed_kmh")
    actual_result = json.loads(workflow["actual_result_json"] or "{}")
    request_data = json.loads(request_row["data"] or "{}")
    requested_services = set(json.loads(request_row["services"] or "[]"))
    actual_by_hour = {
        item.get("hour"): item for item in actual_result.get("hourly", [])
        if isinstance(item, dict)
    }
    requested_hours = sorted(set(json.loads(request_row["hours"] or "[]")))
    comparison = []
    missing = []
    for hour in requested_hours:
        actual_item = actual_by_hour.get(hour)
        actual_speed = actual_item.get("speed_kmh") if actual_item else None
        try:
            valid = (
                math.isfinite(float(candidate_speed))
                and math.isfinite(float(actual_speed))
                and float(candidate_speed) > float(actual_speed)
            )
        except (TypeError, ValueError):
            valid = False
        if actual_speed is None:
            missing.append(f"{hour:02d}:00")
        comparison.append({
            "hour": hour,
            "actual_speed_kmh": actual_speed,
            "freeflow_speed_kmh": candidate_speed,
            "valid": valid if actual_speed is not None else None,
        })
    validation = {
        "rows": comparison,
        "missing_hours": missing,
        "mismatched_hours": [
            f"{item['hour']:02d}:00" for item in comparison
            if item["valid"] is False
        ],
        "is_valid": bool(comparison) and not missing and all(
            item["valid"] is True for item in comparison
        ),
    }
    if validation["missing_hours"]:
        validation_error = "Waiting for actual speed for selected hour(s): " + ", ".join(validation["missing_hours"]) + "."
    elif validation["mismatched_hours"]:
        validation_error = "Needs re-send: actual speed exceeds free-flow at " + ", ".join(validation["mismatched_hours"]) + "."
    else:
        validation_error = None
    db().execute(
        "UPDATE volumes_workflows SET freeflow_speed=?, freeflow_validation_json=?, "
        "freeflow_accepted=?, freeflow_override=0, freeflow_accepted_at=?, error=?, updated_at=? "
        "WHERE request_id=?",
        (
            candidate_speed,
            json.dumps(validation, ensure_ascii=False),
            int(validation["is_valid"]),
            _timestamp_now() if validation["is_valid"] else None,
            validation_error,
            _timestamp_now(),
            request_id,
        ),
    )
    attempt_status = "accepted" if validation["is_valid"] else (
        "waiting_for_actual_speed" if validation["missing_hours"] else "needs_resend"
    )
    db().execute(
        "UPDATE volume_freeflow_attempts SET validation_json=?,status=?,accepted=?,"
        "override=0,accepted_at=?,error=? WHERE job_id=?",
        (
            json.dumps(validation, ensure_ascii=False),
            attempt_status,
            int(validation["is_valid"]),
            _timestamp_now() if validation["is_valid"] else None,
            validation_error,
            workflow["freeflow_job_id"],
        ),
    )
    volume_data = request_data.get("volumes") or {}
    volume_data["freeflow_context"] = source_result.get("context")
    request_data["volumes"] = volume_data
    if validation["is_valid"] and "freeflow_speed" in requested_services:
        request_data["freeflow_speed"] = {
            "rows": [{"value": candidate_speed}],
            "note": "Selected as the highest speed returned for the reviewed A→C weekly-weekend window.",
            "source": "outscraper",
            "updated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
            "outscraper_job_id": workflow["freeflow_job_id"],
            "reference_context": source_result.get("context"),
        }
        _create_notification(
            request_row["user_id"], request_id, "directions_result",
            "Free-flow speed results are ready",
            f"The administrator approved free-flow speed for request REQ-{request_id}. Traffic volumes remain under review.",
        )
    status = "ready_for_review" if workflow["status"] == "ready_for_review" else (
        "sent" if workflow["status"] == "sent" else "waiting_for_inputs"
    )
    db().execute(
        "UPDATE requests SET data=?, status=CASE WHEN status NOT IN ('Rejected','Completed') "
        "THEN 'Processing' ELSE status END WHERE id=?",
        (json.dumps(request_data, ensure_ascii=False), request_id),
    )
    db().execute(
        "UPDATE volumes_workflows SET status=?, error=?, updated_at=? WHERE request_id=?",
        (status, validation_error, _timestamp_now(), request_id),
    )
    db().commit()
    return validation["is_valid"]

def _persist_holiday_result_context(job_id):
    job_row = db().execute("SELECT config_json FROM outscraper_jobs WHERE id=?", (job_id,)).fetchone()
    if not job_row:
        return
    context = json.loads(job_row["config_json"]).get("trafficlens_context")
    if not context:
        return
    request_id, service = context["source_request_id"], context["service"]
    request_row = db().execute("SELECT * FROM requests WHERE id=?", (request_id,)).fetchone()
    if not request_row:
        app.logger.warning("Completed reference-date Directions job %s has no matching source request.", job_id)
        return
    requested_services = set(json.loads(request_row["services"] or "[]"))
    if (
        service == "freeflow_speed"
        and "volumes" in requested_services
        and context.get("volume_input")
    ):
        try:
            route_row = db().execute(
                "SELECT raw_response_json FROM outscraper_job_routes WHERE job_id=? ORDER BY route_index LIMIT 1",
                (job_id,),
            ).fetchone()
            response = json.loads(route_row["raw_response_json"] or "{}") if route_row else {}
            timezone = ZoneInfo(context["timezone"])
            departure = int(context["departure_time"])
            finish = int(context["finish_time"])
            distance_meters = float(context["road_ab_meters"])
            values_by_hour = {}
            for record in _flatten_provider_records(response.get("data")):
                try:
                    timestamp = int(record["timestamp"])
                    if not departure <= timestamp <= finish:
                        continue
                    seconds = _interpolate_route_seconds(
                        record.get("road_distance_timing(meters:seconds)"), distance_meters
                    )
                    if seconds <= 0:
                        continue
                    hour = dt.datetime.fromtimestamp(timestamp, timezone).hour
                    values_by_hour.setdefault(hour, []).append(distance_meters / seconds * 3.6)
                except (KeyError, TypeError, ValueError, OverflowError):
                    continue
            if not values_by_hour:
                raise ValueError("The completed A→C response contains no usable time-stamped free-flow observations.")
            hourly = [
                {
                    "hour": hour,
                    "speed_kmh": max(speeds),
                    "sample_count": len(speeds),
                }
                for hour, speeds in sorted(values_by_hour.items())
            ]
            selected_speed = max(item["speed_kmh"] for item in hourly)
            result = {
                "hourly": hourly,
                "selected_speed_kmh": selected_speed,
                "selection_rule": "highest_speed_in_reviewed_window",
                "context": context,
                "outscraper_job_id": job_id,
            }
        except (KeyError, TypeError, ValueError, OverflowError, sqlite3.Error, ZoneInfoNotFoundError) as exc:
            db().execute(
                "UPDATE volumes_workflows SET freeflow_result_json=NULL, freeflow_accepted=0, "
                "status='waiting_for_inputs', error=?, updated_at=? WHERE request_id=?",
                (str(exc)[:1000], _timestamp_now(), request_id),
            )
            db().execute(
                "UPDATE volume_freeflow_attempts SET status='failed',error=? WHERE job_id=?",
                (str(exc)[:1000], job_id),
            )
            db().commit()
            app.logger.error("Could not extract A-to-C weekly-weekend free-flow speed for volume request %s: %s", request_id, exc)
            return
        request_data = json.loads(request_row["data"] or "{}")
        volume_data = request_data.get("volumes") or {}
        volume_data["freeflow_context"] = {
            key: context[key]
            for key in (
                "study_date", "country_code", "country_name", "timezone",
                "reference_type", "weekend_date", "weekend_day", "departure_time",
                "finish_time", "interval_minutes",
            )
            if key in context
        }
        request_data["volumes"] = volume_data
        db().execute(
            "UPDATE volumes_workflows SET freeflow_speed=?, freeflow_job_id=?, "
            "freeflow_result_json=?, freeflow_accepted=0, freeflow_validation_json=NULL, "
            "freeflow_override=0, freeflow_accepted_at=NULL, status='waiting_for_inputs', "
            "error=NULL, updated_at=? WHERE request_id=?",
            (
                selected_speed,
                job_id,
                json.dumps(result, ensure_ascii=False),
                _timestamp_now(),
                request_id,
            ),
        )
        db().execute(
            "UPDATE volume_freeflow_attempts SET status='waiting_for_actual_speed',result_json=?,error=NULL "
            "WHERE job_id=?",
            (json.dumps(result, ensure_ascii=False), job_id),
        )
        db().execute(
            "UPDATE requests SET data=? WHERE id=?",
            (json.dumps(request_data, ensure_ascii=False), request_id),
        )
        db().commit()
        if _evaluate_volume_freeflow(request_id):
            _start_volume_calculation(request_id)
        return
    if service == "freeflow_speed" and "volumes" in requested_services and service not in requested_services:
        try:
            route_row = db().execute(
                "SELECT raw_response_json FROM outscraper_job_routes WHERE job_id=? ORDER BY route_index LIMIT 1",
                (job_id,),
            ).fetchone()
            response = json.loads(route_row["raw_response_json"] or "{}") if route_row else {}
            target_meters = float(request_row["road_km"]) * 1000
            seconds = []
            for record in _flatten_provider_records(response.get("data")):
                seconds.append(_interpolate_route_seconds(
                    record.get("road_distance_timing(meters:seconds)"), target_meters
                ))
            if not seconds or any(value <= 0 for value in seconds):
                raise ValueError("The completed Directions response has no usable A-to-B timing profile.")
            freeflow_speed = target_meters / (sum(seconds) / len(seconds)) * 3.6
        except (KeyError, TypeError, ValueError, OverflowError, sqlite3.Error) as exc:
            db().execute(
                "UPDATE volumes_workflows SET status='failed', error=?, updated_at=? WHERE request_id=?",
                (str(exc)[:1000], _timestamp_now(), request_id),
            )
            db().commit()
            app.logger.error("Could not extract weekly-weekend free-flow speed for volume request %s: %s", request_id, exc)
            return
        db().execute(
            "UPDATE volumes_workflows SET freeflow_speed=?, freeflow_job_id=?, freeflow_accepted=1, "
            "freeflow_accepted_at=?, status='waiting_for_inputs', "
            "error=NULL, updated_at=? WHERE request_id=?",
            (freeflow_speed, job_id, _timestamp_now(), _timestamp_now(), request_id),
        )
        context = json.loads(job_row["config_json"]).get("trafficlens_context") or {}
        request_data = json.loads(request_row["data"] or "{}")
        volumes_data = request_data.get("volumes") or {}
        volumes_data["freeflow_context"] = {
            key: context[key]
            for key in (
                "study_date", "country_code", "country_name", "timezone",
                "reference_type", "weekend_date", "weekend_day", "departure_time",
            )
            if key in context
        }
        request_data["volumes"] = volumes_data
        db().execute(
            "UPDATE requests SET data=? WHERE id=?",
            (json.dumps(request_data, ensure_ascii=False), request_id),
        )
        db().commit()
        _start_volume_calculation(request_id)
        return
    if service not in requested_services:
        app.logger.warning("Completed reference-date Directions job %s has no matching source service.", job_id)
        return
    data = json.loads(request_row["data"] or "{}")
    saved = data.get(service) or {"rows": [], "note": "", "source": "outscraper"}
    if saved.get("source") == "network":
        saved["rows"] = []
        saved["note"] = ""
    if not saved.get("rows"):
        saved["note"] = "Directions API response is complete and awaiting administrator review of the service value."
        saved["source"] = "outscraper"
        if service == "freeflow_time":
            suggested = _suggest_freeflow_minutes(job_id, request_row["road_km"])
            if suggested is not None:
                saved["suggested_minutes"] = suggested
                saved["note"] += f" Suggested A-to-B free flow time: {suggested:.2f} min (average of the holiday Directions response). Confirm or correct it."
    reference_context = {
        key: context[key] for key in (
            "study_date", "country_code", "country_name", "timezone",
            "reference_type", "holiday_date", "holiday_name",
            "weekend_date", "weekend_day", "departure_time",
        ) if key in context
    }
    saved["holiday_context" if service == "freeflow_time" else "reference_context"] = reference_context
    saved["outscraper_job_id"] = job_id
    data[service] = saved
    db().execute(
        "UPDATE requests SET data=?, status=CASE WHEN status='Pending' THEN 'Processing' ELSE status END WHERE id=?",
        (json.dumps(data, ensure_ascii=False), request_id),
    )
    db().commit()

def _validate_trafficlens_request(configuration):
    context = configuration.get("trafficlens_request")
    if not context:
        return
    row = db().execute("SELECT * FROM requests WHERE id=?", (context["source_request_id"],)).fetchone()
    if not row:
        raise ValueError("The TrafficLens request associated with this Directions preview no longer exists.")
    source = to_dict(row)
    if not {"travel", "speed", "volumes"}.intersection(source["services"]) or not source.get("point_c") or not source.get("road_km"):
        raise ValueError("This request needs Travel time, Speed, or Volumes, a calculated Point C, and a saved A-to-B road distance.")
    expected_a = ", ".join(f"{float(value):.6f}" for value in source["coords"][0])
    expected_b = ", ".join(f"{float(value):.6f}" for value in source["coords"][-1])
    expected_c = ", ".join(f"{float(value):.6f}" for value in source["point_c"])
    coordinates = source["path"] or source["coords"]
    midpoint = coordinates[len(coordinates) // 2]
    country_code, country_name, timezone_name = _route_holiday_location(midpoint)
    route = configuration["routes"][0]
    if (
        context["point_a"] != expected_a
        or context["point_b"] != expected_b
        or context["point_c"] != expected_c
        or route["origin"] != expected_a
        or route["destination"] != expected_c
        or abs(context["road_ab_meters"] - float(source["road_km"]) * 1000) > 1
        or configuration["interval"] != context["interval_minutes"]
        or context["country_code"] != country_code
        or context["country_name"] != country_name
        or context["timezone"] != timezone_name
        or configuration["region"] != country_code
        or configuration["departure_time"] is None
        or configuration["finish_time"] is None
    ):
        raise ValueError("The A-to-C route, A-to-B distance, route country, time zone, or interval no longer matches the request.")

def _flatten_provider_records(value):
    if isinstance(value, list):
        for item in value:
            yield from _flatten_provider_records(item)
    elif isinstance(value, dict):
        yield value

def _interpolate_route_seconds(profile, target_meters):
    if not isinstance(profile, dict):
        raise ValueError("The Scraper response is missing road_distance_timing(meters:seconds).")
    points = []
    for distance, seconds in profile.items():
        try:
            distance_meters, elapsed_seconds = float(distance), float(seconds)
        except (TypeError, ValueError):
            continue
        if math.isfinite(distance_meters) and math.isfinite(elapsed_seconds) and distance_meters >= 0 and elapsed_seconds >= 0:
            points.append((distance_meters, elapsed_seconds))
    points.sort()
    if not points or points[0][0] > 0 or points[-1][0] < target_meters:
        raise ValueError("The road timing profile does not reach Point B on the A-to-C response.")
    for index, (distance, seconds) in enumerate(points):
        if math.isclose(distance, target_meters, abs_tol=0.01):
            return seconds
        if distance > target_meters:
            before_distance, before_seconds = points[index - 1]
            ratio = (target_meters - before_distance) / (distance - before_distance)
            return before_seconds + ratio * (seconds - before_seconds)
    raise ValueError("The road timing profile does not contain the A-to-B distance.")

def _mean_route_minutes(response, target_meters):
    """Average A-to-B travel time, in minutes, across every record of a Directions response."""
    seconds = []
    for record in _flatten_provider_records((response or {}).get("data")):
        try:
            value = _interpolate_route_seconds(
                record.get("road_distance_timing(meters:seconds)"), target_meters
            )
        except ValueError:
            continue
        if value > 0:
            seconds.append(value)
    return sum(seconds) / len(seconds) / 60 if seconds else None

def _suggest_freeflow_minutes(job_id, road_km):
    try:
        target_meters = float(road_km) * 1000
        route_row = db().execute(
            "SELECT raw_response_json FROM outscraper_job_routes WHERE job_id=? ORDER BY route_index LIMIT 1",
            (job_id,),
        ).fetchone()
        response = json.loads(route_row["raw_response_json"] or "{}") if route_row else {}
        return _mean_route_minutes(response, target_meters) if target_meters > 0 else None
    except (TypeError, ValueError, KeyError, sqlite3.Error):
        return None

def _extract_request_directions_result(job_id):
    workflow = db().execute(
        "SELECT request_id FROM request_directions_workflows WHERE job_id=?", (job_id,)
    ).fetchone()
    if not workflow:
        return
    route = db().execute(
        "SELECT raw_response_json, execution_status FROM outscraper_job_routes WHERE job_id=? ORDER BY route_index LIMIT 1",
        (job_id,),
    ).fetchone()
    if not route or route["execution_status"] != "completed" or not route["raw_response_json"]:
        raise ValueError("The Scraper response is not complete yet.")
    source_row = db().execute("SELECT * FROM requests WHERE id=?", (workflow["request_id"],)).fetchone()
    if not source_row:
        raise ValueError("The source request no longer exists.")
    source = to_dict(source_row)
    config_row = db().execute("SELECT config_json FROM outscraper_jobs WHERE id=?", (job_id,)).fetchone()
    config = json.loads(config_row["config_json"])
    context = config["trafficlens_request"]
    start_timestamp, end_timestamp = config["departure_time"], config["finish_time"]
    response = json.loads(route["raw_response_json"])
    records = list(_flatten_provider_records(response.get("data")))
    values_by_hour = {}
    timezone = ZoneInfo(context["timezone"])
    for record in records:
        try:
            timestamp = int(record["timestamp"])
            if timestamp < start_timestamp or timestamp > end_timestamp:
                continue
            elapsed_seconds = _interpolate_route_seconds(
                record.get("road_distance_timing(meters:seconds)"), context["road_ab_meters"]
            )
            hour = dt.datetime.fromtimestamp(timestamp, timezone).hour
        except (KeyError, TypeError, ValueError, OverflowError):
            continue
        values_by_hour.setdefault(hour, []).append(elapsed_seconds)
    requested_hours = sorted({
        int(hour) for hour in source["hours"] if str(hour).isdigit() and 0 <= int(hour) <= 23
    })
    if not requested_hours:
        raise ValueError("The source request has no selected study hours.")
    if not values_by_hour:
        raise ValueError("The completed Scraper response contains no usable timing data for the selected time range.")
    distance_meters = float(context["road_ab_meters"])
    all_seconds = [seconds for hour_values in values_by_hour.values() for seconds in hour_values]
    fallback_seconds = sum(all_seconds) / len(all_seconds)
    reused_hours = [hour for hour in requested_hours if not values_by_hour.get(hour)]
    hourly = []
    for hour in requested_hours:
        hour_values = values_by_hour.get(hour)
        seconds = sum(hour_values) / len(hour_values) if hour_values else fallback_seconds
        if seconds <= 0:
            raise ValueError(f"The Scraper returned an invalid travel time for {hour:02d}:00.")
        hourly.append({
            "hour": hour,
            "distance_m": distance_meters,
            "time_seconds": seconds,
            "travel_minutes": seconds / 60,
            "speed_kmh": distance_meters / seconds * 3.6,
            "reused_observation": not bool(hour_values),
        })
    mean_seconds = fallback_seconds
    result = {
        "route": "A → B",
        "point_a": context["point_a"],
        "point_b": context["point_b"],
        "point_c": context["point_c"],
        "distance_meters": distance_meters,
        "time_seconds": mean_seconds,
        "speed_kmh": distance_meters / mean_seconds * 3.6,
        "hourly": hourly,
        "reused_hours": reused_hours,
        "source_request_id": workflow["request_id"],
        "outscraper_job_id": job_id,
    }
    db().execute(
        "UPDATE request_directions_workflows SET status='ready', result_json=?, error=NULL, updated_at=? "
        "WHERE job_id=? AND status!='confirmed'",
        (json.dumps(result, ensure_ascii=False), _timestamp_now(), job_id),
    )
    db().commit()
    requested_services = set(source["services"])
    if "volumes" in requested_services:
        db().execute(
            "UPDATE volumes_workflows SET actual_result_json=?, status='waiting_for_inputs', "
            "error=NULL, updated_at=? WHERE request_id=? AND status NOT IN ('ready_for_review','sent')",
            (json.dumps(result, ensure_ascii=False), _timestamp_now(), workflow["request_id"]),
        )
        db().commit()
        if db().execute(
            "SELECT freeflow_result_json FROM volumes_workflows WHERE request_id=?",
            (workflow["request_id"],),
        ).fetchone()["freeflow_result_json"]:
            _evaluate_volume_freeflow(workflow["request_id"])
        _start_volume_calculation(workflow["request_id"])

def _volume_result_warnings(result):
    """Plain-language checks the admin sees before approving MATLAB output."""
    lanes = ((result.get("inputs") or {}).get("lanes")) or None
    warnings = []
    for item in result.get("hourly", []):
        label = f"{int(item['hour']):02d}:00"
        value, tolerance = item["value"], item.get("tolerance")
        if value < 0:
            warnings.append(f"{label}: the estimated volume is negative.")
        elif value == 0:
            warnings.append(f"{label}: the estimated volume is zero.")
        elif lanes and value / lanes > MAX_VOLUME_PER_LANE:
            warnings.append(
                f"{label}: {value / lanes:.0f} veh/h per lane is above {MAX_VOLUME_PER_LANE}, "
                "which is higher than a lane normally carries."
            )
        if tolerance is not None and value > 0 and tolerance > value:
            warnings.append(f"{label}: the RMSE tolerance is larger than the estimated volume.")
    return warnings

MAX_VOLUME_LANES = 12
MAX_VOLUME_PER_LANE = 2400  # veh/h/lane, above typical saturation flow

def _start_volume_calculation(request_id):
    _sync_volume_actual_speeds(request_id)
    workflow = db().execute(
        "SELECT * FROM volumes_workflows WHERE request_id=?", (request_id,)
    ).fetchone()
    if not workflow or workflow["status"] in {"running", "ready_for_review", "sent"}:
        return False
    missing_inputs = []
    if not workflow["actual_result_json"]:
        missing_inputs.append("completed actual Directions speeds")
    if workflow["lanes"] is None:
        missing_inputs.append("lane count")
    if workflow["freeflow_speed"] is None:
        missing_inputs.append("free-flow speed")
    elif not workflow["freeflow_accepted"]:
        missing_inputs.append("accepted free-flow speed")
    if missing_inputs:
        error = "Waiting for inputs: " + ", ".join(missing_inputs) + "."
        db().execute(
            "UPDATE volumes_workflows SET status='waiting_for_inputs', error=?, updated_at=? WHERE request_id=?",
            (error, _timestamp_now(), request_id),
        )
        db().commit()
        return False

    try:
        actual_result = json.loads(workflow["actual_result_json"] or "{}")
        actual_by_hour = {
            item.get("hour"): item for item in actual_result.get("hourly", [])
            if isinstance(item, dict)
        }
        request_row = db().execute(
            "SELECT hours,data,user_id FROM requests WHERE id=?", (request_id,)
        ).fetchone()
        if not request_row:
            raise ValueError("The source request no longer exists.")
        requested_hours = sorted(set(json.loads(request_row["hours"] or "[]")))
        rows = []
        missing_hours = []
        for hour in requested_hours:
            actual = actual_by_hour.get(hour)
            speed = actual.get("speed_kmh") if actual else None
            try:
                valid_speed = math.isfinite(float(speed)) and float(speed) > 0
            except (TypeError, ValueError):
                valid_speed = False
            if not valid_speed:
                missing_hours.append(f"{hour:02d}:00")
                continue
            rows.append({
                "hour": hour,
                "lanes": workflow["lanes"],
                "freeflow_speed": workflow["freeflow_speed"],
                "actual_speed": speed,
            })
        if missing_hours:
            error = "Waiting for actual speed for selected hour(s): " + ", ".join(missing_hours) + "."
            db().execute(
                "UPDATE volumes_workflows SET status='waiting_for_inputs', error=?, updated_at=? WHERE request_id=?",
                (error, _timestamp_now(), request_id),
            )
            db().commit()
            return False
    except (TypeError, ValueError, json.JSONDecodeError) as exc:
        error = f"Waiting for valid request inputs: {exc}"
        db().execute(
            "UPDATE volumes_workflows SET status='waiting_for_inputs', error=?, updated_at=? WHERE request_id=?",
            (error[:1000], _timestamp_now(), request_id),
        )
        db().commit()
        return False

    existing_result = volumes_runner.find_existing_volume_result(request_id, rows, request_row["user_id"])
    if existing_result:
        request_data = json.loads(request_row["data"] or "{}")
        existing_result["inputs"] = {
        "lanes": workflow["lanes"],
        "freeflow_speed": workflow["freeflow_speed"],
        "actual_speed_by_hour": [
            {"hour": hour, "speed_kmh": actual_by_hour[hour]["speed_kmh"]}
            for hour in requested_hours
        ],
        "freeflow_context": (request_data.get("volumes") or {}).get("freeflow_context"),
        }
        existing_result["warnings"] = _volume_result_warnings(existing_result)
        db().execute(
        "UPDATE volumes_workflows SET status='ready_for_review', result_json=?, "
        "workbook_path=?, error=NULL, updated_at=? "
        "WHERE request_id=? AND status IN ('waiting_for_inputs','failed')",
        (
            json.dumps(existing_result, ensure_ascii=False),
            existing_result["result_workbook"],
            _timestamp_now(),
            request_id,
        ),
        )
        db().commit()
        return True

    worker_key = f"volumes-{request_id}"
    with _outscraper_workers_lock:
        if worker_key in _outscraper_workers:
            return False
        _outscraper_workers.add(worker_key)

    try:
        prepared_run = volumes_runner.prepare_volume_run(request_id, rows, request_row["user_id"])
    except ValueError as exc:
        with _outscraper_workers_lock:
            _outscraper_workers.discard(worker_key)
        db().execute(
            "UPDATE volumes_workflows SET status='waiting_for_inputs', error=?, updated_at=? WHERE request_id=?",
            (str(exc)[:1000], _timestamp_now(), request_id),
        )
        db().commit()
        return False
    except Exception as exc:
        with _outscraper_workers_lock:
            _outscraper_workers.discard(worker_key)
        app.logger.exception("Could not save MATLAB inputs for volume request %s", request_id)
        db().execute(
            "UPDATE volumes_workflows SET status='failed', error=?, updated_at=? WHERE request_id=?",
            (f"Could not complete and verify Excel input preparation; MATLAB was not started. {str(exc)}"[:1000], _timestamp_now(), request_id),
        )
        db().commit()
        return False

    db().execute(
        "UPDATE volumes_workflows SET status='running', error=NULL, result_json=NULL, workbook_path=?, updated_at=? WHERE request_id=?",
        (prepared_run["input_workbook"], _timestamp_now(), request_id),
    )
    db().commit()
    try:
        threading.Thread(
            target=_run_volume_calculation,
            args=(request_id, prepared_run),
            name=worker_key,
            daemon=True,
        ).start()
    except RuntimeError as exc:
        with _outscraper_workers_lock:
            _outscraper_workers.discard(worker_key)
        db().execute(
            "UPDATE volumes_workflows SET status='failed', error=?, updated_at=? WHERE request_id=?",
            ("Excel inputs were saved and verified, but the local MATLAB worker could not be started.", _timestamp_now(), request_id),
        )
        db().commit()
        app.logger.error("Could not start MATLAB volume worker for request %s: %s", request_id, exc)
        return False
    return True

def _run_volume_calculation(request_id, prepared_run):
    worker_key = f"volumes-{request_id}"
    try:
        with app.app_context():
            workflow = db().execute(
                "SELECT * FROM volumes_workflows WHERE request_id=?", (request_id,)
            ).fetchone()
            request_row = db().execute(
                "SELECT * FROM requests WHERE id=?", (request_id,)
            ).fetchone()
            if not workflow or not request_row:
                raise ValueError("The volume workflow or source request no longer exists.")
            actual_result = json.loads(workflow["actual_result_json"] or "{}")
            actual_by_hour = {item["hour"]: item for item in actual_result.get("hourly", [])}
            result = volumes_runner.run_prepared_volume_estimation(prepared_run)
            request_data = json.loads(request_row["data"] or "{}")
            freeflow_context = (request_data.get("volumes") or {}).get("freeflow_context")
            result["inputs"] = {
                "lanes": workflow["lanes"],
                "freeflow_speed": workflow["freeflow_speed"],
                "actual_speed_by_hour": [
                    {"hour": hour, "speed_kmh": actual_by_hour[hour]["speed_kmh"]}
                    for hour in prepared_run["hours"]
                ],
                "freeflow_context": freeflow_context,
            }
            result["warnings"] = _volume_result_warnings(result)
            db().execute(
                "UPDATE volumes_workflows SET status='ready_for_review', result_json=?, workbook_path=?, "
                "error=NULL, updated_at=? WHERE request_id=? AND status='running'",
                (json.dumps(result, ensure_ascii=False), result["result_workbook"], _timestamp_now(), request_id),
            )
            db().commit()
    except Exception as exc:
        app.logger.exception("MATLAB volume calculation failed for request %s", request_id)
        with app.app_context():
            db().execute(
                "UPDATE volumes_workflows SET status='failed', error=?, updated_at=? "
                "WHERE request_id=? AND status='running'",
                (f"Excel inputs were saved and verified, but the MATLAB calculation did not complete. {str(exc)}"[:1000], _timestamp_now(), request_id),
            )
            db().commit()
    finally:
        with _outscraper_workers_lock:
            _outscraper_workers.discard(worker_key)

def _update_request_directions_workflow(job_id, status, error=None):
    db().execute(
        "UPDATE request_directions_workflows SET status=?, error=?, updated_at=? "
        "WHERE job_id=? AND status NOT IN ('ready','confirmed')",
        (status, error, _timestamp_now(), job_id),
    )
    db().commit()

def _poll_outscraper_route(route):
    db().execute("UPDATE outscraper_job_routes SET last_polled_at=? WHERE id=?",
                 (_timestamp_now(), route["id"]))
    db().commit()
    try:
        response = outscraper.fetch_request(route["provider_request_id"])
    except outscraper.OutscraperError as exc:
        if exc.retryable:
            db().execute("UPDATE outscraper_job_routes SET error=? WHERE id=?", (str(exc), route["id"]))
            db().commit()
            return "running"
        db().execute(
            "UPDATE outscraper_job_routes SET execution_status='failed', completed_at=?, error=? WHERE id=? AND execution_status!='abandoned'",
            (_timestamp_now(), str(exc), route["id"]))
        db().commit()
        return "failed"
    except ValueError:
        db().execute("UPDATE outscraper_job_routes SET error=? WHERE id=?",
                     ("Outscraper returned an invalid request ID.", route["id"]))
        db().commit()
        return "unknown"
    state = outscraper.provider_state(response)
    finished = _timestamp_now() if state in ("completed", "failed") else None
    safe_response = outscraper.redact_provider_data(response)
    provider_error = safe_response.get("errorMessage")
    if state == "failed" and not isinstance(provider_error, str):
        provider_error = "Outscraper marked the request as failed."
    db().execute(
        "UPDATE outscraper_job_routes SET execution_status=?, provider_status=?, raw_response_json=?, result_count=?, completed_at=?, error=? WHERE id=? AND execution_status!='abandoned'",
        (state, str(response.get("status") or "")[:80], json.dumps(safe_response, ensure_ascii=False),
         outscraper.result_count(response), finished,
         provider_error[:500] if state == "failed" else
         "Provider status is unclear; do not resubmit this route." if state == "unknown" else None,
         route["id"]))
    db().commit()
    return state

def _monitor_outscraper_job(job_id):
    try:
        for _ in range(16):
            time.sleep(15)
            with app.app_context():
                routes = db().execute(
                    "SELECT * FROM outscraper_job_routes WHERE job_id=? AND execution_status IN ('running','unknown') AND provider_request_id IS NOT NULL",
                    (job_id,)).fetchall()
                for row in routes:
                    last = row["last_polled_at"]
                    if last:
                        try:
                            elapsed = (dt.datetime.now(dt.timezone.utc) - dt.datetime.fromisoformat(last)).total_seconds()
                        except ValueError:
                            elapsed = 15
                        if elapsed < 15:
                            continue
                    _poll_outscraper_route(dict(row))
                status = _refresh_job_status(job_id)
                remaining = db().execute(
                    "SELECT 1 FROM outscraper_job_routes WHERE job_id=? AND execution_status IN ('running','unknown') AND provider_request_id IS NOT NULL LIMIT 1",
                    (job_id,)).fetchone()
                if not remaining or status in ("completed", "failed"):
                    return
    finally:
        with _outscraper_workers_lock:
            _outscraper_workers.discard(job_id)

def _start_outscraper_monitor(job_id):
    if not app.config["OUTSCRAPER_AUTO_POLL"] or not outscraper.api_key_configured():
        return
    with _outscraper_workers_lock:
        if job_id in _outscraper_workers:
            return
        _outscraper_workers.add(job_id)
    try:
        threading.Thread(
            target=_monitor_outscraper_job, args=(job_id,),
            name=f"outscraper-poll-{job_id}", daemon=True,
        ).start()
    except RuntimeError:
        with _outscraper_workers_lock:
            _outscraper_workers.discard(job_id)
        app.logger.error("Could not start the local Outscraper status monitor; manual refresh remains available.")

def _run_outscraper_job(job_id):
    with app.app_context():
        job = _outscraper_job(job_id)
        if not job: return
        configuration = job["config"]
        for route in job["routes"]:
            if route["execution_status"] != "queued":
                continue
            db().execute(
                "UPDATE outscraper_job_routes SET execution_status='submitting', submitted_at=?, request_details_json=? WHERE id=? AND execution_status='queued'",
                (_timestamp_now(), json.dumps(outscraper.directions_request_details(route, configuration), ensure_ascii=False),
                 route["id"]))
            db().execute("UPDATE outscraper_jobs SET updated_at=? WHERE id=?", (_timestamp_now(), job_id))
            db().commit()
            try:
                response = outscraper.submit_route(route, configuration)
            except outscraper.SubmissionUncertain as exc:
                db().execute("UPDATE outscraper_job_routes SET execution_status='unknown', error=? WHERE id=?",
                             (str(exc), route["id"]))
                db().commit()
                break
            except outscraper.OutscraperError as exc:
                db().execute("UPDATE outscraper_job_routes SET execution_status='failed', error=? WHERE id=?",
                             (str(exc), route["id"]))
                db().commit()
                break
            safe_response = outscraper.redact_provider_data(response)
            provider_id = response.get("id")
            if provider_id is not None and not outscraper.valid_request_id(provider_id):
                db().execute("UPDATE outscraper_job_routes SET execution_status='unknown', provider_status=?, raw_response_json=?, result_count=?, error=? WHERE id=?",
                             (str(response.get("status") or "")[:80], json.dumps(safe_response, ensure_ascii=False),
                              outscraper.result_count(response),
                              "Outscraper returned an invalid request ID; check provider status before retrying.", route["id"]))
                db().commit()
                break
            state = outscraper.provider_state(response)
            if state == "running" and not provider_id:
                state = "unknown"
            finished = _timestamp_now() if state in ("completed", "failed") else None
            provider_error = safe_response.get("errorMessage")
            if state == "failed" and not isinstance(provider_error, str):
                provider_error = "Outscraper marked the request as failed."
            db().execute(
                "UPDATE outscraper_job_routes SET execution_status=?, provider_request_id=?, provider_status=?, raw_response_json=?, result_count=?, completed_at=?, error=? WHERE id=?",
                (state, provider_id, str(response.get("status") or "")[:80],
                 json.dumps(safe_response, ensure_ascii=False), outscraper.result_count(response),
                 finished, provider_error[:500] if state == "failed" else None if state != "unknown" else
                 "Provider status is unclear; do not resubmit this route.", route["id"]))
            db().commit()
            if state in ("failed", "unknown"):
                break
        db().execute(
            "UPDATE outscraper_job_routes SET execution_status='not_submitted', error=? "
            "WHERE job_id=? AND execution_status='queued'",
            ("Not submitted after a preceding route failed or had an uncertain submission.", job_id))
        db().commit()
        status = _refresh_job_status(job_id)
        if status in ("running", "unknown") and app.config["OUTSCRAPER_AUTO_POLL"]:
            _monitor_outscraper_job(job_id)
def _process_outscraper_job(job_id):
    try:
        _run_outscraper_job(job_id)
    except (sqlite3.Error, ValueError, TypeError, KeyError):
        app.logger.error("Local Outscraper processing failed; marking uncertain routes without retrying.")
        with app.app_context():
            now = _timestamp_now()
            db().execute(
                "UPDATE outscraper_job_routes SET execution_status='unknown', error=? "
                "WHERE job_id=? AND execution_status IN ('queued','submitting')",
                ("Local processing failed while submission status may be uncertain. Do not resubmit.", job_id))
            db().execute(
                "UPDATE outscraper_jobs SET execution_status='unknown', error=?, completed_at=?, updated_at=? WHERE id=?",
                ("Local processing failed; check provider status before taking further action.", now, now, job_id))
            db().commit()
    finally:
        with _outscraper_workers_lock:
            _outscraper_workers.discard(job_id)

def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if "uid" not in session: return redirect(url_for("login", next=request.path))
        return f(*a, **k)
    return w

def admin_required(f):
    @wraps(f)
    @login_required
    def w(*a, **k):
        if session.get("role") != "admin": abort(403)
        return f(*a, **k)
    return w

def fmt_hours(hs):
    hs, out, i = sorted(set(hs)), [], 0
    while i < len(hs):
        j = i
        while j + 1 < len(hs) and hs[j + 1] == hs[j] + 1: j += 1
        out.append(f"{hs[i]:02d}:00-{hs[j] + 1:02d}:00"); i = j + 1
    return ", ".join(out) or "-"

def _clean_request_title(value):
    """Client-chosen request name: optional, single line, at most 80 characters."""
    if not isinstance(value, str):
        return None
    cleaned = " ".join("".join(ch for ch in value if ch.isprintable() or ch.isspace()).split())[:80]
    return cleaned or None

def to_dict(r):
    d = dict(r)
    for k in ("services", "bbox", "coords"): d[k] = json.loads(d[k])
    d["services"] = [s for s in d["services"] if s in service_catalog()]
    if len(d["coords"]) > 2: d["coords"] = [d["coords"][0], d["coords"][-1]]
    d["data"] = json.loads(d.get("data") or "{}")
    d["path"] = json.loads(d.get("path") or "null") or d["coords"]   # road shape; falls back to a straight line
    d["hours"] = json.loads(d.get("hours") or "[]")
    d["point_c"] = json.loads(d["point_c"]) if d.get("point_c") else None
    d["hours_label"] = fmt_hours(d["hours"])
    d["time_range_label"] = (
        f"{d['start_time']}–{d['end_time']}"
        if d.get("start_time") and d.get("end_time") else d["hours_label"]
    )
    d.pop("results", None)
    filled = [s for s in d["services"] if d["data"].get(s, {}).get("rows") or d["data"].get(s, {}).get("note")]
    d["progress"] = [len(filled), len(d["services"])]
    return d

def clean_rows(service, rows):
    form = forms_catalog().get(service)
    if not form: return []
    hourly = form["kind"] == "hourly"; out = []
    for row in (rows or [])[: 24 if hourly else 1]:
        if not isinstance(row, dict): continue
        if form.get("value_type") == "mixed":
            value = row.get("value")
            if value is None or not str(value).strip(): continue
            c = {"value": str(value).strip()[:200]}
        else:
            try: c = {"value": float(row.get("value"))}
            except (TypeError, ValueError): continue
        if hourly:
            try: h = int(row.get("hour"))
            except (TypeError, ValueError): continue
            if not 0 <= h <= 23: continue
            c["hour"] = h
        out.append(c)
    return out

def svinfo():
    over = {r["key"]: r["description"] for r in db().execute("SELECT key, description FROM service_info")}
    catalog, info = service_catalog(), {}
    for key, service in catalog.items():
        default = SERVICE_DEFAULTS.get(key, {"icon": "\U0001F9E9", "desc": service.get("description", "Data provided by the administrator.")})
        info[key] = {"icon": default["icon"], "desc": over.get(key) or default["desc"]}
    return info

def service_catalog():
    saved = get_setting("service_catalog", {}) or {}
    catalog = {}
    for key, name in SERVICES.items():
        config = saved.get(key, {})
        form = FORMS[key]
        catalog[key] = {
            "name": name, "unit": form["unit"], "kind": form["kind"],
            "value_type": "number", "custom": False,
            "enabled": config.get("enabled", key not in HIDDEN_SERVICES_BY_DEFAULT),
            "description": SERVICE_DEFAULTS[key]["desc"],
        }
    for key, config in saved.items():
        if not key.startswith("custom_") or not isinstance(config, dict): continue
        catalog[key] = {
            "name": str(config.get("name") or "Custom service"),
            "unit": str(config.get("unit") or ""),
            "kind": config.get("kind") if config.get("kind") in ("static", "hourly") else "static",
            "value_type": "mixed", "custom": True,
            "enabled": bool(config.get("enabled", True)),
            "description": str(config.get("description") or "Data entered by the administrator."),
        }
    return catalog

def active_services():
    return {key: service["name"] for key, service in service_catalog().items() if service["enabled"]}

def all_service_names():
    return {key: service["name"] for key, service in service_catalog().items()}

def forms_catalog():
    forms = {key: dict(form) for key, form in FORMS.items()}
    for key, service in service_catalog().items():
        if service["custom"]:
            forms[key] = _f(service["kind"], service["name"], service["unit"], f"Enter {service['name']} data for this link.")
            forms[key]["value_type"] = "mixed"
        else:
            forms[key]["value_type"] = "number"
    return forms

def get_setting(k, default=None):
    r = db().execute("SELECT value FROM settings WHERE key=?", (k,)).fetchone()
    return json.loads(r["value"]) if r else default

def set_setting(k, v):
    db().execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (k, json.dumps(v)))
    db().commit()

def net_services():
    """Services the static road network can answer without a date/time-specific API request."""
    cfg = get_setting("network", {})
    if not (network.HAVE_GEO and cfg.get("file")): return []
    out = ["length"]
    if cfg.get("lanes"): out.append("lanes")
    if cfg.get("capacity"): out.append("capacity")
    return out

def apply_network(rid):
    """Fill the services the network can answer. Pending request: all answered -> Completed, some answered -> Processing."""
    cfg = get_setting("network", {})
    if not (network.HAVE_GEO and cfg.get("file")): return None
    row = db().execute("SELECT * FROM requests WHERE id=?", (rid,)).fetchone()
    d = to_dict(row)
    res = network.lookup(cfg, d["path"])
    if not res: return None
    data, filled, now = d["data"], [], dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    for s in d["services"]:
        if s in {"freeflow_speed", "freeflow_time"}:
            continue
        v, cur = res["values"].get(s), d["data"].get(s) or {}
        if v and not (cur.get("rows") and cur.get("source") != "network"):     # never overwrite what the admin entered
            data[s] = {"rows": [{"value": v["value"]}], "note": v["note"], "source": "network", "updated": now}
            filled.append(s)
    status, reply, reply_at = row["status"], row["reply"], row["reply_at"]
    if status == "Pending" and filled:
        answered = [s for s in d["services"] if data.get(s, {}).get("rows")]
        names = lambda L: ", ".join(service_catalog()[x]["name"] for x in L)
        if len(answered) == len(d["services"]):
            status, msg = "Completed", "Your results are ready. All the data you requested was taken from our road network."
        else:
            waiting = [s for s in d["services"] if s not in answered]
            status, msg = "Processing", f"The data for {names(answered)} is ready from our road network and you can view it now. We are preparing {names(waiting)} and will update this request when it is ready."
        if not reply: reply, reply_at = msg, now
    db().execute("UPDATE requests SET data=?, status=?, reply=?, reply_at=? WHERE id=?", (json.dumps(data), status, reply, reply_at, rid)); db().commit()
    return {"filled": filled, "coverage": res["coverage"], "status": status}

@app.context_processor
def inject():
    catalog = service_catalog()
    unread_notifications = 0
    latest_notification_id = 0
    if session.get("uid"):
        unread_notifications = db().execute(
            "SELECT COUNT(*) FROM notifications WHERE user_id=? AND read_at IS NULL",
            (session["uid"],),
        ).fetchone()[0]
        latest_notification_id = db().execute(
            "SELECT COALESCE(MAX(id),0) FROM notifications WHERE user_id=?",
            (session["uid"],),
        ).fetchone()[0]
    return {"SERVICES": {k: v["name"] for k, v in catalog.items()},
            "ACTIVE_SERVICES": active_services(), "SERVICE_CATALOG": catalog,
            "me": session, "APP_NAME": APP_NAME, "SVINFO": svinfo(), "NET_SERVICES": net_services(),
            "unread_notifications": unread_notifications,
            "latest_notification_id": latest_notification_id}

@app.before_request
def refresh_session():
    # keep name/role in sync with the database, so a promoted or demoted user is affected immediately
    if "uid" in session and not request.path.startswith("/static"):
        u = db().execute(
            "SELECT name, role, profile_image FROM users WHERE id=?",
            (session["uid"],),
        ).fetchone()
        if not u: session.clear()
        else:
            session["name"], session["role"] = u["name"], u["role"]
            session["profile_image"] = bool(u["profile_image"])

@app.route("/")
def index():
    if session.get("role") == "admin":
        return redirect(url_for("admin"))
    if session.get("uid"):
        return redirect(url_for("dashboard"))
    return render_template("index.html")

@app.get("/dashboard")
@login_required
def dashboard():
    if session.get("role") == "admin":
        return redirect(url_for("admin"))
    rows = db().execute(
        "SELECT * FROM requests WHERE user_id=? ORDER BY id DESC LIMIT 5",
        (session["uid"],),
    ).fetchall()
    request_counts = {
        row["status"]: row["count"]
        for row in db().execute(
            "SELECT status,COUNT(*) AS count FROM requests WHERE user_id=? GROUP BY status",
            (session["uid"],),
        ).fetchall()
    }
    all_rows = db().execute(
        "SELECT * FROM requests WHERE user_id=?",
        (session["uid"],),
    ).fetchall()
    results_available = sum(
        1 for row in all_rows
        if row["status"] == "Completed"
        or (row["status"] == "Processing" and to_dict(row)["progress"][0] > 0)
    )
    return render_template(
        "dashboard.html",
        reqs=[to_dict(row) for row in rows],
        total_requests=db().execute(
            "SELECT COUNT(*) FROM requests WHERE user_id=?", (session["uid"],)
        ).fetchone()[0],
        request_counts=request_counts,
        results_available=results_available,
        active_services=active_services(),
        selected_service=request.args.get("service", ""),
    )

@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        n, e, p = (request.form.get(k, "").strip() for k in ("name", "email", "password"))
        if not (n and e and len(p) >= 6): flash("Fill all fields. Password needs at least 6 characters."); return render_template("auth.html", mode="register")
        try:
            db().execute("INSERT INTO users(name,email,pw) VALUES(?,?,?)", (n, e.lower(), generate_password_hash(p))); db().commit()
        except sqlite3.IntegrityError:
            flash("This email is already registered."); return render_template("auth.html", mode="register")
        flash("Account created. Sign in to continue."); return redirect(url_for("login"))
    return render_template("auth.html", mode="register")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = db().execute("SELECT * FROM users WHERE email=?", (request.form.get("email", "").strip().lower(),)).fetchone()
        if u and check_password_hash(u["pw"], request.form.get("password", "")):
            session.clear(); session.update(
                uid=u["id"], name=u["name"], role=u["role"],
                profile_image=bool(u["profile_image"]),
            )
            nxt = request.args.get("next")
            return redirect(nxt if nxt and nxt.startswith("/") else url_for("admin" if u["role"] == "admin" else "dashboard"))
        flash("Wrong email or password.")
    return render_template("auth.html", mode="login")

@app.route("/logout")
def logout(): session.clear(); return redirect(url_for("index"))

def _profile_redirect():
    return redirect(url_for("admin_profile" if session.get("role") == "admin" else "profile"))


def _profile_values(user):
    values = dict(user)
    if not values.get("first_name") or not values.get("last_name"):
        parts = (values.get("name") or "").strip().split(None, 1)
        values["first_name"] = values.get("first_name") or (parts[0] if parts else "")
        values["last_name"] = values.get("last_name") or (parts[1] if len(parts) > 1 else "")
    for field in (
        "phone_number", "job_title", "engineering_specialization",
        "company_name", "bio",
    ):
        values[field] = values.get(field) or ""
    return values


def _profile_photo_directory():
    return Path(app.instance_path) / "profile_images"


def _delete_profile_photo(filename, user_id):
    if not filename or not re.fullmatch(rf"user-{user_id}-[a-f0-9]{{32}}\.jpg", filename):
        return
    try:
        (_profile_photo_directory() / filename).unlink(missing_ok=True)
    except OSError:
        app.logger.warning(
            "Could not remove superseded profile photo for user %s", user_id,
            exc_info=True,
        )


def _process_profile_photo(upload):
    if not upload or not upload.filename:
        raise ValueError("Choose a JPG, PNG, or WebP image to upload.")
    try:
        raw = upload.stream.read(PROFILE_PHOTO_MAX_BYTES + 1)
    finally:
        upload.close()
    if len(raw) > PROFILE_PHOTO_MAX_BYTES:
        raise ValueError("Profile photos must be 5 MB or smaller.")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(io.BytesIO(raw)) as source:
                if source.format not in {"JPEG", "PNG", "WEBP"}:
                    raise ValueError("Use a JPG, PNG, or WebP image.")
                if source.width * source.height > 20_000_000:
                    raise ValueError("Choose an image with 20 megapixels or fewer.")
                source.verify()
            with Image.open(io.BytesIO(raw)) as source:
                image = ImageOps.exif_transpose(source).convert("RGB")
                image.thumbnail(PROFILE_PHOTO_SIZE, Image.Resampling.LANCZOS)
                output = io.BytesIO()
                image.save(output, format="JPEG", quality=88, optimize=True)
                return output.getvalue()
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
        raise ValueError("The selected image could not be opened. Choose a valid JPG, PNG, or WebP image.") from exc


def _render_profile(values, errors=None, status=200):
    return render_template(
        "profile.html",
        profile_user=values,
        profile_role="Admin" if session.get("role") == "admin" else "User",
        job_titles=PROFILE_JOB_TITLES,
        engineering_specializations=ENGINEERING_SPECIALIZATIONS,
        profile_errors=errors or {},
    ), status


@app.get("/profile")
@login_required
def profile():
    if session.get("role") == "admin":
        return redirect(url_for("admin_profile"))
    user = db().execute(
        "SELECT * FROM users WHERE id=? AND role='client'",
        (session["uid"],),
    ).fetchone()
    if not user:
        abort(404)
    return _render_profile(_profile_values(user))

@app.get("/admin/profile")
@admin_required
def admin_profile():
    user = db().execute(
        "SELECT * FROM users WHERE id=? AND role='admin'",
        (session["uid"],),
    ).fetchone()
    if not user:
        abort(404)
    return _render_profile(_profile_values(user))


@app.post("/profile/photo")
@login_required
def update_profile_photo():
    user_id = session["uid"]
    if (
        request.content_length is not None
        and request.content_length > PROFILE_PHOTO_MAX_BYTES + 128 * 1024
    ):
        flash("Profile photos must be 5 MB or smaller.", "error")
        return _profile_redirect()
    user = db().execute(
        "SELECT profile_image FROM users WHERE id=? AND role=?",
        (user_id, session.get("role")),
    ).fetchone()
    if not user:
        abort(404)

    if request.form.get("remove_photo") == "1":
        old_filename = user["profile_image"]
        db().execute("UPDATE users SET profile_image=NULL WHERE id=?", (user_id,))
        db().commit()
        session.pop("profile_image", None)
        _delete_profile_photo(old_filename, user_id)
        flash("Profile photo removed.", "success")
        return _profile_redirect()

    try:
        image_data = _process_profile_photo(request.files.get("photo"))
    except ValueError as exc:
        flash(str(exc), "error")
        return _profile_redirect()

    photo_dir = _profile_photo_directory()
    photo_dir.mkdir(parents=True, exist_ok=True)
    filename = f"user-{user_id}-{secrets.token_hex(16)}.jpg"
    final_path = photo_dir / filename
    temporary_path = photo_dir / f".{filename}.tmp"
    old_filename = user["profile_image"]
    try:
        temporary_path.write_bytes(image_data)
        os.replace(temporary_path, final_path)
        db().execute(
            "UPDATE users SET profile_image=? WHERE id=?",
            (filename, user_id),
        )
        db().commit()
    except OSError:
        temporary_path.unlink(missing_ok=True)
        final_path.unlink(missing_ok=True)
        app.logger.exception("Unable to save profile photo for user %s", user_id)
        flash("The profile photo could not be saved. Please try again.", "error")
        return _profile_redirect()
    except sqlite3.Error:
        db().rollback()
        temporary_path.unlink(missing_ok=True)
        final_path.unlink(missing_ok=True)
        app.logger.exception("Unable to update profile photo record for user %s", user_id)
        flash("The profile photo could not be saved. Please try again.", "error")
        return _profile_redirect()

    session["profile_image"] = True
    _delete_profile_photo(old_filename, user_id)
    flash("Profile photo updated.", "success")
    return _profile_redirect()


@app.get("/profile/photo")
@login_required
def profile_photo():
    user = db().execute(
        "SELECT profile_image FROM users WHERE id=? AND role=?",
        (session["uid"], session.get("role")),
    ).fetchone()
    if not user or not user["profile_image"]:
        abort(404)
    filename = user["profile_image"]
    if not re.fullmatch(rf"user-{session['uid']}-[a-f0-9]{{32}}\.jpg", filename):
        app.logger.error("Invalid profile image filename stored for user %s", session["uid"])
        abort(404)
    image_path = _profile_photo_directory() / filename
    if not image_path.is_file():
        abort(404)
    response = send_file(image_path, mimetype="image/jpeg", conditional=True, max_age=0)
    response.headers["Cache-Control"] = "private, no-store"
    response.headers["X-Content-Type-Options"] = "nosniff"
    return response


@app.post("/profile")
@login_required
def update_profile():
    first_name = request.form.get("first_name", "").strip()
    last_name = request.form.get("last_name", "").strip()
    if not first_name and not last_name and request.form.get("name"):
        name_parts = request.form.get("name", "").strip().split(None, 1)
        first_name = name_parts[0] if name_parts else ""
        last_name = name_parts[1] if len(name_parts) > 1 else ""
    email = request.form.get("email", "").strip().lower()
    phone_number = request.form.get("phone_number", "").strip()
    job_title = request.form.get("job_title", "").strip()
    specialization = request.form.get("engineering_specialization", "").strip()
    company_name = request.form.get("company_name", "").strip()
    bio = request.form.get("bio", "").strip()
    values = {
        "first_name": first_name,
        "last_name": last_name,
        "name": f"{first_name} {last_name}".strip(),
        "email": email,
        "phone_number": phone_number,
        "job_title": job_title,
        "engineering_specialization": specialization,
        "company_name": company_name,
        "bio": bio,
    }
    errors = {}
    if not first_name or len(first_name) > 60:
        errors["first_name"] = "Enter your first name (up to 60 characters)."
    if not last_name or len(last_name) > 60:
        errors["last_name"] = "Enter your last name (up to 60 characters)."
    if (
        len(email) > 254
        or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email)
    ):
        errors["email"] = "Enter a valid email address."
    phone_digits = "".join(char for char in phone_number if char.isdigit())
    if (
        not 7 <= len(phone_digits) <= 15
        or not re.fullmatch(r"\+?[0-9().\s-]+", phone_number)
    ):
        errors["phone_number"] = "Enter a valid international phone number (7–15 digits)."
    if job_title not in PROFILE_JOB_TITLES:
        errors["job_title"] = "Choose a job title from the list."
    if job_title == "Engineer":
        if specialization not in ENGINEERING_SPECIALIZATIONS:
            errors["engineering_specialization"] = "Choose an engineering specialization."
    else:
        values["engineering_specialization"] = ""
    if len(company_name) > 120:
        errors["company_name"] = "Company name must be 120 characters or fewer."
    if len(bio) > 500:
        errors["bio"] = "Description must be 500 characters or fewer."
    if errors:
        return _render_profile(values, errors, 400)
    try:
        db().execute(
            "UPDATE users SET name=?,first_name=?,last_name=?,email=?,phone_number=?,"
            "job_title=?,engineering_specialization=?,company_name=?,bio=? WHERE id=?",
            (
                values["name"], first_name, last_name, email, phone_number,
                job_title, values["engineering_specialization"], company_name,
                bio, session["uid"],
            ),
        )
        db().commit()
    except sqlite3.IntegrityError:
        db().rollback()
        errors["email"] = "This email address is already in use."
        return _render_profile(values, errors, 400)
    session["name"] = values["name"]
    flash("Profile updated.", "success")
    return _profile_redirect()

@app.post("/profile/password")
@login_required
def change_password():
    user = db().execute(
        "SELECT pw FROM users WHERE id=?", (session["uid"],)
    ).fetchone()
    current_password = request.form.get("current_password", "")
    new_password = request.form.get("new_password", "")
    confirm_password = request.form.get("confirm_password", "")
    if not user or not check_password_hash(user["pw"], current_password):
        flash("The current password is incorrect.", "error")
        return _profile_redirect()
    if len(new_password) < 6:
        flash("The new password must be at least 6 characters.", "error")
        return _profile_redirect()
    if new_password != confirm_password:
        flash("The new password and confirmation do not match.", "error")
        return _profile_redirect()
    db().execute(
        "UPDATE users SET pw=? WHERE id=?",
        (generate_password_hash(new_password), session["uid"]),
    )
    db().commit()
    flash("Password changed successfully.", "success")
    return _profile_redirect()

@app.route("/study")
@login_required
def study():
    repeat_request = None
    request_id = request.args.get("repeat_request_id", "").strip()
    if request_id:
        try:
            rid = int(request_id)
        except ValueError:
            abort(400, "Invalid request ID")
        row = db().execute(
            "SELECT * FROM requests WHERE id=? AND user_id=?",
            (rid, session["uid"]),
        ).fetchone()
        if not row:
            abort(404)
        original = to_dict(row)
        legacy_end_time = (
            f"{original['hours'][-1] + 1:02d}:00:00"
            if original["hours"] and original["hours"][-1] < 23
            else "23:59:59" if original["hours"] else ""
        )
        repeat_request = {
            "id": original["id"],
            "date": original["d1"],
            "coords": original["coords"],
            "services": original["services"],
            "start_time": original.get("start_time") or (f"{original['hours'][0]:02d}:00:00" if original["hours"] else ""),
            "end_time": original.get("end_time") or legacy_end_time,
            "period_label": original.get("period_label"),
        }
    selected_service = request.args.get("service", "")
    if selected_service not in active_services():
        selected_service = ""
    period_labels = [row["label"] for row in db().execute(
        "SELECT label FROM periods WHERE is_active=1 ORDER BY sort_order,id").fetchall()]
    return render_template(
        "study.html",
        repeat_request=repeat_request,
        selected_service=selected_service if not repeat_request else "",
        period_labels=period_labels,
        max_links=MAX_GROUP_LINKS,
    )

@app.get("/api/network/links")
@login_required
def search_network_links():
    cfg = get_setting("network", {})
    if not (network.HAVE_GEO and cfg.get("file") and os.path.isfile(cfg["file"])):
        return jsonify(error="The road network is not available"), 503
    query = request.args.get("q", "").strip()
    if not query: return jsonify(error="Enter a search term"), 400
    try:
        offset = int(request.args.get("offset", "0"))
        limit = int(request.args.get("limit", "20"))
    except ValueError: return jsonify(error="Invalid search page"), 400
    if offset < 0: return jsonify(error="Invalid search page"), 400
    limit = max(1, min(limit, 40))
    result = network.search_links(cfg["file"], query, offset, limit)
    return jsonify(**result, offset=offset, limit=limit)

@app.get("/api/network/links/<int:link_id>")
@login_required
def get_network_link(link_id):
    cfg = get_setting("network", {})
    if not (network.HAVE_GEO and cfg.get("file") and os.path.isfile(cfg["file"])):
        return jsonify(error="The road network is not available"), 503
    link = network.get_link(cfg["file"], link_id)
    if link is None: return jsonify(error="Link not found"), 404
    return jsonify(link)

@app.post("/api/requests")
@login_required
def create_request():
    j = request.get_json(force=True)
    enabled = active_services()
    sv = list(dict.fromkeys(x for x in j.get("services", []) if x in enabled))
    coords = j.get("coords") or []
    try: day = dt.date.fromisoformat(str(j.get("date") or "")).isoformat()
    except ValueError: return jsonify(error="Choose a valid date"), 400
    try:
        start_time = dt.time.fromisoformat(str(j.get("start_time") or ""))
        end_time = dt.time.fromisoformat(str(j.get("end_time") or ""))
    except ValueError:
        return jsonify(error="Choose a valid start and end time, including minutes and seconds"), 400
    if (start_time.tzinfo or end_time.tzinfo or start_time.microsecond or end_time.microsecond):
        return jsonify(error="Enter local times using hours, minutes and whole seconds only"), 400
    start_seconds = start_time.hour * 3600 + start_time.minute * 60 + start_time.second
    end_seconds = end_time.hour * 3600 + end_time.minute * 60 + end_time.second
    if end_seconds <= start_seconds:
        return jsonify(error="End time must be later than start time on the same study day"), 400
    hours = list(range(start_seconds // 3600, (end_seconds - 1) // 3600 + 1))
    start_time, end_time = start_time.strftime("%H:%M:%S"), end_time.strftime("%H:%M:%S")
    if len(coords) != 2: return jsonify(error="Select the start and the end of the link"), 400
    if not sv: return jsonify(error="Choose at least one service"), 400
    if not hours: return jsonify(error="Choose at least one hour"), 400
    lats, lngs = [c[0] for c in coords], [c[1] for c in coords]
    try:
        path = [[round(float(p[0]), 6), round(float(p[1]), 6)] for p in (j.get("path") or [])[:2000]]
        road_km = round(float(j["road_km"]), 2) if j.get("road_km") is not None else None
        junctions = int(j["junctions"]) if j.get("junctions") is not None else None
    except (TypeError, ValueError, IndexError): path, road_km, junctions = [], None, None
    if junctions is None or junctions < 0 or len(path) < 2:
        return jsonify(error="The road route could not be verified for intersections. Adjust points A and B and try again."), 400
    junction_quality_ack = junctions > 0 and j.get("accept_intersection_quality_impact") is True
    if junctions == 0 and j.get("no_junctions") is not True:
        return jsonify(error="Confirm that the link has no intersections, entrances or exits."), 400
    if junctions > 0 and not junction_quality_ack:
        return jsonify(
            error=(
                f"This link contains {junctions} possible intersection(s). "
                "Confirm that you agree to submit it knowing intersections may affect result quality."
            )
        ), 400
    point_c = None
    point_c_distance_km = None
    if {"speed", "travel", "volumes"}.intersection(sv):
        raw_point_c = j.get("point_c")
        try:
            if not isinstance(raw_point_c, list) or len(raw_point_c) != 2:
                raise ValueError
            point_c = [round(float(raw_point_c[0]), 6), round(float(raw_point_c[1]), 6)]
            if not (-90 <= point_c[0] <= 90 and -180 <= point_c[1] <= 180):
                raise ValueError
            point_c_distance_km = float(j.get("point_c_distance_km"))
            if not math.isfinite(point_c_distance_km) or point_c_distance_km <= 1.5:
                raise ValueError
        except (TypeError, ValueError):
            return jsonify(error="A valid road point C more than 1.5 km from A along the road is required for Time, Speed, or Volumes requests"), 400
    junction_quality_ack_at = _timestamp_now() if junction_quality_ack else None
    cur = db().execute("INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,type,hours,path,road_km,junctions,start_time,end_time,point_c,point_c_distance_km,title,junction_quality_ack,junction_quality_ack_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (session["uid"], dt.date.today().isoformat(), day, day, json.dumps(sv), str(j.get("area") or "Link")[:120],
         json.dumps([min(lats), min(lngs), max(lats), max(lngs)]), json.dumps(coords), "link", json.dumps(hours), json.dumps(path), road_km, junctions, start_time, end_time, json.dumps(point_c) if point_c else None, point_c_distance_km, _clean_request_title(j.get("title")), int(junction_quality_ack), junction_quality_ack_at))
    if "volumes" in sv:
        now = _timestamp_now()
        db().execute(
            "INSERT INTO volumes_workflows(request_id,status,created_at,updated_at) VALUES(?,'waiting_for_inputs',?,?)",
            (cur.lastrowid, now, now),
        )
    if session.get("role") != "admin":
        _notify_admins_new_request(cur.lastrowid, session.get("name", "A client"))
    db().commit()
    rid, info = cur.lastrowid, None
    try: info = apply_network(rid)
    except Exception as e: app.logger.warning("network lookup failed: %s", e)      # never block a client's request
    st = db().execute("SELECT status FROM requests WHERE id=?", (rid,)).fetchone()["status"]
    return jsonify(id=rid, status=st, filled=(info or {}).get("filled", []))


# ---------------------------------------------------------------------------
# Grouped orders: several links and/or several periods submitted as one order.
# Every (link x period) is still an ordinary row in `requests`, so the whole
# existing admin workflow (volumes, directions, free flow...) keeps working;
# `request_groups` ties the rows together for intake, review and one report.
# ---------------------------------------------------------------------------
def _parse_local_time(value):
    t = dt.time.fromisoformat(str(value or ""))
    if t.tzinfo or t.microsecond:
        raise ValueError
    return t

def _clean_group_period(raw, active_labels):
    if not isinstance(raw, dict):
        return None, "Invalid period."
    label = str(raw.get("label") or "").strip()
    if label not in active_labels:
        return None, f"Period '{label[:20]}' is not available."
    try:
        start, end = _parse_local_time(raw.get("start_time")), _parse_local_time(raw.get("end_time"))
    except ValueError:
        return None, f"Period {label}: choose a valid start and end time (whole seconds only)."
    start_s = start.hour * 3600 + start.minute * 60 + start.second
    end_s = end.hour * 3600 + end.minute * 60 + end.second
    if end_s <= start_s:
        return None, f"Period {label}: end time must be later than start time."
    return {"label": label, "start_time": start.strftime("%H:%M:%S"), "end_time": end.strftime("%H:%M:%S"),
            "hours": list(range(start_s // 3600, (end_s - 1) // 3600 + 1))}, None

def _clean_group_link(raw, services, number):
    name = f"Link {number}"
    if not isinstance(raw, dict):
        return None, f"{name}: invalid data."
    try:
        coords = raw.get("coords") or []
        if len(coords) != 2:
            return None, f"{name}: select the start and the end of the link."
        coords = [[round(float(c[0]), 6), round(float(c[1]), 6)] for c in coords]
        path = [[round(float(p[0]), 6), round(float(p[1]), 6)] for p in (raw.get("path") or [])[:2000]]
        road_km = round(float(raw["road_km"]), 2) if raw.get("road_km") is not None else None
        junctions = int(raw["junctions"]) if raw.get("junctions") is not None else None
    except (TypeError, ValueError, IndexError, KeyError):
        return None, f"{name}: the road route could not be verified. Adjust points A and B and add the link again."
    if junctions is None or junctions < 0 or len(path) < 2:
        return None, f"{name}: the road route could not be verified for intersections."
    ack = junctions > 0 and raw.get("accept_intersection_quality_impact") is True
    if junctions == 0 and raw.get("no_junctions") is not True:
        return None, f"{name}: confirm that the link has no intersections, entrances or exits."
    if junctions > 0 and not ack:
        return None, f"{name}: confirm that you accept the possible intersection quality impact."
    point_c = point_c_distance_km = None
    if {"speed", "travel", "volumes"}.intersection(services):
        try:
            rc = raw.get("point_c")
            if not isinstance(rc, list) or len(rc) != 2:
                raise ValueError
            point_c = [round(float(rc[0]), 6), round(float(rc[1]), 6)]
            if not (-90 <= point_c[0] <= 90 and -180 <= point_c[1] <= 180):
                raise ValueError
            point_c_distance_km = float(raw.get("point_c_distance_km"))
            if not math.isfinite(point_c_distance_km) or point_c_distance_km <= 1.5:
                raise ValueError
        except (TypeError, ValueError):
            return None, f"{name}: a valid road point C more than 1.5 km from A is required for Time, Speed or Volumes."
    lats, lngs = [c[0] for c in coords], [c[1] for c in coords]
    return {"coords": coords, "path": path, "road_km": road_km, "junctions": junctions, "ack": ack,
            "point_c": point_c, "point_c_distance_km": point_c_distance_km,
            "bbox": [min(lats), min(lngs), max(lats), max(lngs)]}, None

@app.post("/api/request-groups")
@login_required
def create_request_group():
    j = request.get_json(force=True, silent=True) or {}
    enabled = active_services()
    sv = list(dict.fromkeys(x for x in j.get("services", []) if x in enabled))
    if not sv:
        return jsonify(error="Choose at least one service"), 400
    try: day = dt.date.fromisoformat(str(j.get("date") or "")).isoformat()
    except ValueError: return jsonify(error="Choose a valid date"), 400
    mode = j.get("mode")
    if mode not in GROUP_MODES:
        return jsonify(error="Choose the order type first"), 400
    raw_links, raw_periods = j.get("links"), j.get("periods")
    if not isinstance(raw_links, list) or not isinstance(raw_periods, list):
        return jsonify(error="Add at least one link and one period"), 400
    if not raw_links:
        return jsonify(error="Add at least one link"), 400
    if not raw_periods:
        return jsonify(error="Choose at least one period"), 400
    if len(raw_links) > MAX_GROUP_LINKS:
        return jsonify(error=f"An order can contain at most {MAX_GROUP_LINKS} links"), 400
    if mode in ("single", "multi_period") and len(raw_links) != 1:
        return jsonify(error="This order type allows exactly one link"), 400
    if mode in ("single", "multi_link") and len(raw_periods) != 1:
        return jsonify(error="This order type allows exactly one period"), 400
    active_labels = {row["label"] for row in db().execute("SELECT label FROM periods WHERE is_active=1").fetchall()}
    periods, seen = [], set()
    for raw in raw_periods:
        period, error = _clean_group_period(raw, active_labels)
        if error: return jsonify(error=error), 400
        if period["label"] in seen: return jsonify(error=f"Period {period['label']} was chosen twice"), 400
        seen.add(period["label"]); periods.append(period)
    links = []
    for number, raw in enumerate(raw_links, start=1):
        link, error = _clean_group_link(raw, sv, number)
        if error: return jsonify(error=error), 400
        links.append(link)
    title = _clean_request_title(j.get("title"))
    now_day, total = dt.date.today().isoformat(), len(links) * len(periods)
    d = db()
    group_id = None
    try:
        if total > 1:
            group_id = d.execute(
                "INSERT INTO request_groups(user_id,title,created,study_date,services,mode) VALUES(?,?,?,?,?,?)",
                (session["uid"], title, now_day, day, json.dumps(sv), mode)).lastrowid
        request_ids = []
        for number, link in enumerate(links, start=1):
            for period in periods:
                child_title = title if total == 1 else _clean_request_title(
                    " · ".join(x for x in (title, f"Link {number}", period["label"]) if x))
                cur = d.execute(
                    "INSERT INTO requests(user_id,created,d1,d2,services,area,bbox,coords,type,hours,path,road_km,junctions,start_time,end_time,point_c,point_c_distance_km,title,junction_quality_ack,junction_quality_ack_at,group_id,period_label) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (session["uid"], now_day, day, day, json.dumps(sv),
                     f"Link {number} A to B (~{link['road_km']} km by road)" if total > 1 else f"Link A to B (~{link['road_km']} km by road)",
                     json.dumps(link["bbox"]), json.dumps(link["coords"]), "link", json.dumps(period["hours"]),
                     json.dumps(link["path"]), link["road_km"], link["junctions"], period["start_time"], period["end_time"],
                     json.dumps(link["point_c"]) if link["point_c"] else None, link["point_c_distance_km"],
                     child_title, int(link["ack"]), _timestamp_now() if link["ack"] else None, group_id, period["label"]))
                request_ids.append(cur.lastrowid)
                if "volumes" in sv:
                    now = _timestamp_now()
                    d.execute("INSERT INTO volumes_workflows(request_id,status,created_at,updated_at) VALUES(?,'waiting_for_inputs',?,?)",
                              (cur.lastrowid, now, now))
        if session.get("role") != "admin":
            if group_id:
                for admin_user in d.execute("SELECT id FROM users WHERE role='admin'").fetchall():
                    _create_notification(admin_user["id"], request_ids[0], "new_request", "New grouped request",
                        f"{session.get('name', 'A client')} submitted order GRP-{group_id}: {len(links)} link(s) x {len(periods)} period(s) = {total} requests.")
            else:
                _notify_admins_new_request(request_ids[0], session.get("name", "A client"))
        d.commit()
    except sqlite3.Error:
        d.rollback()
        app.logger.exception("Could not save the grouped order")
        return jsonify(error="The order could not be saved. Nothing was created; please try again."), 500
    for rid in request_ids:
        try: apply_network(rid)
        except Exception as e: app.logger.warning("network lookup failed: %s", e)
    status = d.execute("SELECT status FROM requests WHERE id=?", (request_ids[0],)).fetchone()["status"]
    return jsonify(group_id=group_id, request_ids=request_ids, count=total, id=request_ids[0], status=status)

def _group_children(gid):
    rows = db().execute("SELECT * FROM requests WHERE group_id=? ORDER BY id", (gid,)).fetchall()
    return [to_dict(row) for row in rows]

def _group_status(group_row, children):
    if group_row["sent_at"]:
        return "Completed"
    active = [c for c in children if c["status"] != "Rejected"]
    if not active:
        return "Rejected"
    if any(c["status"] in ("Processing", "Completed") for c in active):
        return "Processing"
    return "Pending"

def _group_ready_to_send(children):
    active = [c for c in children if c["status"] != "Rejected"]
    return bool(active) and all(c["status"] == "Completed" for c in active)

def _group_card(group_row):
    children = _group_children(group_row["id"])
    return {"id": group_row["id"], "title": group_row["title"], "created": group_row["created"],
            "d1": group_row["study_date"], "services": [x for x in json.loads(group_row["services"]) if x in service_catalog()],
            "count": len(children), "periods": sorted({c["period_label"] for c in children if c.get("period_label")}),
            "links": len({json.dumps(c["coords"]) for c in children}),
            "status": _group_status(group_row, children), "sent": bool(group_row["sent_at"])}

def _load_group(gid, admin_view):
    row = db().execute(
        "SELECT g.*, u.name AS client, u.email AS email FROM request_groups g JOIN users u ON u.id=g.user_id WHERE g.id=?",
        (gid,)).fetchone()
    if not row or (not admin_view and row["user_id"] != session["uid"]):
        abort(404)
    children = _group_children(gid)
    # label every child by its link and period, preserving the order links were added
    link_index = {}
    for c in children:
        key = json.dumps(c["coords"])
        link_index.setdefault(key, len(link_index) + 1)
        c["link_no"] = link_index[key]
    return row, children

@app.get("/groups/<int:gid>")
@login_required
def client_group(gid):
    if session.get("role") == "admin":
        return redirect(url_for("admin_group", gid=gid))
    row, children = _load_group(gid, admin_view=False)
    return render_template("group.html", g=row, children=children, card=_group_card(row), admin_view=False,
                           ready=_group_ready_to_send(children))

@app.get("/admin/groups/<int:gid>")
@admin_required
def admin_group(gid):
    row, children = _load_group(gid, admin_view=True)
    return render_template("group.html", g=row, children=children, card=_group_card(row), admin_view=True,
                           ready=_group_ready_to_send(children))

@app.post("/api/groups/<int:gid>/send")
@admin_required
def send_group_report(gid):
    row = db().execute("SELECT * FROM request_groups WHERE id=?", (gid,)).fetchone()
    if not row:
        return jsonify(error="Order not found"), 404
    if row["sent_at"]:
        return jsonify(error="This report was already sent."), 409
    children = _group_children(gid)
    if not _group_ready_to_send(children):
        return jsonify(error="Every request in the order must be Completed (or Rejected) before the report can be sent."), 409
    stamp = _timestamp_now()
    db().execute("UPDATE request_groups SET sent_at=? WHERE id=?", (stamp, gid))
    first = children[0]["id"]
    _create_notification(row["user_id"], first, "group_ready", "Your report is ready",
                         f"All results for order GRP-{gid} are ready in one report.")
    db().commit()
    return jsonify(ok=True, sent_at=stamp)

@app.get("/groups/<int:gid>/export.xlsx")
@login_required
def export_group_excel(gid):
    admin_view = session.get("role") == "admin"
    row, children = _load_group(gid, admin_view=admin_view)
    if not admin_view and not row["sent_at"]:
        abort(404)
    forms = forms_catalog()
    services = [x for x in json.loads(row["services"]) if x in service_catalog()]
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Order Summary"
    summary.append(["Field", "Value"])
    link_count = len({json.dumps(c["coords"]) for c in children})
    for label, value in (
        ("Order ID", f"GRP-{gid}"), ("Name", row["title"] or ""), ("Client", row["client"]),
        ("Status", _group_status(row, children)), ("Study date", row["study_date"]),
        ("Links", link_count), ("Periods", ", ".join(sorted({c["period_label"] for c in children if c.get("period_label")}))),
        ("Requests in this order", len(children)),
        ("Requested services", ", ".join(SERVICES.get(x, x) for x in services)),
    ):
        summary.append([_excel_safe_text(label), value if isinstance(value, (int, float)) else _excel_safe_text(value)])
    _style_excel_sheet(summary)

    links_sheet = workbook.create_sheet("Links and Periods")
    links_sheet.append(["Request", "Link", "Period", "Time range", "Status", "Point A", "Point B", "Point C", "Road length (km)"])
    for c in children:
        show_c = admin_view or c.get("show_point_c")
        links_sheet.append([
            f"REQ-{c['id']}", f"Link {c['link_no']}", _excel_safe_text(c.get("period_label") or ""),
            _excel_safe_text(c["time_range_label"]), c["status"],
            ", ".join(map(str, c["coords"][0])), ", ".join(map(str, c["coords"][-1])),
            ", ".join(map(str, c["point_c"])) if (c.get("point_c") and show_c) else "", c.get("road_km")])
    _style_excel_sheet(links_sheet)

    details = workbook.create_sheet("Service Details")
    details.append(["Link", "Period", "Time range", "Service", "Hour", "Value", "Tolerance", "Unit", "Source", "Notes"])
    for c in children:
        if c["status"] == "Rejected":
            continue
        for service in c["services"]:
            saved = c["data"].get(service, {})
            unit = (forms.get(service) or {}).get("unit", "")
            source = ("Road network" if saved.get("source") == "network" else "MATLAB" if saved.get("source") == "matlab"
                      else "Admin" if (saved.get("rows") or saved.get("note")) else "")
            rows = saved.get("rows") or [{}]
            for service_row in rows:
                hour = service_row.get("hour")
                hour_label = f"{int(hour):02d}:00-{int(hour) + 1:02d}:00" if hour is not None else ""
                value = service_row.get("value")
                if value is not None and not isinstance(value, (int, float)) or isinstance(value, bool):
                    value = _excel_safe_text(value)
                tol = service_row.get("tolerance")
                if tol is not None and not isinstance(tol, (int, float)):
                    tol = _excel_safe_text(tol)
                details.append([f"Link {c['link_no']}", _excel_safe_text(c.get("period_label") or ""), _excel_safe_text(c["time_range_label"]),
                                _excel_safe_text(SERVICES.get(service, service)), hour_label, value, tol,
                                _excel_safe_text(unit), _excel_safe_text(source), _excel_safe_text(saved.get("note") or "")])
    _style_excel_sheet(details)

    route_rows = []
    for c in children:
        directions = (c["data"].get("travel") or c["data"].get("speed") or {}).get("directions_result")
        if directions and c["status"] != "Rejected":
            route_rows.append([f"Link {c['link_no']}", _excel_safe_text(c.get("period_label") or ""),
                               directions.get("distance_meters"), directions.get("time_seconds"), directions.get("speed_kmh")])
    if route_rows:
        routes = workbook.create_sheet("A-to-B Results")
        routes.append(["Link", "Period", "A-to-B distance (m)", "A-to-B time (s)", "A-to-B speed (km/h)"])
        for item in route_rows:
            routes.append(item)
        _style_excel_sheet(routes)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return app.response_class(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="trafficlens-order-{gid}.xlsx"'},
    )

@app.post("/api/periods")
@admin_required
def add_period():
    label = " ".join(str((request.get_json(silent=True) or {}).get("label") or "").split())[:12].upper()
    if not label:
        return jsonify(error="Enter a period name, e.g. EVE"), 400
    if db().execute("SELECT 1 FROM periods WHERE label=?", (label,)).fetchone():
        return jsonify(error="This period already exists"), 409
    order = db().execute("SELECT COALESCE(MAX(sort_order),0)+1 FROM periods").fetchone()[0]
    db().execute("INSERT INTO periods(label,sort_order) VALUES(?,?)", (label, order)); db().commit()
    return jsonify(ok=True)

@app.post("/api/periods/<int:pid>/availability")
@admin_required
def set_period_availability(pid):
    enabled = (request.get_json(silent=True) or {}).get("enabled")
    if not isinstance(enabled, bool):
        return jsonify(error="Invalid"), 400
    if not enabled and db().execute("SELECT COUNT(*) FROM periods WHERE is_active=1 AND id!=?", (pid,)).fetchone()[0] == 0:
        return jsonify(error="At least one period must stay available"), 409
    if not db().execute("UPDATE periods SET is_active=? WHERE id=?", (int(enabled), pid)).rowcount:
        return jsonify(error="Period not found"), 404
    db().commit()
    return jsonify(ok=True)


@app.post("/api/requests/<int:rid>/title")
@login_required
def rename_request(rid):
    payload = request.get_json(silent=True) or {}
    if not isinstance(payload.get("title"), str):
        return jsonify(error="Enter a name for the request."), 400
    title = _clean_request_title(payload["title"])
    cursor = db().execute(
        "UPDATE requests SET title=? WHERE id=? AND user_id=?", (title, rid, session["uid"])
    )
    if cursor.rowcount != 1:
        return jsonify(error="Request not found"), 404
    db().commit()
    return jsonify(ok=True, title=title)

@app.route("/requests")
@login_required
def my_requests():
    rows = db().execute("SELECT * FROM requests WHERE user_id=? AND group_id IS NULL ORDER BY id DESC", (session["uid"],)).fetchall()
    reqs = [to_dict(row) for row in rows]
    for item in reqs:
        item["results_available"] = (
            item["status"] == "Completed"
            or (item["status"] == "Processing" and item["progress"][0] > 0)
        )
    groups = [_group_card(row) for row in db().execute(
        "SELECT * FROM request_groups WHERE user_id=? ORDER BY id DESC", (session["uid"],)).fetchall()]
    initial_status = request.args.get("status", "")
    if initial_status not in STATUSES:
        initial_status = ""
    initial_results = request.args.get("results", "")
    if initial_results != "available":
        initial_results = ""
    return render_template(
        "requests.html",
        reqs=reqs,
        groups=groups,
        statuses=STATUSES,
        initial_status=initial_status,
        initial_results=initial_results,
    )

@app.get("/requests/<int:rid>")
@login_required
def user_request_detail(rid):
    row = db().execute(
        "SELECT r.*, u.name AS client, u.email FROM requests r "
        "JOIN users u ON u.id=r.user_id WHERE r.id=? AND r.user_id=?",
        (rid, session["uid"]),
    ).fetchone()
    if not row:
        abort(404)
    item = to_dict(row)
    item["results_available"] = (
        item["status"] == "Completed"
        or (item["status"] == "Processing" and item["progress"][0] > 0)
    )
    return render_template("request_detail.html", r=item)

@app.get("/notifications")
@login_required
def notifications():
    rows = db().execute(
        "SELECT * FROM notifications WHERE user_id=? ORDER BY id DESC LIMIT 100",
        (session["uid"],),
    ).fetchall()
    return render_template("notifications.html", notifications=rows)

@app.get("/notifications/<int:notification_id>/open")
@login_required
def open_notification(notification_id):
    notification = db().execute(
        "SELECT * FROM notifications WHERE id=? AND user_id=?",
        (notification_id, session["uid"]),
    ).fetchone()
    if not notification:
        abort(404)
    if notification["read_at"] is None:
        db().execute(
            "UPDATE notifications SET read_at=? WHERE id=? AND user_id=? AND read_at IS NULL",
            (dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), notification_id, session["uid"]),
        )
        db().commit()
    return redirect(_notification_destination(notification["request_id"], session.get("role") == "admin"))

def _notification_destination(request_id, is_admin):
    row = db().execute("SELECT group_id FROM requests WHERE id=?", (request_id,)).fetchone()
    group_id = row["group_id"] if row else None
    if group_id:
        return url_for("admin_group", gid=group_id) if is_admin else url_for("client_group", gid=group_id)
    return url_for("admin_request", rid=request_id) if is_admin else url_for("result", rid=request_id)

@app.get("/api/notifications/unread-count")
@login_required
def notification_unread_count():
    count = db().execute(
        "SELECT COUNT(*) FROM notifications WHERE user_id=? AND read_at IS NULL",
        (session["uid"],),
    ).fetchone()[0]
    return jsonify(count=count)

@app.get("/api/notifications/device-feed")
@login_required
def notification_device_feed():
    try:
        after_id = int(request.args.get("after", "0"))
    except ValueError:
        return jsonify(error="Invalid notification cursor"), 400
    if after_id < 0:
        return jsonify(error="Invalid notification cursor"), 400
    rows = db().execute(
        "SELECT id,request_id,kind,title,message,created_at FROM notifications "
        "WHERE user_id=? AND id>? ORDER BY id ASC LIMIT 30",
        (session["uid"], after_id),
    ).fetchall()
    notifications_out = []
    for row in rows:
        destination = _notification_destination(row["request_id"], session.get("role") == "admin")
        notifications_out.append({
            "id": row["id"],
            "title": row["title"],
            "message": row["message"],
            "created_at": row["created_at"],
            "url": destination,
        })
    return jsonify(notifications=notifications_out)

@app.route("/admin")
@admin_required
def admin():
    rows = db().execute("SELECT r.*, u.name AS client FROM requests r JOIN users u ON u.id=r.user_id ORDER BY r.id DESC").fetchall()
    reqs = [to_dict(r) for r in rows]
    review_request_ids = {
        row["request_id"]
        for row in db().execute(
            "SELECT w.request_id FROM request_directions_workflows w JOIN requests r ON r.id=w.request_id "
            "WHERE w.status='ready' AND r.status NOT IN ('Completed','Rejected') "
            "UNION SELECT v.request_id FROM volumes_workflows v JOIN requests r ON r.id=v.request_id "
            "WHERE v.status='ready_for_review' AND r.status NOT IN ('Completed','Rejected')"
        ).fetchall()
    }
    for item in reqs:
        item["waiting_for_review"] = item["id"] in review_request_ids
        item["dashboard_status"] = "Waiting for review" if item["waiting_for_review"] else item["status"]
    today = dt.date.today(); counts = {}
    for r in reqs: counts[r["created"]] = counts.get(r["created"], 0) + 1
    days = [(today - dt.timedelta(days=i)).isoformat() for i in range(29, -1, -1)]
    stats = {"total": len(reqs),
             "clients": db().execute("SELECT COUNT(*) FROM users WHERE role='client'").fetchone()[0],
             "by_status": {s: sum(1 for r in reqs if r["status"] == s) for s in STATUSES},
             "sent_to_users": sum(1 for r in reqs if r["status"] == "Completed"),
             "waiting_review": len(review_request_ids),
             "by_service": {v: sum(1 for r in reqs if k in r["services"]) for k, v in all_service_names().items()},
             "by_hour": {"labels": [f"{h:02d}" for h in range(24)], "values": [sum(1 for r in reqs if h in r["hours"]) for h in range(24)]},
             "per_day": {"labels": days, "values": [counts.get(d, 0) for d in days]}}
    period_rows = [dict(x) for x in db().execute("SELECT * FROM periods ORDER BY sort_order,id").fetchall()]
    users = [dict(u) for u in db().execute(
        "SELECT u.id,u.name,u.email,u.role,COUNT(r.id) AS n,COALESCE(SUM(r.status='Completed'),0) AS done,MAX(r.created) AS last "
        "FROM users u LEFT JOIN requests r ON r.user_id=u.id GROUP BY u.id ORDER BY n DESC, u.id").fetchall()]
    return render_template("admin.html", reqs=reqs, stats=stats, users=users, statuses=STATUSES, periods=period_rows,
                           net={"have_geo": network.HAVE_GEO, "cfg": get_setting("network", {}), "services": net_services()})

def _delete_requests(request_ids=None):
    d = db()
    d.execute("BEGIN IMMEDIATE")
    if request_ids is None:
        rows = d.execute("SELECT id FROM requests").fetchall()
    else:
        rows = d.execute(
            f"SELECT id FROM requests WHERE id IN ({','.join('?' for _ in request_ids)})",
            request_ids,
        ).fetchall() if request_ids else []
    ids = [row["id"] for row in rows]
    if not ids:
        d.commit()
        return 0
    placeholders = ",".join("?" for _ in ids)
    d.execute(f"DELETE FROM notifications WHERE request_id IN ({placeholders})", ids)
    d.execute(f"DELETE FROM volume_freeflow_attempts WHERE request_id IN ({placeholders})", ids)
    d.execute(f"DELETE FROM volumes_workflows WHERE request_id IN ({placeholders})", ids)
    d.execute(f"DELETE FROM request_directions_workflows WHERE request_id IN ({placeholders})", ids)
    d.execute(f"DELETE FROM requests WHERE id IN ({placeholders})", ids)
    d.execute("DELETE FROM request_groups WHERE id NOT IN (SELECT DISTINCT group_id FROM requests WHERE group_id IS NOT NULL)")
    d.commit()
    return len(ids)

@app.post("/api/requests/<int:rid>/delete")
@admin_required
def delete_request(rid):
    deleted = _delete_requests([rid])
    if not deleted:
        return jsonify(error="Request not found"), 404
    return jsonify(ok=True, deleted=deleted)

@app.post("/api/requests/delete-all")
@admin_required
def delete_all_requests():
    payload = request.get_json(silent=True) or {}
    if payload.get("confirmation") != "DELETE_ALL_REQUESTS":
        return jsonify(error="Explicit confirmation is required to delete all requests."), 400
    deleted = _delete_requests()
    return jsonify(ok=True, deleted=deleted)

@app.route("/admin/requests/<int:rid>")
@admin_required
def admin_request(rid):
    r = db().execute("SELECT r.*, u.name AS client, u.email AS email FROM requests r JOIN users u ON u.id=r.user_id WHERE r.id=?", (rid,)).fetchone()
    if not r: abort(404)
    volumes_row = db().execute(
        "SELECT * FROM volumes_workflows WHERE request_id=?", (rid,)
    ).fetchone()
    volumes_workflow = dict(volumes_row) if volumes_row else None
    if volumes_workflow and volumes_workflow.get("result_json"):
        volumes_workflow["result"] = json.loads(volumes_workflow["result_json"])
    if volumes_workflow and volumes_workflow.get("freeflow_result_json"):
        volumes_workflow["freeflow_result"] = json.loads(volumes_workflow["freeflow_result_json"])
    if volumes_workflow and volumes_workflow.get("freeflow_validation_json"):
        volumes_workflow["freeflow_validation"] = json.loads(volumes_workflow["freeflow_validation_json"])
    if volumes_workflow and volumes_workflow.get("freeflow_job_id"):
        job_row = db().execute(
            "SELECT execution_status FROM outscraper_jobs WHERE id=?",
            (volumes_workflow["freeflow_job_id"],),
        ).fetchone()
        volumes_workflow["freeflow_job_status"] = job_row["execution_status"] if job_row else None
    if volumes_workflow:
        volumes_workflow["freeflow_attempts"] = []
        for attempt_row in db().execute(
            "SELECT job_id AS id,submitted_at,status,result_json,validation_json,error,accepted,override,accepted_at "
            "FROM volume_freeflow_attempts WHERE request_id=? ORDER BY submitted_at DESC LIMIT 100",
            (rid,),
        ).fetchall():
            attempt = dict(attempt_row)
            attempt["result"] = json.loads(attempt.pop("result_json") or "null")
            attempt["validation"] = json.loads(attempt.pop("validation_json") or "null")
            volumes_workflow["freeflow_attempts"].append(attempt)
    if volumes_workflow:
        request_data = json.loads(r["data"] or "{}")
        volumes_workflow["freeflow_context"] = (
            request_data.get("volumes") or {}
        ).get("freeflow_context")
    workflow_row = db().execute(
        "SELECT * FROM request_directions_workflows WHERE request_id=? ORDER BY created_at DESC LIMIT 1",
        (rid,),
    ).fetchone()
    directions_workflow = None
    if workflow_row:
        directions_workflow = dict(workflow_row)
        directions_workflow["result"] = json.loads(directions_workflow.pop("result_json") or "null")
    return render_template("admin_request.html", r=to_dict(r), forms=forms_catalog(), statuses=STATUSES,
                           directions_workflow=directions_workflow, volumes_workflow=volumes_workflow)

_timezone_finder = TimezoneFinder() if TimezoneFinder is not None else None

def _nearest_public_holiday(country_code, target_date):
    if holidays is None:
        raise ValueError("Install the project requirements to enable free-flow holiday dates.")
    calendar = holidays.country_holidays(
        country_code, years=range(target_date.year - 1, target_date.year + 2),
        observed=True, language="en",
    )
    candidates = sorted(calendar.items(), key=lambda item: (abs((item[0] - target_date).days), item[0]))
    if not candidates:
        raise ValueError(f"No public holidays are available for country code {country_code}.")
    holiday_date, holiday_name = candidates[0]
    return holiday_date, str(holiday_name)

def _latest_previous_weekly_weekend(country_code, timezone_name, now=None):
    if holidays is None:
        raise ValueError("Install the project requirements to identify weekly weekend dates.")
    try:
        timezone = ZoneInfo(timezone_name)
    except (TypeError, ZoneInfoNotFoundError) as exc:
        raise ValueError("The route time zone is not available for weekend date selection.") from exc
    current = now or dt.datetime.now(dt.timezone.utc)
    if current.tzinfo is None:
        current = current.replace(tzinfo=timezone)
    else:
        current = current.astimezone(timezone)
    try:
        calendar = holidays.country_holidays(
            country_code,
            years=(current.year - 1, current.year),
            observed=True,
            language="en",
        )
    except (NotImplementedError, ValueError) as exc:
        raise ValueError(f"No weekly weekend calendar is available for country code {country_code}.") from exc
    weekend_days = calendar.weekend
    if not weekend_days or any(
        isinstance(day, bool) or not isinstance(day, int) or not 0 <= day <= 6
        for day in weekend_days
    ):
        raise ValueError(f"No valid weekly weekend days are available for country code {country_code}.")
    candidate = current.date() - dt.timedelta(days=1)
    while candidate.weekday() not in weekend_days:
        candidate -= dt.timedelta(days=1)
    weekday_names = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
    return candidate, weekday_names[candidate.weekday()]

def _nearest_previous_weekly_weekend(country_code, timezone_name, study_date, now=None):
    try:
        timezone = ZoneInfo(timezone_name)
        study_day = dt.date.fromisoformat(study_date) if isinstance(study_date, str) else study_date
    except (TypeError, ValueError, ZoneInfoNotFoundError) as exc:
        raise ValueError("The study date or route time zone is not valid.") from exc
    current = now or dt.datetime.now(dt.timezone.utc)
    if current.tzinfo is None:
        current = current.replace(tzinfo=timezone)
    else:
        current = current.astimezone(timezone)
    try:
        calendar = holidays.country_holidays(
            country_code,
            years=(study_day.year - 1, study_day.year, current.year),
            observed=True,
            language="en",
        )
    except (NotImplementedError, ValueError) as exc:
        raise ValueError(f"No weekly weekend calendar is available for country code {country_code}.") from exc
    weekend_days = set(calendar.weekend or ())
    if not weekend_days:
        raise ValueError(f"No weekly weekend calendar is available for country code {country_code}.")
    candidates = [
        study_day + dt.timedelta(days=offset)
        for offset in range(-7, 8)
        if (study_day + dt.timedelta(days=offset)).weekday() in weekend_days
        and (study_day + dt.timedelta(days=offset)) < current.date()
    ]
    if not candidates:
        return _latest_previous_weekly_weekend(country_code, timezone_name, now=current)
    selected = min(candidates, key=lambda day: (abs((day - study_day).days), day))
    weekday_names = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
    return selected, weekday_names[selected.weekday()]

def _validate_holiday_source_request(configuration):
    context = configuration.get("trafficlens_context")
    if not context:
        return
    row = db().execute("SELECT * FROM requests WHERE id=?", (context["source_request_id"],)).fetchone()
    if not row:
        raise ValueError("The TrafficLens request associated with this reference preview no longer exists.")
    source = to_dict(row)
    try:
        expected_origin = ", ".join(f"{float(value):.6f}" for value in source["coords"][0])
        expected_destination = ", ".join(
            f"{float(value):.6f}" for value in (
                source["point_c"] if context.get("volume_input") else source["coords"][1]
            )
        )
        study_date = dt.date.fromisoformat(source["d1"])
        if context["service"] == "freeflow_speed":
            if context.get("volume_input"):
                configured_date = dt.date.fromisoformat(context.get("weekend_date", ""))
                today = dt.datetime.now(ZoneInfo(context["timezone"])).date()
                try:
                    calendar = holidays.country_holidays(
                        context["country_code"],
                        years=(configured_date.year, today.year),
                        observed=True,
                        language="en",
                    )
                except (NotImplementedError, ValueError) as exc:
                    raise ValueError("The route's weekly weekend calendar is unavailable.") from exc
                if configured_date >= today or configured_date.weekday() not in set(calendar.weekend or ()):
                    raise ValueError("Choose a weekly weekend date strictly before today.")
                reference_date = configured_date
                reference_label = configured_date.strftime("%A")
            else:
                reference_date, reference_label = _latest_previous_weekly_weekend(
                    context["country_code"], context["timezone"]
                )
            configured_date = context.get("weekend_date")
            configured_label = context.get("weekend_day")
        else:
            reference_date, reference_label = _nearest_public_holiday(
                context["country_code"], study_date
            )
            configured_date = context.get("holiday_date")
            configured_label = context.get("holiday_name")
    except (TypeError, ValueError, IndexError) as exc:
        raise ValueError("The source request no longer has a valid route or reference date.") from exc
    expected_region = context["country_code"] if context["country_code"] in outscraper.REGIONS else ""
    service_is_requested = context["service"] in source["services"]
    service_is_volume_input = context["service"] == "freeflow_speed" and "volumes" in source["services"]
    if (
        not (service_is_requested or service_is_volume_input)
        or context["study_date"] != study_date.isoformat()
        or configured_date != reference_date.isoformat()
        or configured_label != reference_label
        or configuration["region"] != expected_region
        or configuration["routes"][0]["origin"] != expected_origin
        or configuration["routes"][0]["destination"] != expected_destination
    ):
        reference_name = "weekly weekend date" if context["service"] == "freeflow_speed" else "public holiday"
        raise ValueError(f"The source route, study date, {reference_name}, or region changed. Prepare a new preview.")
    if context.get("volume_input"):
        try:
            timezone = ZoneInfo(context["timezone"])
            departure = dt.datetime.fromtimestamp(configuration["departure_time"], timezone)
            finish = dt.datetime.fromtimestamp(configuration["finish_time"], timezone)
            if (
                departure.date() != reference_date
                or finish.date() != reference_date
                or departure >= finish
                or configuration["routes"][0].get("departure_time") != configuration["departure_time"]
                or configuration["interval"] != context.get("interval_minutes")
                or context.get("departure_time") != configuration["departure_time"]
                or context.get("finish_time") != configuration["finish_time"]
            ):
                raise ValueError
        except (TypeError, ValueError, OverflowError, OSError, ZoneInfoNotFoundError) as exc:
            raise ValueError("The free-flow A→C date, time range, or interval no longer matches the reviewed request.") from exc

def _route_holiday_location(coordinate):
    latitude, longitude = coordinate
    return _route_holiday_location_cached(float(latitude), float(longitude))

@lru_cache(maxsize=256)
def _route_holiday_location_cached(latitude, longitude):
    if _timezone_finder is None:
        raise ValueError("Install the project requirements to identify route countries and time zones.")
    query = urlencode({
        "format": "jsonv2", "lat": f"{latitude:.6f}", "lon": f"{longitude:.6f}",
        "zoom": "10", "addressdetails": "1",
    })
    url = f"https://nominatim.openstreetmap.org/reverse?{query}"
    last_error = None
    for attempt in range(3):
        try:
            outbound = UrlRequest(url, headers={
                "User-Agent": "TrafficLens/1.0 (support@trafficlens.example)",
                "Accept": "application/json",
                "Accept-Language": "en",
            })
            with urlopen(outbound, timeout=8) as response:
                result = json.loads(response.read().decode("utf-8"))
            address = result.get("address") or {}
            country_code = str(address.get("country_code") or result.get("country_code") or "").upper()
            country_name = str(address.get("country") or result.get("country") or "").strip()
            if not country_name:
                display_name = str(result.get("display_name") or "").strip()
                if display_name:
                    country_name = display_name.split(",")[-1].strip()
            timezone_name = _timezone_finder.timezone_at(lat=latitude, lng=longitude)
            if len(country_code) != 2 or not country_name or not (timezone_name and isinstance(timezone_name, str)):
                raise ValueError("Incomplete route country or time zone metadata")
            ZoneInfo(timezone_name)
            return country_code, country_name, timezone_name
        except (OSError, UnicodeDecodeError, json.JSONDecodeError, AttributeError, KeyError, TypeError, ValueError, ZoneInfoNotFoundError, urllib.error.HTTPError) as exc:
            last_error = exc
            if attempt < 2:
                time.sleep(0.5 * (attempt + 1))
                continue
    raise ValueError("Could not identify the route country and local time zone from the route coordinates. Try again later or verify the route location.") from last_error

def _request_directions_location(request_row):
    source = to_dict(request_row)
    if not {"travel", "speed", "volumes"}.intersection(source["services"]):
        raise ValueError("This request does not include Travel time, Speed, or Traffic volumes.")
    if not source.get("point_c") or not source.get("road_km"):
        raise ValueError("This request has no calculated Point C or saved A-to-B road distance.")
    coordinates = source["path"] or source["coords"]
    midpoint = coordinates[len(coordinates) // 2]
    country_code, country_name, timezone_name = _route_holiday_location(midpoint)
    if country_code not in outscraper.REGIONS:
        raise ValueError(f"The detected country {country_name} ({country_code}) is not a supported Outscraper region.")
    return {
        "country_code": country_code,
        "country_name": country_name,
        "timezone": timezone_name,
    }

def _outscraper_prefill_from_request(request_row, service="travel"):
    source = to_dict(request_row)
    service_requested = service in source["services"]
    service_for_volumes = service in {"travel", "freeflow_speed"} and "volumes" in source["services"]
    volume_freeflow = service == "freeflow_speed" and "volumes" in source["services"]
    if service not in {"travel", "freeflow_time", "freeflow_speed"} or not (service_requested or service_for_volumes):
        abort(400, "This request did not include the selected Directions service.")
    if len(source["coords"]) != 2:
        abort(400, "This request does not contain valid start and end coordinates.")
    try:
        origin = ", ".join(f"{float(value):.6f}" for value in source["coords"][0])
        destination_point = source["point_c"] if volume_freeflow else source["coords"][1]
        if volume_freeflow and not destination_point:
            abort(400, "This request has no calculated Point C for a free-flow A→C request.")
        destination = ", ".join(f"{float(value):.6f}" for value in destination_point)
        study_day = dt.date.fromisoformat(source["d1"])
        route = {"origin": origin, "destination": destination}
        region = "KW"
        trafficlens_context = None
        if service == "travel":
            hours = sorted({
                int(hour) for hour in source["hours"]
                if str(hour).isdigit() and 0 <= int(hour) <= 23
            })
            if not hours:
                abort(400, "This request does not contain selected study hours.")
            timezone = ZoneInfo("Asia/Kuwait")
            departure_time = dt.time.fromisoformat(source.get("start_time") or f"{hours[0]:02d}:00:00")
            departure = dt.datetime.combine(study_day, departure_time, tzinfo=timezone)
            if source.get("end_time"):
                finish = dt.datetime.combine(study_day, dt.time.fromisoformat(source["end_time"]), tzinfo=timezone)
            else:
                finish = dt.datetime.combine(study_day, dt.time(), tzinfo=timezone) + dt.timedelta(hours=hours[-1] + 1)
        else:
            coordinates = source["path"] or source["coords"]
            midpoint = coordinates[len(coordinates) // 2]
            country_code, country_name, timezone_name = _route_holiday_location(midpoint)
            timezone = ZoneInfo(timezone_name)
            if service == "freeflow_speed":
                try:
                    if volume_freeflow:
                        reference_date, reference_label = _nearest_previous_weekly_weekend(
                            country_code, timezone_name, study_day
                        )
                    else:
                        reference_date, reference_label = _latest_previous_weekly_weekend(
                            country_code, timezone_name
                        )
                except ValueError as exc:
                    abort(400, str(exc))
                departure = dt.datetime.combine(reference_date, dt.time(7), tzinfo=timezone)
                finish = dt.datetime.combine(reference_date, dt.time(8), tzinfo=timezone)
                reference_context = {
                    "reference_type": "weekly_weekend",
                    "weekend_date": reference_date.isoformat(),
                    "weekend_day": reference_label,
                }
            else:
                try:
                    reference_date, reference_label = _nearest_public_holiday(country_code, study_day)
                except (NotImplementedError, ValueError) as exc:
                    abort(400, str(exc) or f"No supported public holiday calendar is available for route country {country_code}.")
                departure = dt.datetime.combine(reference_date, dt.time(7), tzinfo=timezone)
                reference_context = {
                    "reference_type": "public_holiday",
                    "holiday_date": reference_date.isoformat(),
                    "holiday_name": reference_label,
                }
            departure_timestamp = int(departure.timestamp())
            route["departure_time"] = departure_timestamp
            region = country_code if country_code in outscraper.REGIONS else ""
            trafficlens_context = {
                "source_request_id": source["id"], "service": service,
                "study_date": study_day.isoformat(), "country_code": country_code,
                "country_name": country_name, "timezone": timezone_name,
                "departure_time": departure_timestamp,
                "origin": origin, "destination": destination,
                **reference_context,
            }
            if volume_freeflow:
                trafficlens_context.update({
                    "volume_input": True,
                    "road_ab_meters": float(source["road_km"]) * 1000,
                    "finish_time": int(finish.timestamp()),
                    "interval_minutes": 60,
                })
    except (TypeError, ValueError, IndexError, ZoneInfoNotFoundError):
        abort(400, "This request's route, date, or selected hours cannot be used to prepare directions.")
    prefill_config = {
        "routes": [route],
        "departure_time": int(departure.timestamp()) if service == "travel" or volume_freeflow else None,
        "finish_time": int(finish.timestamp()) if service == "travel" or volume_freeflow else None,
        "interval": 1 if service == "travel" else 60 if volume_freeflow else None,
        "travel_mode": "car",
        "language": "en",
        "region": region,
    }
    if trafficlens_context:
        prefill_config["trafficlens_context"] = trafficlens_context
    return {
        "source_request": {
            "id": source["id"],
            "date": source["d1"],
            "service": service,
            "start_time": source.get("start_time"),
            "end_time": source.get("end_time"),
            "point_c": source.get("point_c"),
            "hours": source["hours"],
            "hours_label": source["hours_label"],
            "area": source["area"],
        },
        "prefill_config": prefill_config,
    }

@app.get("/api/requests/<int:rid>/volumes/freeflow-preview")
@admin_required
def prepare_volumes_freeflow_preview(rid):
    source_row = db().execute("SELECT * FROM requests WHERE id=?", (rid,)).fetchone()
    if not source_row:
        return jsonify(error="Request not found."), 404
    if "volumes" not in json.loads(source_row["services"] or "[]"):
        return jsonify(error="This request does not include Traffic volumes."), 404
    workflow = db().execute(
        "SELECT freeflow_accepted,freeflow_job_id FROM volumes_workflows WHERE request_id=?",
        (rid,),
    ).fetchone()
    if not workflow:
        return jsonify(error="The Traffic volumes workflow is not available."), 404
    if workflow["freeflow_accepted"]:
        return jsonify(error="A free-flow speed has already been accepted for this request."), 409
    previous_configuration = None
    if workflow["freeflow_job_id"]:
        active_job = db().execute(
            "SELECT execution_status,config_json FROM outscraper_jobs WHERE id=?",
            (workflow["freeflow_job_id"],),
        ).fetchone()
        if active_job and active_job["execution_status"] in {
            "submitting", "running", "unknown",
        }:
            return jsonify(error="A free-flow speed request is already in progress."), 409
        if active_job:
            previous_configuration = json.loads(active_job["config_json"])
    if not outscraper.api_key_configured():
        return jsonify(error="OUTSCRAPER_API_KEY is not configured on this server."), 503
    try:
        prepared = _outscraper_prefill_from_request(source_row, "freeflow_speed")
    except (ValueError, TypeError) as exc:
        return jsonify(error=str(exc)), 400
    configuration = prepared["prefill_config"]
    if (
        previous_configuration
        and previous_configuration.get("trafficlens_context", {}).get("volume_input")
    ):
        configuration = previous_configuration
    return jsonify(
        source_request=prepared["source_request"],
        configuration=configuration,
    )

@app.route("/admin/outscraper")
@admin_required
def admin_outscraper():
    start, end = request.args.get("from", "").strip(), request.args.get("to", "").strip()
    status, query = request.args.get("status", "").strip(), request.args.get("q", "").strip()[:200]
    where, params = [], []
    if start:
        try: dt.date.fromisoformat(start)
        except ValueError: abort(400, "Invalid start date")
        where.append("created_at >= ?"); params.append(start + "T00:00:00")
    if end:
        try: dt.date.fromisoformat(end)
        except ValueError: abort(400, "Invalid end date")
        where.append("created_at <= ?"); params.append(end + "T23:59:59")
    if status:
        if status not in {"pending", "approved", "rejected", "invalidated", "awaiting_approval",
                          "submitting", "running", "completed", "failed", "unknown"}:
            abort(400, "Invalid status filter")
        where.append("(approval_status=? OR execution_status=?)"); params.extend((status, status))
    if query:
        where.append("EXISTS (SELECT 1 FROM outscraper_job_routes r WHERE r.job_id=outscraper_jobs.id AND (r.origin LIKE ? OR r.destination LIKE ?))")
        params.extend((f"%{query}%", f"%{query}%"))
    clause = " WHERE " + " AND ".join(where) if where else ""
    jobs = db().execute(
        "SELECT * FROM outscraper_jobs" + clause + " ORDER BY created_at DESC LIMIT 250", params).fetchall()
    summaries = []
    for row in jobs:
        job = dict(row)
        job["config"] = json.loads(job.pop("config_json"))
        job["routes"] = [dict(route) for route in db().execute(
            "SELECT route_index, origin, destination, execution_status FROM outscraper_job_routes WHERE job_id=? ORDER BY route_index",
            (job["id"],)).fetchall()]
        job["provider_ids"] = json.loads(job["provider_ids"] or "[]")
        summaries.append(job)
    source_request = None
    prefill_config = None
    copied_job_id = None
    source_request_id = request.args.get("source_request_id", "").strip()
    if source_request_id:
        try: source_id = int(source_request_id)
        except ValueError: abort(400, "Invalid source request ID")
        source_row = db().execute("SELECT * FROM requests WHERE id=?", (source_id,)).fetchone()
        if not source_row: abort(404)
        source_service = request.args.get("service", "travel").strip()
        prefill = _outscraper_prefill_from_request(source_row, source_service)
        source_request, prefill_config = prefill["source_request"], prefill["prefill_config"]
    duplicate_job_id = request.args.get("duplicate_job_id", "").strip()
    if duplicate_job_id:
        previous = db().execute("SELECT config_json FROM outscraper_jobs WHERE id=?", (duplicate_job_id,)).fetchone()
        if not previous: abort(404)
        try: prefill_config = outscraper.validate_configuration(json.loads(previous["config_json"]))
        except (ValueError, json.JSONDecodeError): abort(400, "This saved request configuration cannot be edited.")
        copied_job_id = duplicate_job_id
    return render_template("outscraper.html", jobs=summaries, api_configured=outscraper.api_key_configured(),
                           languages=sorted(outscraper.LANGUAGES), regions=sorted(outscraper.REGIONS),
                           source_request=source_request, prefill_config=prefill_config,
                           copied_job_id=copied_job_id,
                           filters={"from": start, "to": end, "status": status, "q": query})

@app.post("/api/outscraper/csv-preview")
@admin_required
def outscraper_csv_preview():
    if request.content_length and request.content_length > 7_000_000:
        return jsonify(error="CSV request body is too large"), 413
    payload = request.get_json(silent=True) or {}
    try:
        routes = outscraper.parse_routes_csv(payload.get("csv", ""))
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    return jsonify(routes=routes, count=len(routes))

@app.post("/api/outscraper/preview")
@admin_required
def outscraper_preview():
    payload = request.get_json(silent=True)
    try:
        config = outscraper.validate_configuration(payload)
        _validate_holiday_source_request(config)
        _validate_trafficlens_request(config)
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    holiday_context = config.get("trafficlens_context") or {}
    if holiday_context.get("service") == "freeflow_speed":
        volume_workflow = db().execute(
            "SELECT freeflow_accepted,freeflow_job_id FROM volumes_workflows WHERE request_id=?",
            (holiday_context["source_request_id"],),
        ).fetchone()
        if not volume_workflow:
            return jsonify(error="The Traffic volumes workflow is not available."), 404
        if holiday_context.get("volume_input") and volume_workflow["freeflow_accepted"]:
            return jsonify(error="The free-flow speed has already been accepted for this request."), 409
        if volume_workflow["freeflow_job_id"]:
            active_job = db().execute(
                "SELECT execution_status FROM outscraper_jobs WHERE id=?",
                (volume_workflow["freeflow_job_id"],),
            ).fetchone()
            if active_job and active_job["execution_status"] in {
                "submitting", "running", "unknown",
            }:
                return jsonify(error="A free-flow speed request is already in progress."), 409
    if config.get("trafficlens_request") and not outscraper.api_key_configured():
        return jsonify(error="OUTSCRAPER_API_KEY is not configured on this server."), 503
    job_id, token, now = secrets.token_hex(16), secrets.token_urlsafe(32), _timestamp_now()
    encoded = json.dumps(config, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    fingerprint = outscraper.configuration_fingerprint(config)
    d = db()
    d.execute("BEGIN IMMEDIATE")
    request_context = config.get("trafficlens_request")
    if request_context:
        active_workflow = d.execute(
            "SELECT job_id,status FROM request_directions_workflows WHERE request_id=? "
            "AND status NOT IN ('failed','rejected','invalidated') ORDER BY created_at DESC LIMIT 1",
            (request_context["source_request_id"],),
        ).fetchone()
        if active_workflow:
            d.rollback()
            return jsonify(
                error=f"This client request already has a Directions workflow ({active_workflow['status']}). "
                      "Open its current review or Scraper details before starting another request."
            ), 409
    for old_id, old_token in session.get("outscraper_tokens", {}).items():
        d.execute(
            "UPDATE outscraper_jobs SET approval_status='invalidated', execution_status='invalidated', updated_at=? "
            "WHERE id=? AND approval_status='pending' AND execution_status='awaiting_approval' AND approval_token_hash=?",
            (now, old_id, _token_hash(old_token)))
        if d.execute("SELECT changes()").fetchone()[0]:
            d.execute("UPDATE outscraper_job_routes SET execution_status='invalidated' WHERE job_id=?", (old_id,))
            d.execute(
                "UPDATE request_directions_workflows SET status='invalidated', updated_at=? WHERE job_id=?",
                (now, old_id),
            )
    d.execute(
        "INSERT INTO outscraper_jobs(id,created_at,updated_at,config_json,config_hash,approval_token_hash,estimated_cost,estimated_currency,estimate_confidence,actual_cost,actual_currency) "
        "VALUES(?,?,?,?,?,?,NULL,NULL,'unknown',NULL,NULL)",
        (job_id, now, now, encoded, fingerprint, _token_hash(token)))
    d.executemany(
        "INSERT INTO outscraper_job_routes(job_id,route_index,origin,destination) VALUES(?,?,?,?)",
        [(job_id, index, route["origin"], route["destination"])
         for index, route in enumerate(config["routes"], 1)])
    if request_context:
        d.execute(
            "INSERT INTO request_directions_workflows(job_id,request_id,status,interval_seconds,interval_minutes,created_at,updated_at) "
            "VALUES(?,?,'awaiting_approval',?,?,?,?)",
            (job_id, request_context["source_request_id"], request_context["interval_minutes"] * 60,
             request_context["interval_minutes"], now, now),
        )
    holiday_context = config.get("trafficlens_context") or {}
    if holiday_context.get("service") == "freeflow_speed":
        if holiday_context.get("volume_input"):
            d.execute(
                "INSERT INTO volume_freeflow_attempts(job_id,request_id,submitted_at,status) "
                "VALUES(?,?,?,'awaiting_approval')",
                (job_id, holiday_context["source_request_id"], now),
            )
        d.execute(
            "UPDATE volumes_workflows SET freeflow_job_id=?, freeflow_speed=NULL, "
            "freeflow_result_json=NULL, freeflow_validation_json=NULL, freeflow_accepted=0, "
            "freeflow_override=0, freeflow_accepted_at=NULL, status='waiting_for_inputs', "
            "error=NULL, updated_at=? WHERE request_id=? AND freeflow_accepted=0",
            (job_id, now, holiday_context["source_request_id"]),
        )
    d.commit()
    session["outscraper_tokens"] = {job_id: token}
    return jsonify(id=job_id, approval_token=token, configuration=config,
                   route_count=len(config["routes"]), planned_http_submissions=len(config["routes"]),
                   cost_estimate=None, currency=None, estimate_confidence="unknown"), 201

@app.post("/api/outscraper/<job_id>/invalidate")
@admin_required
def invalidate_outscraper_preview(job_id):
    payload = request.get_json(silent=True) or {}
    token = payload.get("approval_token", "")
    if not isinstance(token, str): return jsonify(error="Invalid approval token"), 400
    cursor = db().execute(
        "UPDATE outscraper_jobs SET approval_status='invalidated', execution_status='invalidated', updated_at=? "
        "WHERE id=? AND approval_status='pending' AND execution_status='awaiting_approval' AND approval_token_hash=?",
        (_timestamp_now(), job_id, _token_hash(token)))
    if cursor.rowcount != 1: return jsonify(error="Preview is no longer awaiting approval"), 409
    db().execute("UPDATE outscraper_job_routes SET execution_status='invalidated' WHERE job_id=?", (job_id,))
    db().execute(
        "UPDATE request_directions_workflows SET status='invalidated', updated_at=? WHERE job_id=?",
        (_timestamp_now(), job_id),
    )
    db().commit()
    _forget_outscraper_token(job_id)
    return jsonify(ok=True)

@app.post("/api/outscraper/<job_id>/reject")
@admin_required
def reject_outscraper_job(job_id):
    payload = request.get_json(silent=True) or {}
    token = payload.get("approval_token", "")
    if not isinstance(token, str): return jsonify(error="Invalid approval token"), 400
    now = _timestamp_now()
    cursor = db().execute(
        "UPDATE outscraper_jobs SET approval_status='rejected', execution_status='rejected', rejected_at=?, updated_at=? "
        "WHERE id=? AND approval_status='pending' AND execution_status='awaiting_approval' AND approval_token_hash=?",
        (now, now, job_id, _token_hash(token)))
    if cursor.rowcount != 1: return jsonify(error="Request is no longer awaiting approval"), 409
    db().execute("UPDATE outscraper_job_routes SET execution_status='rejected' WHERE job_id=?", (job_id,))
    db().execute(
        "UPDATE request_directions_workflows SET status='rejected', updated_at=? WHERE job_id=?",
        (now, job_id),
    )
    db().commit()
    _forget_outscraper_token(job_id)
    return jsonify(ok=True, rejected_at=now)

@app.post("/api/outscraper/<job_id>/approve")
@admin_required
def approve_outscraper_job(job_id):
    payload = request.get_json(silent=True) or {}
    token, proposed = payload.get("approval_token", ""), payload.get("configuration")
    if not isinstance(token, str): return jsonify(error="Invalid approval token"), 400
    if payload.get("acknowledge_unknown_cost") is not True:
        return jsonify(error="Cost is unknown. Explicitly approve execution despite the unknown cost."), 409
    try:
        config = outscraper.validate_configuration(proposed)
        _validate_holiday_source_request(config)
        _validate_trafficlens_request(config)
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    if not outscraper.api_key_configured():
        return jsonify(error="OUTSCRAPER_API_KEY is not configured on this server"), 503
    now = _timestamp_now()
    d = db()
    d.execute("BEGIN IMMEDIATE")
    job = d.execute("SELECT config_hash, approval_token_hash FROM outscraper_jobs WHERE id=? AND approval_status='pending' AND execution_status='awaiting_approval'", (job_id,)).fetchone()
    if not job:
        d.rollback()
        return jsonify(error="Request is no longer awaiting approval; duplicate submission was prevented."), 409
    proposed_hash = outscraper.configuration_fingerprint(config)
    if not secrets.compare_digest(job["config_hash"], proposed_hash) or not secrets.compare_digest(job["approval_token_hash"], _token_hash(token)):
        d.rollback()
        return jsonify(error="Configuration changed or approval expired. Preview and approve it again."), 409
    d.execute(
        "UPDATE outscraper_jobs SET approval_status='approved', execution_status='submitting', approved_at=?, started_at=?, updated_at=? WHERE id=?",
        (now, now, now, job_id))
    d.execute(
        "UPDATE volume_freeflow_attempts SET status='submitting',error=NULL WHERE job_id=?",
        (job_id,),
    )
    d.execute("UPDATE outscraper_job_routes SET execution_status='queued' WHERE job_id=?", (job_id,))
    d.execute(
        "UPDATE request_directions_workflows SET status='running',error=NULL,updated_at=? WHERE job_id=?",
        (now, job_id),
    )
    d.commit()
    _forget_outscraper_token(job_id)
    with _outscraper_workers_lock:
        if job_id in _outscraper_workers:
            return jsonify(error="Request is already being processed"), 409
        _outscraper_workers.add(job_id)
    worker = threading.Thread(target=_process_outscraper_job, args=(job_id,), name=f"outscraper-{job_id}", daemon=True)
    try:
        worker.start()
    except RuntimeError:
        with _outscraper_workers_lock:
            _outscraper_workers.discard(job_id)
        d.execute("UPDATE outscraper_jobs SET execution_status='unknown', error=?, updated_at=? WHERE id=?",
                  ("Could not start the local worker. Do not resubmit automatically.", _timestamp_now(), job_id))
        d.execute("UPDATE outscraper_job_routes SET execution_status='unknown', error=? WHERE job_id=?",
                  ("Local submission did not start; verify request status before taking further action.", job_id))
        d.commit()
        return jsonify(error="Execution could not start safely; request marked unknown and will not be resubmitted."), 503
    return jsonify(ok=True, status="submitting", id=job_id), 202

@app.post("/api/outscraper/<job_id>/refresh")
@admin_required
def refresh_outscraper_job(job_id):
    job = _outscraper_job(job_id)
    if not job: return jsonify(error="Request not found"), 404
    polled = []
    for route in job["routes"]:
        if not route["provider_request_id"] or route["execution_status"] not in ("running", "unknown"):
            continue
        last = route.get("last_polled_at")
        if last:
            try:
                elapsed = (dt.datetime.now(dt.timezone.utc) - dt.datetime.fromisoformat(last)).total_seconds()
            except ValueError:
                elapsed = 15
            if elapsed < 15:
                polled.append({"route": route["route_index"], "status": route["execution_status"], "throttled": True})
                continue
        state = _poll_outscraper_route(route)
        polled.append({"route": route["route_index"], "status": state})
    status = _refresh_job_status(job_id)
    return jsonify(ok=True, status=status, routes=polled)

ABANDON_MESSAGE = (
    "Stopped waiting: the Scraper returned no data. The original request is kept in history; "
    "a new request can be prepared and approved."
)

@app.post("/api/outscraper/<job_id>/abandon")
@admin_required
def abandon_outscraper_job(job_id):
    """Stop waiting for a Directions job that never returned data so a NEW request can be prepared.

    This never resubmits anything: the old job stays in history, and a re-send still needs its own
    preview, cost acknowledgement and approval. The provider may still bill for the original call.
    """
    payload = request.get_json(silent=True) or {}
    if payload.get("acknowledge_possible_charge") is not True:
        return jsonify(error="Confirm that the provider may still bill the original request before stopping the wait."), 409
    d = db()
    d.execute("BEGIN IMMEDIATE")
    job = d.execute(
        "SELECT execution_status, updated_at FROM outscraper_jobs WHERE id=?", (job_id,)
    ).fetchone()
    if not job:
        d.rollback()
        return jsonify(error="Request not found"), 404
    state = job["execution_status"]
    stale_submit = False
    if state == "submitting":
        try:
            age = (dt.datetime.now(dt.timezone.utc) - dt.datetime.fromisoformat(job["updated_at"])).total_seconds()
        except (TypeError, ValueError):
            age = 0
        stale_submit = age > 300
    if state not in ("running", "unknown") and not stale_submit:
        d.rollback()
        return jsonify(error="Only a request that is still waiting for the provider, or whose status is uncertain, can be stopped."), 409
    now = _timestamp_now()
    d.execute(
        "UPDATE outscraper_job_routes SET execution_status='abandoned', error=?, completed_at=COALESCE(completed_at,?) "
        "WHERE job_id=? AND execution_status IN ('queued','submitting','running','unknown')",
        (ABANDON_MESSAGE, now, job_id),
    )
    d.execute(
        "UPDATE outscraper_jobs SET execution_status='abandoned', error=?, completed_at=COALESCE(completed_at,?), updated_at=? WHERE id=?",
        (ABANDON_MESSAGE, now, now, job_id),
    )
    d.execute(
        "UPDATE request_directions_workflows SET status='failed', error=?, updated_at=? "
        "WHERE job_id=? AND status NOT IN ('ready','confirmed')",
        (ABANDON_MESSAGE, now, job_id),
    )
    d.execute(
        "UPDATE volume_freeflow_attempts SET status='abandoned', error=? WHERE job_id=? AND accepted=0",
        (ABANDON_MESSAGE, job_id),
    )
    d.execute(
        "UPDATE volumes_workflows SET status=CASE WHEN status IN ('running','ready_for_review','sent') THEN status ELSE 'waiting_for_inputs' END, "
        "error=?, updated_at=? WHERE freeflow_job_id=? AND freeflow_accepted=0",
        (ABANDON_MESSAGE, now, job_id),
    )
    d.commit()
    return jsonify(ok=True, status="abandoned", message=ABANDON_MESSAGE)

@app.get("/api/requests/<int:rid>/directions")
@admin_required
def request_directions_status(rid):
    workflow = db().execute(
        "SELECT w.*, j.execution_status AS job_status, j.error AS job_error "
        "FROM request_directions_workflows w JOIN outscraper_jobs j ON j.id=w.job_id "
        "WHERE w.request_id=? ORDER BY w.created_at DESC LIMIT 1",
        (rid,),
    ).fetchone()
    if not workflow:
        return jsonify(workflow=None)
    result = json.loads(workflow["result_json"]) if workflow["result_json"] else None
    return jsonify({
        "workflow": {
            "job_id": workflow["job_id"],
            "status": workflow["status"],
            "job_status": workflow["job_status"],
            "interval_minutes": workflow["interval_minutes"] or max(1, math.ceil(workflow["interval_seconds"] / 60)),
            "result": result,
            "error": workflow["error"] or workflow["job_error"],
            "confirmed_at": workflow["confirmed_at"],
            "view_url": url_for("outscraper_job_detail", job_id=workflow["job_id"]),
        }
    })

@app.get("/api/requests/<int:rid>/directions/location")
@admin_required
def request_directions_location(rid):
    row = db().execute("SELECT * FROM requests WHERE id=?", (rid,)).fetchone()
    if not row:
        return jsonify(error="Source request not found."), 404
    try:
        location = _request_directions_location(row)
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    return jsonify(location=location)

@app.post("/api/requests/<int:rid>/volumes/lanes")
@admin_required
def save_volumes_lanes(rid):
    worker_key = f"volumes-{rid}"
    with _outscraper_workers_lock:
        if worker_key in _outscraper_workers:
            return jsonify(error="Lane count cannot be changed while MATLAB inputs are being saved or calculated."), 409
    payload = request.get_json(silent=True) or {}
    try:
        lanes = float(payload.get("lanes"))
    except (TypeError, ValueError):
        return jsonify(error="Enter a valid lane count."), 400
    if not math.isfinite(lanes) or not lanes.is_integer() or not 1 <= lanes <= MAX_VOLUME_LANES:
        return jsonify(error=f"Lane count must be a whole number from 1 to {MAX_VOLUME_LANES}."), 400
    lanes = int(lanes)
    workflow = db().execute(
        "SELECT status FROM volumes_workflows WHERE request_id=?", (rid,)
    ).fetchone()
    if not workflow:
        return jsonify(error="This request does not include Traffic volumes."), 404
    if workflow["status"] in {"running", "ready_for_review", "sent"}:
        return jsonify(error="Lane count cannot be changed while MATLAB is running or after a result is ready for review."), 409
    request_row = db().execute(
        "SELECT services,data FROM requests WHERE id=?", (rid,)
    ).fetchone()
    if not request_row:
        return jsonify(error="Source request not found."), 404
    requested_services = set(json.loads(request_row["services"] or "[]"))
    request_data = json.loads(request_row["data"] or "{}")
    _save_volume_lanes_service(request_data, requested_services, lanes)
    db().execute("BEGIN IMMEDIATE")
    db().execute(
        "UPDATE volumes_workflows SET lanes=?, status='waiting_for_inputs', error=NULL, updated_at=? WHERE request_id=?",
        (lanes, _timestamp_now(), rid),
    )
    if "lanes" in requested_services:
        db().execute("UPDATE requests SET data=? WHERE id=?", (json.dumps(request_data, ensure_ascii=False), rid))
    db().commit()
    _start_volume_calculation(rid)
    return jsonify(ok=True)

@app.get("/api/requests/<int:rid>/volumes")
@admin_required
def volumes_workflow_status(rid):
    workflow = db().execute(
        "SELECT * FROM volumes_workflows WHERE request_id=?", (rid,)
    ).fetchone()
    if not workflow:
        return jsonify(error="This request does not include Traffic volumes."), 404
    actual_hours_complete = _sync_volume_actual_speeds(rid)
    validation = json.loads(workflow["freeflow_validation_json"] or "{}")
    if (
        workflow["freeflow_result_json"]
        and validation.get("missing_hours")
        and actual_hours_complete
    ):
        if _evaluate_volume_freeflow(rid):
            _start_volume_calculation(rid)
        workflow = db().execute(
            "SELECT * FROM volumes_workflows WHERE request_id=?", (rid,)
        ).fetchone()
    if workflow["freeflow_speed"] is None and workflow["freeflow_job_id"]:
        completed_job = db().execute(
            "SELECT execution_status FROM outscraper_jobs WHERE id=?",
            (workflow["freeflow_job_id"],),
        ).fetchone()
        if completed_job and completed_job["execution_status"] == "completed":
            _persist_holiday_result_context(workflow["freeflow_job_id"])
            workflow = db().execute(
                "SELECT * FROM volumes_workflows WHERE request_id=?", (rid,)
            ).fetchone()
    result = json.loads(workflow["result_json"]) if workflow["result_json"] else None
    request_row = db().execute("SELECT data FROM requests WHERE id=?", (rid,)).fetchone()
    request_data = json.loads(request_row["data"] or "{}") if request_row else {}
    freeflow_job = None
    if workflow["freeflow_job_id"]:
        freeflow_job = db().execute(
            "SELECT execution_status FROM outscraper_jobs WHERE id=?",
            (workflow["freeflow_job_id"],),
        ).fetchone()
    attempts = []
    for attempt_row in db().execute(
        "SELECT job_id AS id,submitted_at,status,result_json,validation_json,error,accepted,override,accepted_at "
        "FROM volume_freeflow_attempts WHERE request_id=? ORDER BY submitted_at DESC LIMIT 100",
        (rid,),
    ).fetchall():
        attempt = dict(attempt_row)
        attempt["result"] = json.loads(attempt.pop("result_json") or "null")
        attempt["validation"] = json.loads(attempt.pop("validation_json") or "null")
        attempts.append(attempt)
    return jsonify(
        workflow={
            "status": workflow["status"],
            "lanes": workflow["lanes"],
            "freeflow_speed": workflow["freeflow_speed"],
            "freeflow_context": (request_data.get("volumes") or {}).get("freeflow_context"),
            "freeflow_job_id": workflow["freeflow_job_id"],
            "freeflow_job_status": freeflow_job["execution_status"] if freeflow_job else None,
            "freeflow_accepted": bool(workflow["freeflow_accepted"]),
            "freeflow_override": bool(workflow["freeflow_override"]),
            "freeflow_validation": json.loads(workflow["freeflow_validation_json"] or "null"),
            "freeflow_result": json.loads(workflow["freeflow_result_json"] or "null"),
            "freeflow_attempts": attempts,
            "error": workflow["error"],
            "result": result,
        }
    )

@app.post("/api/requests/<int:rid>/volumes/freeflow/accept")
@admin_required
def accept_volume_freeflow(rid):
    payload = request.get_json(silent=True) or {}
    use_anyway = payload.get("use_anyway") is True
    workflow = db().execute(
        "SELECT * FROM volumes_workflows WHERE request_id=?", (rid,)
    ).fetchone()
    if not workflow:
        return jsonify(error="This request does not include Traffic volumes."), 404
    if workflow["freeflow_accepted"]:
        return jsonify(error="The free-flow speed has already been accepted."), 409
    actual_hours_complete = _sync_volume_actual_speeds(rid)
    validation = json.loads(workflow["freeflow_validation_json"] or "{}")
    if (
        workflow["freeflow_result_json"]
        and validation.get("missing_hours")
        and actual_hours_complete
        and _evaluate_volume_freeflow(rid)
    ):
        started = _start_volume_calculation(rid)
        current = db().execute(
            "SELECT status FROM volumes_workflows WHERE request_id=?", (rid,)
        ).fetchone()
        return jsonify(
            ok=True,
            accepted=True,
            overridden=False,
            status=current["status"] if current else "waiting_for_inputs",
            matlab_started=started,
        )
    workflow = db().execute(
        "SELECT * FROM volumes_workflows WHERE request_id=?", (rid,)
    ).fetchone()
    if not workflow["freeflow_result_json"] or not workflow["freeflow_validation_json"]:
        return jsonify(error="A complete free-flow response and comparison are required before acceptance."), 409
    validation = json.loads(workflow["freeflow_validation_json"])
    if validation.get("missing_hours"):
        return jsonify(error="Actual Directions speeds are missing for selected hours. Complete and confirm every requested hour before accepting free-flow speed."), 409
    if not validation.get("is_valid") and not use_anyway:
        return jsonify(error="The free-flow speed is not higher than actual speed for every selected hour. Explicitly confirm Use anyway to continue."), 409
    if use_anyway and not validation.get("mismatched_hours"):
        return jsonify(error="Use anyway is available only when the comparison has mismatched hours."), 409
    now = _timestamp_now()
    source_result = json.loads(workflow["freeflow_result_json"])
    request_row = db().execute(
        "SELECT services,data,user_id FROM requests WHERE id=?", (rid,)
    ).fetchone()
    if not request_row:
        return jsonify(error="Source request not found."), 404
    request_data = json.loads(request_row["data"] or "{}")
    requested_services = set(json.loads(request_row["services"] or "[]"))
    if "freeflow_speed" in requested_services:
        request_data["freeflow_speed"] = {
            "rows": [{"value": source_result["selected_speed_kmh"]}],
            "note": "Selected as the highest speed returned for the reviewed A→C weekly-weekend window; accepted with an explicit administrator override.",
            "source": "outscraper",
            "updated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
            "outscraper_job_id": workflow["freeflow_job_id"],
            "reference_context": source_result.get("context"),
        }
        _create_notification(
            request_row["user_id"], rid, "directions_result",
            "Free-flow speed results are ready",
            f"The administrator approved free-flow speed for request REQ-{rid}. Traffic volumes remain under review.",
        )
        db().execute(
            "UPDATE requests SET data=?,status=CASE WHEN status NOT IN ('Rejected','Completed') "
            "THEN 'Processing' ELSE status END WHERE id=?",
            (json.dumps(request_data, ensure_ascii=False), rid),
        )
    db().execute(
        "UPDATE volumes_workflows SET freeflow_accepted=1,freeflow_override=?,"
        "freeflow_accepted_at=?,error=NULL,updated_at=? WHERE request_id=? AND freeflow_accepted=0",
        (int(use_anyway), now, now, rid),
    )
    if db().execute("SELECT changes()").fetchone()[0] != 1:
        db().rollback()
        return jsonify(error="The free-flow workflow changed before acceptance. Reload and review the latest response."), 409
    db().execute(
        "UPDATE volume_freeflow_attempts SET status='accepted',accepted=1,override=?,"
        "accepted_at=?,validation_json=? WHERE job_id=?",
        (
            int(use_anyway),
            now,
            json.dumps(validation, ensure_ascii=False),
            workflow["freeflow_job_id"],
        ),
    )
    db().commit()
    started = _start_volume_calculation(rid)
    return jsonify(
        ok=True,
        accepted=True,
        overridden=use_anyway,
        status=db().execute(
            "SELECT status FROM volumes_workflows WHERE request_id=?", (rid,)
        ).fetchone()["status"],
        matlab_started=started,
    )

@app.post("/api/requests/<int:rid>/volumes/run")
@admin_required
def retry_volumes_calculation(rid):
    workflow = db().execute(
        "SELECT status FROM volumes_workflows WHERE request_id=?", (rid,)
    ).fetchone()
    if not workflow:
        return jsonify(error="This request does not include Traffic volumes."), 404
    if workflow["status"] != "failed":
        return jsonify(error="Only a failed MATLAB calculation can be retried."), 409
    db().execute(
        "UPDATE volumes_workflows SET status='waiting_for_inputs', error=NULL, updated_at=? WHERE request_id=?",
        (_timestamp_now(), rid),
    )
    db().commit()
    started = _start_volume_calculation(rid)
    current = db().execute(
        "SELECT status FROM volumes_workflows WHERE request_id=?", (rid,)
    ).fetchone()
    status = current["status"] if current else "waiting_for_inputs"
    return jsonify(ok=started, status=status)

@app.get("/api/requests/<int:rid>/volumes/workbook")
@login_required
def download_volumes_workbook(rid):
    """Serve this request's own MATLAB result workbook: its client after approval, or an admin."""
    row = db().execute(
        "SELECT r.user_id, w.status, w.workbook_path FROM requests r "
        "JOIN volumes_workflows w ON w.request_id=r.id WHERE r.id=?",
        (rid,),
    ).fetchone()
    is_admin = session.get("role") == "admin"
    if not row or (row["user_id"] != session["uid"] and not is_admin):
        abort(404)
    if not is_admin and row["status"] != "sent":
        abort(404)
    try:
        resolved = Path(row["workbook_path"] or "").resolve(strict=True)
        resolved.relative_to(volumes_runner.volume_runs_root().resolve())
    except (OSError, ValueError):
        abort(404)
    if not resolved.is_file() or not resolved.name.endswith("-results.xlsx"):
        abort(404)
    return send_file(
        resolved,
        as_attachment=True,
        download_name=f"REQ-{rid}-volumes.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

def _directions_service_data(requested_services, result, job_id):
    """Speed / Travel service payloads built from a reviewed A->C Directions result."""
    updated = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    hourly = result["hourly"]
    reuse_note = (
        " No separate result was returned for "
        + ", ".join(f"{hour:02d}:00" for hour in result.get("reused_hours", []))
        + "; the available completed result was reused for those hours."
        if result.get("reused_hours") else ""
    )
    payloads = {}
    if "travel" in requested_services:
        payloads["travel"] = {
            "rows": [{"hour": item["hour"], "value": item["travel_minutes"]} for item in hourly],
            "note": "Automatically extracted from the A→C Scraper response at the saved A→B road distance." + reuse_note,
            "source": "outscraper", "updated": updated,
            "outscraper_job_id": job_id, "directions_result": result,
        }
    if "speed" in requested_services:
        payloads["speed"] = {
            "rows": [{"hour": item["hour"], "value": item["speed_kmh"]} for item in hourly],
            "note": "Automatically calculated as A→B distance / time and converted to km/h." + reuse_note,
            "source": "outscraper", "updated": updated,
            "outscraper_job_id": job_id, "directions_result": result,
        }
    return payloads

@app.post("/api/requests/<int:rid>/volumes/confirm")
@admin_required
def confirm_volumes_result(rid):
    workflow = db().execute(
        "SELECT * FROM volumes_workflows WHERE request_id=?", (rid,)
    ).fetchone()
    if not workflow:
        return jsonify(error="This request does not include Traffic volumes."), 404
    if workflow["status"] == "sent":
        return jsonify(error="These volume results have already been sent."), 409
    if workflow["status"] != "ready_for_review" or not workflow["result_json"]:
        return jsonify(error="The MATLAB result is not ready for review."), 409
    row = db().execute("SELECT * FROM requests WHERE id=?", (rid,)).fetchone()
    if not row:
        return jsonify(error="Source request not found."), 404
    result = json.loads(workflow["result_json"])
    outputs = result.get("hourly") or []
    requested_hours = sorted(set(json.loads(row["hours"] or "[]")))
    if [item.get("hour") for item in outputs] != requested_hours:
        return jsonify(error="The MATLAB result does not match the requested hours. Recalculate before sending."), 409
    if any(item.get("value") is None or item["value"] < 0 for item in outputs):
        return jsonify(error="The MATLAB result contains a negative volume. Check the inputs and recalculate before sending."), 409
    data = json.loads(row["data"] or "{}")
    volume_data = {
        "rows": [
            {"hour": item["hour"], "value": int(math.floor(item["value"] + 0.5)), "tolerance": round(item["tolerance"], 1)}
            for item in outputs
        ],
        "note": "Estimated by the MATLAB traffic model using the saved lanes, actual hourly speed, and weekly-weekend free-flow speed.",
        "source": "matlab",
        "updated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
    }
    freeflow_context = result.get("inputs", {}).get("freeflow_context")
    if freeflow_context:
        volume_data["freeflow_context"] = freeflow_context
    data["volumes"] = volume_data
    requested_services = set(json.loads(row["services"] or "[]"))
    _save_volume_lanes_service(data, requested_services, workflow["lanes"])
    # One approval sends everything that is ready: the reviewed Directions result (Speed / Travel)
    # goes out together with Traffic volumes and Lanes.
    directions_row = db().execute(
        "SELECT job_id,result_json FROM request_directions_workflows "
        "WHERE request_id=? AND status='ready' AND result_json IS NOT NULL "
        "ORDER BY updated_at DESC,created_at DESC LIMIT 1",
        (rid,),
    ).fetchone()
    directions_job_id = None
    if directions_row:
        directions_job_id = directions_row["job_id"]
        data.update(_directions_service_data(
            requested_services, json.loads(directions_row["result_json"]), directions_job_id
        ))
    not_ready = [
        SERVICES[service] for service in ("speed", "travel")
        if service in requested_services and not (data.get(service) or {}).get("rows")
    ]
    if not_ready:
        return jsonify(error="These results are not ready yet: " + ", ".join(not_ready) + ". Wait for the Directions result before sending."), 409
    all_services_ready = _all_requested_services_ready(requested_services, data)
    now = _timestamp_now()
    d = db()
    d.execute("BEGIN IMMEDIATE")
    current = d.execute(
        "SELECT status FROM volumes_workflows WHERE request_id=?", (rid,)
    ).fetchone()
    if not current or current["status"] != "ready_for_review":
        d.rollback()
        return jsonify(error="The volume workflow changed before confirmation. Reload and review the current status."), 409
    d.execute(
        "UPDATE requests SET data=?, status=CASE WHEN status='Rejected' THEN status "
        "WHEN ? THEN 'Completed' ELSE 'Processing' END WHERE id=?",
        (json.dumps(data, ensure_ascii=False), int(all_services_ready), rid),
    )
    d.execute(
        "UPDATE volumes_workflows SET status='sent', reviewed_by=?, reviewed_at=?, sent_at=?, updated_at=? "
        "WHERE request_id=? AND status='ready_for_review'",
        (session["uid"], now, now, now, rid),
    )
    if directions_job_id:
        d.execute(
            "UPDATE request_directions_workflows SET status='confirmed', confirmed_by=?, confirmed_at=?, updated_at=? "
            "WHERE request_id=? AND job_id=? AND status='ready'",
            (session["uid"], now, now, rid, directions_job_id),
        )
    sent_directions = bool(directions_job_id and {"speed", "travel"}.intersection(requested_services))
    _create_notification(
        row["user_id"], rid, "volume_result",
        "Request results are ready" if sent_directions else "Traffic volume results are ready",
        f"The administrator approved the results for request REQ-{rid}." if sent_directions
        else f"The administrator approved the traffic volume results for request REQ-{rid}.",
    )
    d.commit()
    return jsonify(
        ok=True,
        status="sent",
        request_status="Completed" if all_services_ready else "Processing",
        result_url=url_for("result", rid=rid),
    )

@app.post("/api/requests/<int:rid>/directions/confirm")
@admin_required
def confirm_request_directions(rid):
    payload = request.get_json(silent=True) or {}
    job_id = payload.get("job_id")
    workflow = db().execute(
        "SELECT * FROM request_directions_workflows WHERE request_id=? AND job_id=?",
        (rid, job_id),
    ).fetchone() if isinstance(job_id, str) else None
    if not workflow:
        return jsonify(error="No Scraper workflow is linked to this request."), 404
    if workflow["status"] == "confirmed":
        return jsonify(error="These results have already been confirmed and sent."), 409
    if workflow["status"] != "ready" or not workflow["result_json"]:
        return jsonify(error="The Scraper response is not complete and ready for review."), 409
    row = db().execute("SELECT * FROM requests WHERE id=?", (rid,)).fetchone()
    if not row:
        return jsonify(error="Source request not found."), 404
    result = json.loads(workflow["result_json"])
    requested_services = set(json.loads(row["services"] or "[]"))
    hourly = result["hourly"]
    data = json.loads(row["data"] or "{}")
    if "volumes" in requested_services:
        now = _timestamp_now()
        d = db()
        d.execute("BEGIN IMMEDIATE")
        current = d.execute(
            "SELECT status FROM request_directions_workflows WHERE request_id=? AND job_id=?",
            (rid, job_id),
        ).fetchone()
        volume_workflow = d.execute(
            "SELECT status,lanes,freeflow_result_json,freeflow_accepted FROM volumes_workflows WHERE request_id=?", (rid,)
        ).fetchone()
        if not current or current["status"] != "ready" or not volume_workflow:
            d.rollback()
            return jsonify(error="The volume or Directions workflow changed before confirmation. Reload and review the current status."), 409
        if volume_workflow["status"] not in {"running", "ready_for_review", "sent"}:
            d.execute(
                "UPDATE volumes_workflows SET actual_result_json=?, status='waiting_for_inputs', "
                "error=NULL, updated_at=? WHERE request_id=?",
                (json.dumps(result, ensure_ascii=False), now, rid),
            )
        updated = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
        if "travel" in requested_services:
            data["travel"] = {
                "rows": [{"hour": item["hour"], "value": item["travel_minutes"]} for item in hourly],
                "note": "Automatically extracted from the A→C Scraper response at the saved A→B road distance.",
                "source": "outscraper",
                "updated": updated,
                "outscraper_job_id": job_id,
                "directions_result": result,
            }
        if "speed" in requested_services:
            data["speed"] = {
                "rows": [{"hour": item["hour"], "value": item["speed_kmh"]} for item in hourly],
                "note": "Automatically calculated as A→B distance / time and converted to km/h.",
                "source": "outscraper",
                "updated": updated,
                "outscraper_job_id": job_id,
                "directions_result": result,
            }
        _save_volume_lanes_service(data, requested_services, volume_workflow["lanes"])
        all_services_ready = _all_requested_services_ready(requested_services, data)
        d.execute(
            "UPDATE requests SET data=?, status=CASE WHEN status='Rejected' THEN status "
            "WHEN ? THEN 'Completed' WHEN status!='Completed' THEN 'Processing' ELSE status END WHERE id=?",
            (json.dumps(data, ensure_ascii=False), int(all_services_ready), rid),
        )
        d.execute(
            "UPDATE request_directions_workflows SET status='confirmed', confirmed_by=?, confirmed_at=?, updated_at=? "
            "WHERE request_id=? AND job_id=? AND status='ready'",
            (session["uid"], now, now, rid, job_id),
        )
        if {"travel", "speed"}.intersection(requested_services):
            visible_services = [
                SERVICES[service] for service in ("travel", "speed")
                if service in requested_services
            ]
            _create_notification(
                row["user_id"], rid, "directions_result",
                "Request results are ready",
                f"The administrator approved {', '.join(visible_services)} for request REQ-{rid}. Traffic volumes are still being prepared.",
            )
        d.commit()
        if volume_workflow["status"] != "sent":
            if volume_workflow["freeflow_result_json"]:
                if _evaluate_volume_freeflow(rid):
                    _start_volume_calculation(rid)
            elif volume_workflow["freeflow_accepted"]:
                _start_volume_calculation(rid)
        return jsonify(
            ok=True, status="confirmed", confirmed_at=now,
            message=(
                "The approved directions results are available to the client; actual speeds were also saved internally for Traffic volumes."
                if {"travel", "speed"}.intersection(requested_services)
                else "Actual speed was saved internally for Traffic volumes; no unrequested service data was sent to the client."
            ),
        )
    updated = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    reuse_note = (
        " No separate result was returned for "
        + ", ".join(f"{hour:02d}:00" for hour in result.get("reused_hours", []))
        + "; the available completed result was reused for those hours."
        if result.get("reused_hours") else ""
    )
    if "travel" in requested_services:
        data["travel"] = {
            "rows": [{"hour": item["hour"], "value": item["travel_minutes"]} for item in hourly],
            "note": "Automatically extracted from the A→C Scraper response at the saved A→B road distance." + reuse_note,
            "source": "outscraper",
            "updated": updated,
            "outscraper_job_id": job_id,
            "directions_result": result,
        }
    if "speed" in requested_services:
        data["speed"] = {
            "rows": [{"hour": item["hour"], "value": item["speed_kmh"]} for item in hourly],
            "note": "Automatically calculated as A→B distance / time and converted to km/h." + reuse_note,
            "source": "outscraper",
            "updated": updated,
            "outscraper_job_id": job_id,
            "directions_result": result,
        }
    all_services_ready = _all_requested_services_ready(requested_services, data)
    now = _timestamp_now()
    d = db()
    d.execute("BEGIN IMMEDIATE")
    current = d.execute(
        "SELECT status FROM request_directions_workflows WHERE request_id=? AND job_id=?",
        (rid, job_id),
    ).fetchone()
    if not current or current["status"] != "ready":
        d.rollback()
        return jsonify(error="The workflow changed before confirmation. Reload and review the current status."), 409
    d.execute(
        "UPDATE requests SET data=?, status=CASE WHEN status!='Rejected' AND ? THEN 'Completed' "
        "WHEN status NOT IN ('Rejected','Completed') THEN 'Processing' ELSE status END WHERE id=?",
        (json.dumps(data, ensure_ascii=False), int(all_services_ready), rid),
    )
    d.execute(
        "UPDATE request_directions_workflows SET status='confirmed', confirmed_by=?, confirmed_at=?, updated_at=? "
        "WHERE request_id=? AND job_id=? AND status='ready'",
        (session["uid"], now, now, rid, job_id),
    )
    _create_notification(
        row["user_id"], rid, "directions_result",
        "Request results are ready",
        f"The administrator approved the directions results for request REQ-{rid}.",
    )
    d.commit()
    return jsonify(ok=True, status="confirmed", confirmed_at=now, result_url=url_for("result", rid=rid))

@app.route("/admin/outscraper/<job_id>")
@admin_required
def outscraper_job_detail(job_id):
    job = _outscraper_job(job_id)
    if not job: abort(404)
    token = session.get("outscraper_tokens", {}).get(job_id)
    if (job["approval_status"] != "pending" or job["execution_status"] != "awaiting_approval"
            or not token or not secrets.compare_digest(job["approval_token_hash"], _token_hash(token))):
        token = None
    job["approval_token"] = token
    for route in job["routes"]:
        route["result_pretty"] = json.dumps(route["result"], ensure_ascii=False, indent=2) if route["result"] is not None else ""
    request_workflow = db().execute(
        "SELECT request_id,status,error,confirmed_at FROM request_directions_workflows WHERE job_id=?",
        (job_id,),
    ).fetchone()
    return render_template("outscraper_detail.html", job=job, api_configured=outscraper.api_key_configured(),
                           request_workflow=request_workflow)

def _csv_safe_cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    elif value is None:
        value = ""
    else:
        value = str(value)
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value

def _result_records(value):
    if isinstance(value, list):
        for item in value:
            yield from _result_records(item)
    elif isinstance(value, dict):
        yield value
    elif value is not None:
        yield {"value": value}

@app.get("/admin/outscraper/<job_id>/results.csv")
@admin_required
def export_outscraper_results(job_id):
    job = _outscraper_job(job_id)
    if not job: abort(404)
    records = []
    for route in job["routes"]:
        response = route["result"] or {}
        context = job["config"].get("trafficlens_context") or {}
        for record in _result_records(response.get("data")):
            records.append({
                **record,
                "configured_route_index": route["route_index"],
                "configured_origin": route["origin"], "configured_destination": route["destination"],
                "provider_request_id": route["provider_request_id"] or "", "provider_status": route["provider_status"] or "",
                "requested_service": SERVICES.get(context.get("service"), ""),
                "study_date": context.get("study_date", ""),
                "reference_date_used": context.get("weekend_date") or context.get("holiday_date", ""),
                "reference_day_or_holiday": context.get("weekend_day") or context.get("holiday_name", ""),
                "holiday_date_used": context.get("holiday_date", ""),
                "holiday_name": context.get("holiday_name", ""),
                "route_country": context.get("country_name", ""),
                "route_timezone": context.get("timezone", ""),
                "departure_time_local": "07:00" if context else "",
            })
    fields = list(dict.fromkeys(key for record in records for key in record)) or [
        "configured_route_index", "configured_origin", "configured_destination", "provider_request_id", "provider_status"]
    output = io.StringIO(newline="")
    safe_fields = [_csv_safe_cell(field) for field in fields]
    writer = csv.DictWriter(output, fieldnames=safe_fields, extrasaction="ignore")
    writer.writeheader()
    for record in records:
        writer.writerow({_csv_safe_cell(key): _csv_safe_cell(value) for key, value in record.items()})
    return app.response_class(
        output.getvalue(), mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="outscraper-{job_id}.csv"'}
    )

@app.post("/api/requests/<int:rid>/data/<service>")
@admin_required
def save_service(rid, service):
    r = db().execute("SELECT * FROM requests WHERE id=?", (rid,)).fetchone()
    if not r or service not in forms_catalog() or service not in json.loads(r["services"]): return jsonify(error="Invalid"), 400
    if service == "volumes":
        return jsonify(error="Traffic volumes are managed by the MATLAB review workflow and cannot be edited manually."), 409
    j = request.get_json(force=True) or {}
    data = json.loads(r["data"] or "{}")
    saved = {"rows": clean_rows(service, j.get("rows")), "note": str(j.get("note") or "")[:5000], "source": "admin",
             "updated": dt.datetime.now().strftime("%Y-%m-%d %H:%M")}
    for context_key in ("holiday_context", "reference_context"):
        if data.get(service, {}).get(context_key):
            saved[context_key] = data[service][context_key]
            saved["outscraper_job_id"] = data[service].get("outscraper_job_id")
    data[service] = saved
    db().execute("UPDATE requests SET data=? WHERE id=?", (json.dumps(data), rid)); db().commit()
    return jsonify(ok=True, rows=len(data[service]["rows"]))

@app.post("/api/requests/<int:rid>/save-all")
@admin_required
def save_request_changes(rid):
    row = db().execute("SELECT * FROM requests WHERE id=?", (rid,)).fetchone()
    if not row: return jsonify(error="Request not found"), 404
    payload = request.get_json(force=True) or {}
    services = payload.get("services", {})
    if not isinstance(services, dict): return jsonify(error="Invalid service data"), 400
    requested_services = set(json.loads(row["services"] or "[]"))
    forms = forms_catalog()
    clean_services = {}
    for service, values in services.items():
        if service not in requested_services or service not in forms:
            return jsonify(error=f"Service is not part of this request: {service}"), 400
        if service == "volumes":
            return jsonify(error="Traffic volumes are managed by the MATLAB review workflow and cannot be edited manually."), 409
        if not isinstance(values, dict) or not isinstance(values.get("rows", []), list):
            return jsonify(error=f"Invalid data for service: {service}"), 400
        note = values.get("note", "")
        if not isinstance(note, str): return jsonify(error=f"Invalid note for service: {service}"), 400
        clean_services[service] = {
            "rows": clean_rows(service, values.get("rows", [])),
            "note": note[:5000],
            "source": "admin",
            "updated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        }

    status = payload.get("status")
    if status is not None and status not in STATUSES:
        return jsonify(error="Invalid request status"), 400
    reply = payload.get("reply")
    if reply is not None and not isinstance(reply, str):
        return jsonify(error="Invalid reply"), 400
    reply = reply.strip() if reply is not None else None
    point_c_visible = payload.get("point_c_visible")
    if point_c_visible is not None and (
        not isinstance(point_c_visible, bool)
        or (point_c_visible and not row["point_c"])
    ):
        return jsonify(error="Point C visibility must be true or false, and this request must have a point C"), 400

    data = json.loads(row["data"] or "{}")
    for service, values in clean_services.items():
        for context_key in ("holiday_context", "reference_context"):
            existing_context = (data.get(service) or {}).get(context_key)
            if existing_context:
                values[context_key] = existing_context
                values["outscraper_job_id"] = data[service].get("outscraper_job_id")
    data.update(clean_services)
    if status == "Completed" and not _all_requested_services_ready(requested_services, data):
        return jsonify(error="A request cannot be marked Completed until every selected service has a result ready for the client."), 409
    updates, params = [], []
    if clean_services:
        updates.append("data=?")
        params.append(json.dumps(data))
    if status is not None:
        updates.append("status=?")
        params.append(status)
    if reply is not None:
        updates.extend(("reply=?", "reply_at=?"))
        params.extend((reply or None, dt.datetime.now().strftime("%Y-%m-%d %H:%M") if reply else None))
    if point_c_visible is not None:
        updates.append("show_point_c=?")
        params.append(int(point_c_visible))
    if updates:
        params.append(rid)
        db().execute(f"UPDATE requests SET {', '.join(updates)} WHERE id=?", params)
        if reply is not None:
            _notify_request_owner_reply(rid, row["user_id"], row["reply"] or "", reply)
        db().commit()
    return jsonify(ok=True, saved_services=list(clean_services), status=status or row["status"])

@app.post("/api/requests/<int:rid>/status")
@admin_required
def set_status(rid):
    s = (request.get_json(force=True) or {}).get("status")
    row = db().execute(    "SELECT services,data,user_id FROM requests WHERE id=?", (rid,)).fetchone()
    if not row or s not in STATUSES: return jsonify(error="Invalid"), 400
    if s == "Completed" and not _all_requested_services_ready(
        json.loads(row["services"] or "[]"), json.loads(row["data"] or "{}")
    ):
        return jsonify(error="A request cannot be marked Completed until every selected service has a result ready for the client."), 409
    db().execute("UPDATE requests SET status=? WHERE id=?", (s, rid)); db().commit()
    return jsonify(ok=True, status=s)

@app.post("/api/requests/<int:rid>/reply")
@admin_required
def set_reply(rid):
    text = ((request.get_json(force=True) or {}).get("reply") or "").strip()
    row = db().execute("SELECT user_id,reply FROM requests WHERE id=?", (rid,)).fetchone()
    if not row: return jsonify(error="Not found"), 404
    db().execute("UPDATE requests SET reply=?, reply_at=? WHERE id=?", (text or None, dt.datetime.now().strftime("%Y-%m-%d %H:%M") if text else None, rid))
    _notify_request_owner_reply(rid, row["user_id"], row["reply"] or "", text)
    db().commit()
    return jsonify(ok=True)

@app.post("/api/users/<int:uid>/role")
@admin_required
def set_role(uid):
    role = (request.get_json(force=True) or {}).get("role")
    if role not in ("admin", "client") or not db().execute("SELECT 1 FROM users WHERE id=?", (uid,)).fetchone(): return jsonify(error="Invalid"), 400
    if uid == session["uid"]: return jsonify(error="You cannot change your own role"), 400
    db().execute("UPDATE users SET role=? WHERE id=?", (role, uid)); db().commit()
    return jsonify(ok=True, role=role)

@app.post("/api/services/<key>")
@admin_required
def save_service_info(key):
    catalog = service_catalog()
    if key not in catalog: return jsonify(error="Invalid"), 400
    text = str((request.get_json(force=True) or {}).get("description") or "").strip()[:500]
    if text:
        db().execute("INSERT INTO service_info(key,description) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET description=excluded.description", (key, text))
    else:
        db().execute("DELETE FROM service_info WHERE key=?", (key,))
    db().commit()
    return jsonify(ok=True, description=text or catalog[key]["description"])

@app.post("/api/services/<key>/availability")
@admin_required
def set_service_availability(key):
    catalog = service_catalog()
    if key not in catalog: return jsonify(error="Invalid service"), 404
    enabled = (request.get_json(force=True) or {}).get("enabled")
    if not isinstance(enabled, bool): return jsonify(error="Choose whether the service is enabled"), 400
    saved = get_setting("service_catalog", {}) or {}
    saved.setdefault(key, {})["enabled"] = enabled
    set_setting("service_catalog", saved)
    return jsonify(ok=True, enabled=enabled)

@app.post("/api/services/custom")
@admin_required
def create_custom_service():
    j = request.get_json(force=True) or {}
    name = str(j.get("name") or "").strip()
    unit = str(j.get("unit") or "").strip()
    kind = j.get("kind")
    if not name or len(name) > 80: return jsonify(error="Enter a service name (up to 80 characters)"), 400
    if len(unit) > 24: return jsonify(error="The unit must be 24 characters or fewer"), 400
    if kind not in ("static", "hourly"): return jsonify(error="Choose whether the service is per link or per hour"), 400
    saved = get_setting("service_catalog", {}) or {}
    key = "custom_" + os.urandom(8).hex()
    saved[key] = {"name": name, "unit": unit, "kind": kind, "enabled": True,
                  "description": "Data entered by the administrator."}
    set_setting("service_catalog", saved)
    return jsonify(ok=True, key=key, name=name)

@app.post("/api/network/upload")
@admin_required
def net_upload():
    if not network.HAVE_GEO: return jsonify(error="geopandas is not installed. Run: pip install geopandas pyogrio"), 400
    f = request.files.get("file")
    if not f or not f.filename.lower().endswith(network.EXTS): return jsonify(error="Upload a .zip (Shapefile), .geojson or .gpkg file"), 400
    os.makedirs(network.NET_DIR, exist_ok=True)
    for old in os.listdir(network.NET_DIR): os.remove(os.path.join(network.NET_DIR, old))
    path = os.path.join(network.NET_DIR, "network" + os.path.splitext(f.filename)[1].lower()); f.save(path)
    try: meta = network.inspect(path)
    except Exception as e:
        os.remove(path); return jsonify(error=f"Could not read the file: {e}"), 400
    old = get_setting("network", {})
    cfg = {"file": path, "filename": f.filename, "uploaded": dt.datetime.now().strftime("%Y-%m-%d %H:%M"), "tol": old.get("tol", 25), "min_cov": old.get("min_cov", 80),
           "speed_unit": old.get("speed_unit", "km/h"), **meta}
    guessed = network.guess(meta["columns"])
    for k in ("lanes", "speed", "capacity", "road_name"):
        cfg[k] = old.get(k) if old.get(k) in meta["columns"] else guessed.get(k, "")
    set_setting("network", cfg)
    return jsonify(ok=True, features=meta["features"])

@app.post("/api/network/config")
@admin_required
def net_config():
    cfg, j = get_setting("network", {}), request.get_json(force=True) or {}
    if not cfg.get("file"): return jsonify(error="Upload a network file first"), 400
    for k in ("lanes", "speed", "capacity", "road_name"):
        v = j.get(k) or ""
        if v and v not in cfg["columns"]: return jsonify(error=f"Unknown column: {v}"), 400
        cfg[k] = v
    cfg["speed_unit"] = j.get("speed_unit") if j.get("speed_unit") in network.UNITS else "km/h"
    try: cfg["tol"], cfg["min_cov"] = max(5, min(200, float(j.get("tol") or 25))), max(10, min(100, float(j.get("min_cov") or 80)))
    except (TypeError, ValueError): return jsonify(error="Distance and coverage must be numbers"), 400
    set_setting("network", cfg)
    return jsonify(ok=True, services=net_services())

@app.post("/api/requests/<int:rid>/network")
@admin_required
def rerun_network(rid):
    if not db().execute("SELECT 1 FROM requests WHERE id=?", (rid,)).fetchone(): return jsonify(error="Not found"), 404
    try: info = apply_network(rid)
    except Exception as e: return jsonify(error=f"Network lookup failed: {e}"), 400
    if info is None: return jsonify(error="No road network is configured yet (Admin > Network)"), 400
    return jsonify(ok=True, **info)

def _redact_client_directions_context(result_request):
    for service_data in result_request["data"].values():
        if not isinstance(service_data, dict):
            continue
        service_data.pop("outscraper_job_id", None)
        direction_result = service_data.get("directions_result")
        if isinstance(direction_result, dict):
            for key in ("source_request_id", "outscraper_job_id", "point_a", "point_b", "point_c"):
                direction_result.pop(key, None)

@app.route("/results/<int:rid>")
@login_required
def result(rid):
    r = db().execute("SELECT r.*, u.name AS client FROM requests r JOIN users u ON u.id=r.user_id WHERE r.id=?", (rid,)).fetchone()
    if not r or (r["user_id"] != session["uid"] and session["role"] != "admin"): abort(404)
    if r["group_id"] and session["role"] != "admin":
        grp = db().execute("SELECT sent_at FROM request_groups WHERE id=?", (r["group_id"],)).fetchone()
        if grp and not grp["sent_at"]:
            flash("This order is still being prepared. All its results will be sent together in one report.")
            return redirect(url_for("client_group", gid=r["group_id"]))
    if r["status"] == "Pending" and session["role"] != "admin": flash("Your request is still pending. Results will appear here soon."); return redirect(url_for("my_requests"))
    result_request = to_dict(r)
    if session.get("role") != "admin" and not result_request.get("show_point_c"):
        result_request["point_c"] = None
        result_request["point_c_distance_km"] = None
    if session.get("role") != "admin":
        _redact_client_directions_context(result_request)
    return render_template("result.html", r=result_request, forms=forms_catalog())

def _excel_safe_text(value):
    if value is None:
        return ""
    text = str(value)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text

def _style_excel_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="5632A8")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        column_letter = column[0].column_letter
        width = max((len(str(cell.value or "")) for cell in column), default=10)
        sheet.column_dimensions[column_letter].width = min(max(width + 2, 14), 48)

@app.get("/results/<int:rid>/export.xlsx")
@login_required
def export_result_excel(rid):
    row = db().execute(
        "SELECT r.*, u.name AS client FROM requests r JOIN users u ON u.id=r.user_id WHERE r.id=?",
        (rid,),
    ).fetchone()
    if not row or (row["user_id"] != session["uid"] and session["role"] != "admin"):
        abort(404)
    result_request = to_dict(row)
    if session.get("role") != "admin" and not result_request.get("show_point_c"):
        result_request["point_c"] = None
        result_request["point_c_distance_km"] = None
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Request Summary"
    summary.append(["Field", "Value"])
    for label, value in (
        ("Request ID", f"REQ-{result_request['id']}"),
        ("Client", result_request["client"]),
        ("Status", result_request["status"]),
        ("Study date", result_request["d1"]),
        ("Time range", result_request["time_range_label"]),
        ("Hourly periods", result_request["hours_label"]),
        ("Link", result_request["area"]),
        ("Road length (km)", result_request.get("road_km")),
        ("Requested services", ", ".join(SERVICES.get(service, service) for service in result_request["services"])),
        ("Point A", ", ".join(map(str, result_request["coords"][0])) if result_request["coords"] else ""),
        ("Point B", ", ".join(map(str, result_request["coords"][-1])) if result_request["coords"] else ""),
        ("Point C", ", ".join(map(str, result_request["point_c"])) if result_request.get("point_c") else ""),
        ("Admin reply", result_request.get("reply") or ""),
    ):
        summary.append([_excel_safe_text(label), value if isinstance(value, (int, float)) else _excel_safe_text(value)])
    directions_result = (
        (result_request["data"].get("travel") or result_request["data"].get("speed") or {}).get("directions_result")
    )
    if directions_result:
        for label, value in (
            ("Directions route", directions_result["route"]),
            ("A-to-B distance (m)", directions_result["distance_meters"]),
            ("A-to-B time (seconds)", directions_result["time_seconds"]),
            ("A-to-B speed (km/h)", directions_result["speed_kmh"]),
        ):
            summary.append([_excel_safe_text(label), value if isinstance(value, (int, float)) else _excel_safe_text(value)])
    for service in result_request["services"]:
        service_data = result_request["data"].get(service) or {}
        reference_context = (
            service_data.get("reference_context")
            or service_data.get("holiday_context")
            or service_data.get("freeflow_context")
        )
        if reference_context:
            service_name = SERVICES.get(service, service)
            weekend_reference = (
                reference_context.get("reference_type") == "weekly_weekend"
                or bool(reference_context.get("weekend_date"))
            )
            reference_date = reference_context.get("weekend_date") or reference_context.get("holiday_date", "")
            reference_label = reference_context.get("weekend_day") or reference_context.get("holiday_name", "")
            reference_kind = "weekly weekend used" if weekend_reference else "public holiday used"
            summary.append([_excel_safe_text(f"{service_name} study date"), _excel_safe_text(reference_context["study_date"])])
            summary.append([_excel_safe_text(f"{service_name} {reference_kind}"), _excel_safe_text(
                f"{reference_label} · {reference_date}"
            )])
            summary.append([_excel_safe_text(f"{service_name} route country / time zone"), _excel_safe_text(
                f'{reference_context["country_name"]} ({reference_context["country_code"]}) · {reference_context["timezone"]}'
            )])
            summary.append([_excel_safe_text(f"{service_name} departure"), _excel_safe_text("07:00 local time")])
    _style_excel_sheet(summary)

    details = workbook.create_sheet("Service Details")
    details.append(["Service", "Hour", "Value", "Tolerance", "Unit", "Source", "Notes", "Updated"])
    forms = forms_catalog()
    for service in result_request["services"]:
        definition = forms.get(service, {})
        saved = result_request["data"].get(service, {})
        rows = saved.get("rows") or []
        if rows:
            for service_row in rows:
                hour = service_row.get("hour")
                hour_label = f"{int(hour):02d}:00-{int(hour) + 1:02d}:00" if hour is not None else ""
                value = service_row.get("value")
                if isinstance(value, bool):
                    value = _excel_safe_text(value)
                elif not isinstance(value, (int, float)):
                    value = _excel_safe_text(value)
                tolerance = service_row.get("tolerance")
                if tolerance is not None and not isinstance(tolerance, (int, float)):
                    tolerance = _excel_safe_text(tolerance)
                details.append([
                    _excel_safe_text(SERVICES.get(service, service)), hour_label, value, tolerance,
                    _excel_safe_text(definition.get("unit", "")),
                    _excel_safe_text(
                        "Road network" if saved.get("source") == "network"
                        else "MATLAB" if saved.get("source") == "matlab"
                        else "Admin"
                    ),
                    _excel_safe_text(saved.get("note") or ""), _excel_safe_text(saved.get("updated") or ""),
                ])
        else:
            details.append([
                _excel_safe_text(SERVICES.get(service, service)), "", "", "",
                _excel_safe_text(definition.get("unit", "")),
                _excel_safe_text(
                    "Road network" if saved.get("source") == "network"
                    else "MATLAB" if saved.get("source") == "matlab"
                    else ""
                ),
                _excel_safe_text(saved.get("note") or ""), _excel_safe_text(saved.get("updated") or ""),
            ])
    _style_excel_sheet(details)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return app.response_class(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="trafficlens-request-{rid}.xlsx"'},
    )

init()   # also runs when started with "flask run" or gunicorn
if app.config["OUTSCRAPER_AUTO_POLL"]:
    with app.app_context():
        recoverable_jobs = db().execute(
            "SELECT DISTINCT job_id FROM outscraper_job_routes WHERE execution_status IN ('running','unknown') AND provider_request_id IS NOT NULL"
        ).fetchall()
    for recoverable in recoverable_jobs:
        _start_outscraper_monitor(recoverable["job_id"])

if __name__ == "__main__":
    app.run(debug=True)